#!/usr/bin/env python3
"""Offline tests for scripts/ppi_sa_probe.py.

Run:
    python3 scripts/ppi_sa_probe_test.py -v
"""

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

import ppi_sa_probe as probe


class HelpersTest(unittest.TestCase):
    def test_year_with_footnote(self):
        self.assertEqual(probe._year("20221)"), 2022)
        self.assertEqual(probe._year("2013"), 2013)
        self.assertEqual(probe._year(2026), 2026)
        self.assertIsNone(probe._year(""))

    def test_month_with_footnote(self):
        self.assertEqual(probe._month("январь1)"), 1)
        self.assertEqual(probe._month("Декабрь"), 12)
        self.assertIsNone(probe._month(""))

    def test_number(self):
        self.assertEqual(probe._num("−4,1"), -4.1)
        self.assertEqual(probe._num("8,1"), 8.1)

    def test_shift_month(self):
        self.assertEqual(probe._shift_month("2026-01-01", -1), "2025-12-01")
        self.assertEqual(probe._shift_month("2026-03-01", -3), "2025-12-01")
        self.assertEqual(probe._shift_month("2026-12-01", 1), "2027-01-01")

    def test_chain(self):
        points = [(date(2020, 1, 1), 100.0), (date(2020, 2, 1), 110.0), (date(2020, 3, 1), 99.0)]
        self.assertEqual([round(value, 4) for _, value in probe.chain(points)], [100.0, 110.0, 108.9])


class TaxonomyTest(unittest.TestCase):
    def test_food(self):
        self.assertTrue(probe.is_food("10.71.11.100"))
        self.assertTrue(probe.is_food("11.07.19"))
        self.assertFalse(probe.is_food("11.01.10"))
        self.assertFalse(probe.is_food("14.13.1"))

    def test_nonfood(self):
        self.assertTrue(probe.is_nonfood("13.10"))
        self.assertTrue(probe.is_nonfood("20.41"))
        self.assertFalse(probe.is_nonfood("19.20.21.100"))
        self.assertFalse(probe.is_nonfood("12.00.11.130"))
        self.assertFalse(probe.is_nonfood("10.71.11"))

    def test_nonfood_narrow(self):
        self.assertTrue(probe.is_nonfood_narrow("14.13.1"))
        self.assertTrue(probe.is_nonfood_narrow("31.09.12"))
        self.assertFalse(probe.is_nonfood_narrow("24.10"))
        self.assertFalse(probe.is_nonfood_narrow("28.11"))


class WorkbookParsersTest(unittest.TestCase):
    def _ved_months(self) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "2.1"
        ws.append(["", "Индексы цен производителей"])
        ws.append(["на конец периода, в %"])
        ws.append(["", "2020", "2021", "2022", "2023", "2024"])
        ws.append(["К предыдущему месяцу"])
        ws.append([" Январь", "100,5", "101,0", "102,0", "103,0", "104,0"])
        ws.append([" Февраль", "99,5", "102,0", "103,0", "104,0", "105,0"])
        ws.append(["Декабрь в % к декабрю предыдущего года", "100", "100", "100", "100", "100"])
        path = Path(tempfile.mkdtemp()) / "ved.xlsx"
        wb.save(path)
        return path

    def test_parse_ved_months(self):
        points = probe.parse_ved_months(self._ved_months(), "2.1")
        self.assertEqual(len(points), 10)
        self.assertEqual(points[0], (date(2020, 1, 1), 100.5))

    def _ved_matrix(self) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "8.1"
        ws.append(["", "Индексы цен производителей"])
        ws.append(["", "на конец периода, в %"])
        ws.append(["", "ОКВЭД 2", "январь", "февраль"])
        ws.append(["Промышленность", "1323500.029.31"])
        ws.append(["Российская Федерация", "", "100,5", "101,0"])
        path = Path(tempfile.mkdtemp()) / "ved8.xlsx"
        wb.save(path)
        return path

    def test_parse_ved_matrix(self):
        result = probe.parse_ved_matrix(self._ved_matrix(), "8.1")
        self.assertEqual(result["1323500.029.31"], {1: 100.5, 2: 101.0})

    def _products(self, with_rf: bool) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "4"
        ws.append(["", "Индексы цен производителей"])
        ws.append(["", "", "на конец периода, в %"])
        ws.append(["", "Код товара \nна основе ОКПД2", "январь", "февраль"])
        if with_rf:
            ws.append(["Товар", "10.71.11.100"])
            ws.append(["Российская Федерация", "", "100,5", "101,0"])
        else:
            ws.append(["Товар", "10.71.11.100", "101,0", "102,0"])
        path = Path(tempfile.mkdtemp()) / "tov.xlsx"
        wb.save(path)
        return path

    def test_parse_products_product_row(self):
        points = probe.parse_products(self._products(False), "4")
        self.assertEqual(len(points), 2)
        self.assertEqual(points[0]["code"], "10.71.11.100")

    def test_parse_products_rf_row(self):
        points = probe.parse_products(self._products(True), "4")
        self.assertEqual(len(points), 2)
        self.assertEqual(points[0]["value"], 100.5)


if __name__ == "__main__":
    unittest.main()
