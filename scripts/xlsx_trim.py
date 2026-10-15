#!/usr/bin/env python3
"""Урезает XLSX-фикстуру до объявленного набора листов и проверяет инвариант §11.2.

Зачем zip-хирургия, а не библиотека
-----------------------------------
`docs/IMPORTER_PACKAGE_STANDARD.md` §11.1 требует, чтобы фикстура была **реальным
урезанным артефактом источника**, а §11.2 — чтобы «урезание фикстуры не меняло
разметку: те же листы, шапки, якоря, форматы ячеек». Отсюда два запрета:

* синтетические книги запрещены — синтетика кодирует допущения автора парсера
  и остаётся зелёной ровно тогда, когда парсер разошёлся с источником;
* пересборка книги сторонней библиотекой (`openpyxl`, `excelize`) запрещена —
  они переписывают XML листов, и разметка перестаёт быть разметкой источника.

XLSX — это zip-контейнер OPC (ECMA-376 Part 2). Скрипт делает **хирургию
архива**: XML оставляемых листов копируется **побайтово как есть**, а из
служебных файлов вычищаются записи об удаляемых листах. Так сохраняются и
разметка, и форматы ячеек, и подписи колонок.

Что правится в служебных файлах
--------------------------------
1. `xl/workbook.xml` — список `<sheet>` сужается до оставленных листов;
   `<definedNames>` фильтруется (правило ниже); блок `<externalReferences>`
   уходит целиком вместе с частями `xl/externalLinks/**`.
2. `xl/_rels/workbook.xml.rels` — уходят отношения удаляемых листов, внешних
   книг и `calcChain`.
3. `[Content_Types].xml` — остаются все `Default` плюс `Override` только тех
   частей, которые реально остались в архиве.
4. `xl/calcChain.xml` (если был) — удаляется: цепочка пересчёта ссылается на
   удалённые листы, а Excel и `excelize` её не требуют.
5. `<workbookView activeTab firstSheet>` — ссылки на удалённые листы снимаются,
   ссылки на оставленные перенумеровываются; без этого `openpyxl.wb.active`
   возвращает `None`, потому что индекс указывает за пределы списка листов.

Правило фильтрации `<definedNames>`: остаётся только имя, тело которого
ссылается **и только** на оставленные листы, а `localSheetId` перенумеровывается
под новый порядок листов. Причина: в бюллетенях ЦБ `workbook.xml` несёт десятки
тысяч имён от чужого шаблона (`#REF!`, именованные константы-массивы, ссылки на
внешние книги `[N]справочник!`), их тела не относятся ни к одному оставленному
листу, а весу дают мегабайты. Если же на оставляемых листах есть формулы, скрипт
печатает предупреждение: кэшированные значения при этом не меняются (значения
лежат в XML листа, который копируется побайтово), но живые формулы на имена и на
внешние книги опираться уже не будут.

Воспроизводимость
-----------------
Архив пишется самим скриптом, без shell-`zip`: `[Content_Types].xml` первым
элементом (требование OPC), `ZIP_DEFLATED`, без записей каталогов, время всех
записей прибито к 1980-01-01, внешние атрибуты — как у `zip -X -D`. Два прогона
на одном входе дают побайтово одинаковый файл.

Запуск (из корня репозитория)
-----------------------------
    # урезать; --report пишет JSON-отчёт о составе архива
    python3 scripts/xlsx_trim.py trim \\
        /tmp/dip/f_62343_bulletin_26-06.xlsx cbr/testdata/dip_bulletin_26-06.xlsx \\
        --keep 2,23 --report /tmp/trim_26-06.json

    # проверить готовую фикстуру против исходника
    python3 scripts/xlsx_trim.py verify \\
        /tmp/dip/f_62343_bulletin_26-06.xlsx cbr/testdata/dip_bulletin_26-06.xlsx \\
        --keep 2,23

`verify` сверяет: (а) XML оставленных листов побайтово совпадает с исходником;
(б) архив читается и все `Override`/отношения резолвятся в существующие части;
(в) дамп значений оставленных листов (`values_only`) совпадает с исходником —
последняя проверка делается через `openpyxl` и **пропускается**, если его нет
(конвейер от `openpyxl` не зависит). Любое расхождение — ненулевой код возврата.

Когда скрипт нужен снова
-------------------------
Каждый раз, когда объявленный набор листов (рядов) импортёра `cbr_dip_bulletin`
расширяется: тот же исходник ЦБ урезается тем же скриптом с новым `--keep`, и
фикстура остаётся воспроизводимой. Обзор листов, подписей и дат в скачанных
выпусках (чем выбирать, какие листы добавлять) — `scripts/xlsx_inspect.py`.

Чего скрипт НЕ делает
---------------------
* не скачивает выпуски с сайта ЦБ — это работа импортёра и `make`;
* не ходит в сеть, не пишет в ClickHouse, не меняет файлы-источники;
* не правит XML самих листов — они копируются побайтово;
* не «санирует» метаданные книги (`docProps/core.xml`, `x15ac:absPath`): они
  остаются такими, какими их опубликовал ЦБ.

Только стандартная библиотека (Python 3.9).
"""

import argparse
import hashlib
import json
import os
import posixpath
import re
import sys
import zipfile

# Часть, которая обязана лежать первым элементом архива (требование OPC).
FIRST_PART = "[Content_Types].xml"

# Служебные части, которые скрипт переписывает; всё остальное копируется байт в байт.
REWRITTEN_PARTS = (FIRST_PART, "xl/workbook.xml", "xl/_rels/workbook.xml.rels")

# Отношения уровня книги, которые остаются; остальные типы удаляются.
WB_KEEP_TYPES = ("worksheet", "styles", "sharedStrings", "theme", "customXml")

# Отношения уровня книги, которые удаляются осознанно (а не как «неизвестный тип»).
WB_DROP_TYPES = ("externalLink", "calcChain")

# Метаданные записей архива — ровно то, что писал `zip -X -D`: без extra-полей,
# без записей каталогов, права 0o644 в старших 16 битах external_attr.
ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)
ZIP_EXTERNAL_ATTR = 0o100644 << 16
ZIP_INTERNAL_ATTR = 0
ZIP_CREATE_SYSTEM = 3
ZIP_CREATE_VERSION = 30
ZIP_EXTRACT_VERSION = 20


class TrimError(Exception):
    """Ошибка урезания: архив писать нельзя."""


# --------------------------------------------------------------------------
# OPC: отношения и пути
# --------------------------------------------------------------------------


def rels_part_for(part):
    """Часть с отношениями: `xl/workbook.xml` -> `xl/_rels/workbook.xml.rels`."""
    folder, name = part.rsplit("/", 1)
    return folder + "/_rels/" + name + ".rels"


def resolve_target(source_part, target):
    """Цель отношения OPC -> путь части в архиве.

    Цели относительны каталога исходной части; у пакетных отношений
    (`_rels/.rels`) исходная часть — корень пакета.
    """
    if source_part == "_rels/.rels" or target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    folder = source_part.rsplit("/", 1)[0]
    return posixpath.normpath(posixpath.join(folder, target))


def rels_source_folder(rels_part):
    """Каталог, относительно которого считаются цели файла `.rels`."""
    if rels_part == "_rels/.rels":
        return ""
    folder, _ = rels_part.rsplit("/_rels/", 1)
    return folder


def parse_rels(xml):
    """Разбор `<Relationships>` в список словарей (порядок файла сохраняется)."""
    out = []
    for match in re.finditer(r"<Relationship\b[^>]*/>", xml):
        tag = match.group(0)
        out.append(
            {
                "raw": tag,
                "rid": re.search(r'Id="([^"]*)"', tag).group(1),
                "rtype": re.search(r'Type="([^"]*)"', tag).group(1),
                "target": re.search(r'Target="([^"]*)"', tag).group(1),
                "external": 'TargetMode="External"' in tag,
            }
        )
    return out


def short_type(rtype):
    """`.../relationships/worksheet` -> `worksheet`."""
    return rtype.rsplit("/", 1)[-1]


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def sheet_parts(zf):
    """Имя листа -> часть с его XML (через `xl/workbook.xml` и его отношения)."""
    workbook = zf.read("xl/workbook.xml").decode("utf-8")
    rels = {r["rid"]: r for r in parse_rels(zf.read("xl/_rels/workbook.xml.rels").decode("utf-8"))}
    out = {}
    for tag in re.findall(r"<sheet\b[^>]*/>", workbook):
        name = re.search(r'name="([^"]*)"', tag).group(1)
        rid = re.search(r'r:id="([^"]*)"', tag).group(1)
        if rid in rels:
            out[name] = resolve_target("xl/workbook.xml", rels[rid]["target"])
    return out


def parse_keep(raw_values):
    """`["2,23"]` -> `["2", "23"]`; поддерживает и `--keep 2 --keep 23`."""
    names = []
    for raw in raw_values:
        for name in raw.split(","):
            name = name.strip()
            if name and name not in names:
                names.append(name)
    if not names:
        raise TrimError("--keep пуст: не указано ни одного листа")
    return names


# --------------------------------------------------------------------------
# Урезание
# --------------------------------------------------------------------------


def _body_sheet_refs(body):
    """Имена листов из тела `<definedName>` плюс признак внешней книги `[N]`."""
    refs = []
    for match in re.finditer(r"(?:'([^']*)'|([^\s!'\"\[\],\(\)\{\};:]+))!", body):
        name = match.group(1) if match.group(1) is not None else match.group(2)
        name = name.strip("'")
        if "[" in name:
            name = re.sub(r"^\[[^\]]*\]", "", name)
        if name:
            refs.append(name)
    return refs, ("[" in body)


def plan_trim(zf, src_path, keep):
    """Разбирает открытый архив и возвращает (отчёт, {часть: байты} нового архива)."""
    names = set(zf.namelist())
    if "xl/workbook.xml" not in names or FIRST_PART not in names:
        raise TrimError(src_path + ": нет xl/workbook.xml или " + FIRST_PART + " — это не xlsx")

    report = {
        "src": src_path,
        "src_size": os.path.getsize(src_path),
        "src_parts": len(names),
        "kept_sheets": list(keep),
    }

    workbook = zf.read("xl/workbook.xml").decode("utf-8")
    wb_rels_xml = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    ct_xml = zf.read(FIRST_PART).decode("utf-8")
    wb_rels = parse_rels(wb_rels_xml)

    sheet_tags = re.findall(r"<sheet\b[^>]*/>", workbook)
    sheet_names = [re.search(r'name="([^"]*)"', tag).group(1) for tag in sheet_tags]
    report["all_sheets"] = sheet_names

    missing = [name for name in keep if name not in sheet_names]
    if missing:
        raise TrimError(
            "в " + src_path + " нет листов " + repr(missing) + "; есть: " + repr(sheet_names)
        )
    parts_of_sheets = sheet_parts(zf)
    unresolvable = [name for name in keep if name not in parts_of_sheets]
    if unresolvable:
        raise TrimError("не удалось найти часть листа для " + repr(unresolvable))

    old_index = {name: i for i, name in enumerate(sheet_names)}
    kept_old_index = set(old_index[name] for name in keep)
    new_index_of_old = {old: i for i, old in enumerate(sorted(kept_old_index))}
    sheet_rid_of_name = {
        re.search(r'name="([^"]*)"', tag).group(1): re.search(r'r:id="([^"]*)"', tag).group(1)
        for tag in sheet_tags
    }
    kept_rids = set(sheet_rid_of_name[name] for name in keep)

    # ---------------------------------------------------------------- состав архива
    retained = set()
    report["missing_targets"] = []

    def add(part, source):
        if part in retained:
            return
        if part not in names:
            report["missing_targets"].append({"part": part, "referenced_by": source})
            return
        retained.add(part)
        rp = rels_part_for(part)
        if rp in names:
            retained.add(rp)
            for rel in parse_rels(zf.read(rp).decode("utf-8")):
                if not rel["external"]:
                    add(resolve_target(part, rel["target"]), rp + " -> " + rel["target"])

    # корневые отношения пакета: сама книга и свойства документа
    retained.add(FIRST_PART)
    retained.add("_rels/.rels")
    for rel in parse_rels(zf.read("_rels/.rels").decode("utf-8")):
        if rel["external"]:
            continue
        target = resolve_target("_rels/.rels", rel["target"])
        if short_type(rel["rtype"]) == "officeDocument":
            retained.add(target)
            retained.add(rels_part_for(target))
        else:
            add(target, "_rels/.rels")

    # отношения уровня книги
    kept_wb_rels = []
    dropped_wb_rels = []
    for rel in wb_rels:
        kind = short_type(rel["rtype"])
        if kind in WB_DROP_TYPES or kind not in WB_KEEP_TYPES:
            dropped_wb_rels.append([kind, rel["target"]])
            continue
        if kind == "worksheet" and rel["rid"] not in kept_rids:
            dropped_wb_rels.append([kind, rel["target"]])
            continue
        add(resolve_target("xl/workbook.xml", rel["target"]), "xl/_rels/workbook.xml.rels")
        kept_wb_rels.append(rel)
    report["dropped_wb_rels"] = dropped_wb_rels
    kept_rids_set = set(rel["rid"] for rel in kept_wb_rels)

    # ---------------------------------------------------------------- <sheets> и <definedNames>
    new_sheet_tags = [
        tag for tag in sheet_tags if re.search(r'name="([^"]*)"', tag).group(1) in keep
    ]
    dn_block = re.search(r"<definedNames>.*?</definedNames>", workbook, re.S)
    dn_items = re.findall(
        r"<definedName\b[^>]*>.*?</definedName>|<definedName\b[^>]*/>", workbook, re.S
    )
    dn_stat = {
        "total": len(dn_items),
        "kept": 0,
        "dropped_scope": 0,
        "dropped_broken": 0,
        "dropped_external": 0,
        "dropped_noref": 0,
        "dropped_othersheet": 0,
    }

    kept_dn = []
    for item in dn_items:
        attrs = re.match(r"<definedName\b([^>]*)>", item).group(1)
        body = re.sub(r"^<definedName\b[^>]*>|</definedName>$", "", item)
        local_id = re.search(r'localSheetId="(\d+)"', attrs)
        refs, has_external = _body_sheet_refs(body)
        if has_external:
            dn_stat["dropped_external"] += 1
            continue
        if body.strip() == "#REF!":
            dn_stat["dropped_broken"] += 1
            continue
        if local_id and int(local_id.group(1)) not in kept_old_index:
            dn_stat["dropped_scope"] += 1
            continue
        if not refs:
            dn_stat["dropped_noref"] += 1
            continue
        if [ref for ref in refs if ref not in keep]:
            dn_stat["dropped_othersheet"] += 1
            continue
        if local_id:
            item = re.sub(
                r'localSheetId="\d+"',
                'localSheetId="%d"' % new_index_of_old[int(local_id.group(1))],
                item,
            )
        kept_dn.append(item)
        dn_stat["kept"] += 1
    report["definedNames"] = dn_stat

    # ---------------------------------------------------------------- переписывание служебного
    new_workbook = re.sub(
        r"<sheets[^>]*>.*?</sheets>",
        "<sheets>" + "".join(new_sheet_tags) + "</sheets>",
        workbook,
        flags=re.S,
    )

    dropped_bookview = []

    def fix_workbook_view(match):
        """activeTab/firstSheet нумеруют листы по порядку; индекс удалённого листа снимаем."""
        tag = match.group(0)
        for attr in ("activeTab", "firstSheet"):
            found = re.search(attr + r'="(\d+)"', tag)
            if not found:
                continue
            old = int(found.group(1))
            if old in new_index_of_old:
                tag = tag.replace(found.group(0), '%s="%d"' % (attr, new_index_of_old[old]))
            else:
                tag = tag.replace(" " + found.group(0), "")
                dropped_bookview.append("%s=%d" % (attr, old))
        return re.sub(r"\s+", " ", tag).replace(" >", ">")

    new_workbook = re.sub(r"<workbookView\b[^>]*/>", fix_workbook_view, new_workbook)
    if dropped_bookview:
        report["dropped_bookview_attrs"] = dropped_bookview

    if dn_block:
        replacement = ""
        if kept_dn:
            replacement = "<definedNames>" + "".join(kept_dn) + "</definedNames>"
        new_workbook = new_workbook.replace(dn_block.group(0), replacement)
    external_block = re.search(r"<externalReferences>.*?</externalReferences>", new_workbook, re.S)
    if external_block:
        new_workbook = new_workbook.replace(external_block.group(0), "")
    report["removed_blocks"] = {
        "definedNames": bool(dn_block),
        "definedNames_kept_elements": len(kept_dn),
        "externalReferences": bool(external_block),
        "externalReferences_count": len(re.findall(r"<externalReference\b", workbook)),
    }

    new_wb_rels = re.sub(
        r"<Relationship\b[^>]*/>",
        lambda m: m.group(0) if re.search(r'Id="([^"]*)"', m.group(0)).group(1) in kept_rids_set else "",
        wb_rels_xml,
    )

    ct_head = ct_xml[: ct_xml.index(">", ct_xml.index("<Types")) + 1]
    ct_defaults = re.findall(r"<Default\b[^>]*/>", ct_xml)
    ct_overrides_kept = []
    ct_overrides_dropped = 0
    for override in re.findall(r"<Override\b[^>]*/>", ct_xml):
        part_name = re.search(r'PartName="([^"]*)"', override).group(1).lstrip("/")
        if part_name in retained:
            ct_overrides_kept.append(override)
        else:
            ct_overrides_dropped += 1
    new_ct = ct_head + "".join(ct_defaults) + "".join(ct_overrides_kept) + "</Types>"
    report["content_types"] = {
        "overrides_total": len(ct_overrides_kept) + ct_overrides_dropped,
        "overrides_kept": len(ct_overrides_kept),
        "overrides_dropped": ct_overrides_dropped,
    }

    # ---------------------------------------------------------------- части нового архива
    rewritten = {
        "xl/workbook.xml": new_workbook.encode("utf-8"),
        "xl/_rels/workbook.xml.rels": new_wb_rels.encode("utf-8"),
        FIRST_PART: new_ct.encode("utf-8"),
    }
    parts = {}
    for part in retained:
        parts[part] = rewritten.get(part, zf.read(part))

    # ---------------------------------------------------------------- инвариант §11.2
    report["sheet_parts"] = {}
    report["sheet_identity"] = {}
    formulas = 0
    for name in keep:
        part = parts_of_sheets[name]
        src_bytes = zf.read(part)
        out_bytes = parts[part]
        found = re.findall(r"<f\b[^>]*>(.*?)</f>", out_bytes.decode("utf-8"), re.S)
        formulas += len(found)
        report["sheet_parts"][name] = part
        report["sheet_identity"][name] = {
            "part": part,
            "bytes": len(src_bytes),
            "sha256": sha256(src_bytes),
            "identical": src_bytes == out_bytes,
        }
        if found:
            report.setdefault("sheets_with_formulas", {})[name] = len(found)
    report["formulas_in_kept_sheets"] = formulas
    broken = [name for name in keep if not report["sheet_identity"][name]["identical"]]
    if broken:
        raise TrimError("инвариант §11.2 нарушен: XML листов изменился — " + repr(broken))

    dropped = sorted(names - retained)
    report["retained_parts"] = len(retained)
    report["dropped_parts"] = len(dropped)
    by_dir = {}
    for part in dropped:
        folder = re.sub(r"[^/]*$", "", part)
        by_dir[folder] = by_dir.get(folder, 0) + 1
    report["dropped_by_dir"] = dict(sorted(by_dir.items()))
    return report, parts


def write_zip(out_path, parts):
    """Пишет архив детерминированно: `[Content_Types].xml` первым, остальное по имени."""
    order = [FIRST_PART] + sorted(part for part in parts if part != FIRST_PART)
    folder = os.path.dirname(os.path.abspath(out_path))
    if folder and not os.path.isdir(folder):
        os.makedirs(folder)
    tmp_path = out_path + ".tmp"
    try:
        with zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for part in order:
                info = zipfile.ZipInfo(part, date_time=ZIP_DATE_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = ZIP_CREATE_SYSTEM
                info.create_version = ZIP_CREATE_VERSION
                info.extract_version = ZIP_EXTRACT_VERSION
                info.external_attr = ZIP_EXTERNAL_ATTR
                info.internal_attr = ZIP_INTERNAL_ATTR
                zf.writestr(info, parts[part])
        os.replace(tmp_path, out_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def cmd_trim(args):
    keep = parse_keep(args.keep)
    with zipfile.ZipFile(args.src) as zf:
        report, parts = plan_trim(zf, args.src, keep)
    write_zip(args.out, parts)

    with open(args.out, "rb") as fh:
        out_bytes = fh.read()
    report["out"] = args.out
    report["out_size"] = len(out_bytes)
    report["out_parts"] = len(parts)
    report["out_sha256"] = sha256(out_bytes)
    report["rewritten_parts"] = sorted(REWRITTEN_PARTS)
    report["zip_contract"] = {
        "first_entry": FIRST_PART,
        "compression": "ZIP_DEFLATED",
        "directory_entries": False,
        "date_time": "1980-01-01T00:00:00",
        "external_attr": hex(ZIP_EXTERNAL_ATTR),
    }

    print("вход:  %s (%d Б, %d частей, листов %d)"
          % (args.src, report["src_size"], report["src_parts"], len(report["all_sheets"])))
    print("листы: оставлены %s -> изъято %d"
          % (report["kept_sheets"], len(report["all_sheets"]) - len(keep)))
    print("части: %d -> %d (удалено %d)"
          % (report["src_parts"], report["out_parts"], report["dropped_parts"]))
    blocks = report["removed_blocks"]
    print("служебное: definedNames %d -> %d; <externalReference> %d -> 0; calcChain в книге %s"
          % (report["definedNames"]["total"], blocks["definedNames_kept_elements"],
             blocks["externalReferences_count"],
             "был, удалён" if [r for r in report["dropped_wb_rels"] if r[0] == "calcChain"] else "не было"))
    if report.get("dropped_bookview_attrs"):
        print("служебное: сняты атрибуты закладки %s" % ", ".join(report["dropped_bookview_attrs"]))
    for name in report["kept_sheets"]:
        item = report["sheet_identity"][name]
        print("  лист %-4s %s: %d Б, sha256=%s, побайтово=%s"
              % (repr(name), item["part"], item["bytes"], item["sha256"][:16], item["identical"]))
    if report["formulas_in_kept_sheets"]:
        print("ПРЕДУПРЕЖДЕНИЕ: на оставляемых листах %d формул (%s) — кэш значений не меняется, "
              "но имена и внешние книги из книги удалены"
              % (report["formulas_in_kept_sheets"], report.get("sheets_with_formulas")))
    if report["missing_targets"]:
        print("ПРЕДУПРЕЖДЕНИЕ: %d целей отношений нет в архиве: %s"
              % (len(report["missing_targets"]), report["missing_targets"]))
    print("выход: %s (%d Б, %d частей, sha256=%s)"
          % (args.out, report["out_size"], report["out_parts"], report["out_sha256"]))
    print("инвариант §11.2: XML оставленных листов побайтово совпал с исходником")

    if args.report:
        with open(args.report, "w", encoding="utf-8") as fh:
            json.dump(report, fh, ensure_ascii=False, indent=1, sort_keys=True)
            fh.write("\n")
        print("отчёт: %s" % args.report)
    return 0


# --------------------------------------------------------------------------
# Проверка
# --------------------------------------------------------------------------


def _values_dump(path, sheet_names):
    """Дамп значений листов через openpyxl; None, если openpyxl не установлен."""
    try:
        import openpyxl
    except ImportError:
        return None
    dump = {}
    for data_only in (True, False):
        book = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
        try:
            for name in sheet_names:
                if name not in book.sheetnames:
                    dump[(data_only, name)] = None
                    continue
                dump[(data_only, name)] = [tuple(row) for row in book[name].iter_rows(values_only=True)]
        finally:
            book.close()
    return dump


def cmd_verify(args):
    keep = parse_keep(args.keep)
    problems = []

    def fail(message):
        problems.append(message)
        print("      FAIL " + message)

    print("проверка: %s" % args.src)
    print("     против: %s" % args.out)
    print("     листы: %s" % keep)
    for path in (args.src, args.out):
        if not os.path.exists(path):
            print("verify: FAIL — файл не найден: %s" % path)
            return 1

    with zipfile.ZipFile(args.src) as src_zip, zipfile.ZipFile(args.out) as out_zip:
        out_names = set(out_zip.namelist())

        # (а) главный инвариант §11.2: XML оставленных листов не изменился
        print("  (а) побайтовое совпадение XML оставленных листов")
        src_parts = sheet_parts(src_zip)
        out_parts = sheet_parts(out_zip)
        if list(out_parts.keys()) != keep:
            fail("в фикстуре листы %s, ожидались %s" % (list(out_parts.keys()), keep))
        for name in keep:
            if name not in src_parts:
                fail("в исходнике нет листа %s" % repr(name))
                continue
            if name not in out_parts:
                fail("в фикстуре нет листа %s" % repr(name))
                continue
            if src_parts[name] != out_parts[name]:
                fail("лист %s: в исходнике %s, в фикстуре %s"
                     % (repr(name), src_parts[name], out_parts[name]))
                continue
            src_bytes = src_zip.read(src_parts[name])
            out_bytes = out_zip.read(out_parts[name])
            if src_bytes != out_bytes:
                fail("лист %s (%s): XML разошёлся — %d Б против %d Б"
                     % (repr(name), out_parts[name], len(src_bytes), len(out_bytes)))
            else:
                print("      OK лист %-4s %s: %d Б, sha256=%s"
                      % (repr(name), out_parts[name], len(src_bytes), sha256(src_bytes)[:16]))

        # (б) архив читается, ссылки и content types резолвятся
        print("  (б) целостность архива и ссылок")
        broken = out_zip.testzip()
        if broken is None:
            print("      OK zipfile.testzip() -> None (%d частей)" % len(out_names))
        else:
            fail("битая часть архива: %s" % broken)

        content_types = out_zip.read(FIRST_PART).decode("utf-8")
        overrides = [m.group(1).lstrip("/") for m in re.finditer(r'PartName="([^"]*)"', content_types)]
        lost = [part for part in overrides if part not in out_names]
        if lost:
            fail("Override без части: %s" % lost)
        else:
            print("      OK все %d Override из %s резолвятся" % (len(overrides), FIRST_PART))

        dangling = []
        checked = 0
        for part in sorted(name for name in out_names if name.endswith(".rels")):
            folder = rels_source_folder(part)
            for rel in parse_rels(out_zip.read(part).decode("utf-8")):
                if rel["external"]:
                    continue
                if rel["target"].startswith("/"):
                    resolved = posixpath.normpath(rel["target"].lstrip("/"))
                else:
                    resolved = posixpath.normpath(posixpath.join(folder, rel["target"]))
                checked += 1
                if resolved not in out_names:
                    dangling.append(part + " -> " + rel["target"])
        if dangling:
            fail("отношения в никуда: %s" % dangling)
        else:
            print("      OK все %d внутренних отношений резолвятся" % checked)

    # (в) значения ячеек
    print("  (в) дамп значений оставленных листов")
    src_dump = _values_dump(args.src, keep)
    if src_dump is None:
        print("      values-сравнение пропущено: нет openpyxl (pip install openpyxl)")
    else:
        out_dump = _values_dump(args.out, keep)
        for key in sorted(src_dump, key=lambda k: (k[1], k[0])):
            data_only, name = key
            rows_src, rows_out = src_dump[key], out_dump.get(key)
            if rows_src is None:
                fail("дамп исходника пуст для листа %s" % repr(name))
                continue
            if rows_out is None:
                fail("не удалось прочитать лист %s из фикстуры" % repr(name))
                continue
            cells_src = sum(1 for row in rows_src for cell in row if cell is not None)
            cells_out = sum(1 for row in rows_out for cell in row if cell is not None)
            if rows_src == rows_out:
                print("      OK data_only=%-5s лист %-4s: строк %d/%d, непустых ячеек %d/%d"
                      % (data_only, repr(name), len(rows_src), len(rows_out), cells_src, cells_out))
            else:
                where = "число строк %d против %d" % (len(rows_src), len(rows_out))
                for i, pair in enumerate(zip(rows_src, rows_out), 1):
                    if pair[0] != pair[1]:
                        where = "строка %d: %r против %r" % (i, pair[0], pair[1])
                        break
                fail("data_only=%s лист %s разошёлся: %s" % (data_only, repr(name), where))

    print("verify: %s" % ("PASS" if not problems else "FAIL (%d)" % len(problems)))
    return 0 if not problems else 1


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Урезание XLSX-фикстур zip-хирургией (docs/IMPORTER_PACKAGE_STANDARD.md §11.2).",
    )
    subparsers = parser.add_subparsers(dest="command")

    trim = subparsers.add_parser("trim", help="урезать книгу до листов --keep и записать архив")
    trim.add_argument("src", help="исходный xlsx, скачанный с сайта ЦБ (не меняется)")
    trim.add_argument("out", help="куда записать урезанную фикстуру")
    trim.add_argument("--keep", action="append", required=True,
                      help="имена оставляемых листов через запятую, например 2,23")
    trim.add_argument("--report", default=None, help="файл для JSON-отчёта о составе архива")
    trim.set_defaults(func=cmd_trim)

    verify = subparsers.add_parser("verify", help="проверить фикстуру против исходника")
    verify.add_argument("src", help="исходный xlsx")
    verify.add_argument("out", help="урезанная фикстура")
    verify.add_argument("--keep", action="append", required=True,
                        help="имена оставленных листов через запятую, например 2,23")
    verify.set_defaults(func=cmd_verify)

    args = parser.parse_args(argv)
    if not getattr(args, "func", None):
        parser.print_help()
        return 2
    try:
        return args.func(args)
    except TrimError as exc:
        print("ОШИБКА: %s" % exc, file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
