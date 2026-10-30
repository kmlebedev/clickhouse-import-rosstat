#!/usr/bin/env python3
"""T3 research tool: OKPD2 -> KIPC mapping, CPI weights and PPI parity.

Stages:
  mapping  build the versioned OKPD2 -> KIPC mapping and coverage report
  series   build monthly PPI aggregates for equal / KIPC / narrow schemes
  pilot    run the R X-13 script on the weighted series
  parity   compare the seasonally adjusted series with CBR text anchors
  report   print coverage and parity summaries

The script writes only under build/ppi-sa-research/ and docs/research/.
It is research tooling: it never writes to ClickHouse and never changes the schema.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import subprocess
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

import ppi_sa_probe as probe

ROOT = probe.ROOT
BUILD = probe.BUILD
RAW = probe.RAW
OUT = BUILD / "out"
RESEARCH = ROOT / "docs" / "research"

TAXONOMY_VERSION = "okpd2-kipc-2026-v2"
WEIGHT_YEAR = 2026
WEIGHTS_FILE = RAW / "vesa_kIPC.xlsx"
PRODUCTS_FILE = RAW / "tov.xlsx"
RULES_FILE = RESEARCH / "2026-10-20-ppi-kipc-rules.csv"
MAPPING_FILE = RESEARCH / "2026-10-20-ppi-kipc-mapping.csv"
COVERAGE_FILE = RESEARCH / "2026-10-20-ppi-kipc-coverage.json"
WEIGHTS_PANEL_FILE = RESEARCH / "2026-10-20-ppi-kipc-weights-panel.csv"
PARITY_FILE = RESEARCH / "2026-10-20-ppi-kipc-parity.csv"
PARITY_SUMMARY_FILE = RESEARCH / "2026-10-20-ppi-kipc-parity-summary.json"

NARROW_CLASSES = {"13", "14", "15", "17", "18", "21", "31", "32"}

SCOPE_GROUPS = {
    "food": ["01.1.1", "01.1.2", "01.1.3", "01.1.4", "01.1.5", "01.1.8", "01.1.9", "01.2"],
    "alcohol": ["02.1"],
    "tobacco": ["02.3"],
    "nonfood": [
        "04.3.1.1",
        "03.1",
        "03.2",
        "05.1",
        "05.2",
        "05.3",
        "05.4",
        "05.5",
        "05.6",
        "06.1",
        "07.1",
        "07.2.1",
        "08.1",
        "09.1",
        "09.2",
        "09.3",
        "09.7",
        "13.1.2",
        "13.2",
    ],
    "fuel": ["07.2.2"],
    "utilities": ["04.4", "04.5"],
}
SUBTRACT_GROUPS = {
    "03.1": ["03.1.4"],
    "03.2": ["03.2.2"],
    "05.6": ["05.6.2"],
}
ALL_GROUPS = (
    SCOPE_GROUPS["food"]
    + SCOPE_GROUPS["alcohol"]
    + SCOPE_GROUPS["tobacco"]
    + SCOPE_GROUPS["nonfood"]
    + SCOPE_GROUPS["fuel"]
)
CORE_GROUPS = SCOPE_GROUPS["food"] + SCOPE_GROUPS["nonfood"]


@dataclass(frozen=True)
class Rule:
    priority: int
    match_type: str
    pattern: str
    kipc_group: str
    status: str
    note: str


def _text(value) -> str:
    return str(value or "").strip()


def _rows(path: Path, sheet: str) -> list[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def _header_row(rows: list[tuple]) -> int:
    for index, row in enumerate(rows):
        if any("Код товара" in _text(cell) for cell in row):
            return index
    raise ValueError("KIPC header 'Код товара' not found")


def _sheet_year(sheet: str) -> int | None:
    match = re.search(r"(20\d{2})", sheet)
    return int(match.group(1)) if match else None


def parse_kipc_single(path: Path, sheet: str) -> dict[str, dict]:
    """Parse a single-year KIPC sheet into code -> {name, weight}."""
    rows = _rows(path, sheet)
    header = _header_row(rows)
    weight_col = None
    for col, cell in enumerate(rows[header]):
        text = _text(cell).lower()
        if text in {"спр", "вес", "доля"} or re.fullmatch(r"20\d{2}", text):
            weight_col = col
            break
    if weight_col is None:
        weight_col = 2
    result: dict[str, dict] = {}
    for row in rows[header + 1 :]:
        code = _text(row[0]) if row else ""
        name = _text(row[1]) if len(row) > 1 else ""
        if not code or not name:
            continue
        value = probe._float(row[weight_col]) if weight_col < len(row) else None
        if value is None:
            continue
        result[code] = {"name": name, "weight": value}
    return result


def parse_kipc_panel(path: Path) -> list[dict]:
    """Parse all KIPC sheets into year/period rows for the weights panel."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = list(wb.sheetnames)
    wb.close()
    result: list[dict] = []
    for sheet in sheets:
        if sheet in {"Содержание"}:
            continue
        rows = _rows(path, sheet)
        try:
            header = _header_row(rows)
        except ValueError:
            continue
        year_cols: list[tuple[int, int]] = []
        for col, cell in enumerate(rows[header]):
            text = _text(cell)
            if re.fullmatch(r"20\d{2}", text):
                year_cols.append((col, int(text)))
        if year_cols:
            periods = [(col, str(year)) for col, year in year_cols]
        else:
            year = _sheet_year(sheet)
            if year is None:
                continue
            period = f"{year}-04" if "01.04" in sheet else str(year)
            periods = [(2, period)]
        for row in rows[header + 1 :]:
            code = _text(row[0]) if row else ""
            name = _text(row[1]) if len(row) > 1 else ""
            if not code or not name:
                continue
            for col, period in periods:
                value = probe._float(row[col]) if col < len(row) else None
                if value is None:
                    continue
                result.append({"period": period, "code": code, "name": name, "weight": value})
    return result


def load_weights() -> dict[str, dict]:
    weights = parse_kipc_single(WEIGHTS_FILE, str(WEIGHT_YEAR))
    if not weights:
        raise ValueError(f"{WEIGHTS_FILE}: no {WEIGHT_YEAR} weights")
    return weights


def _is_child(code: str, parent: str) -> bool:
    """OKPD2 hierarchy: parent is a prefix followed by '.' or a digit."""
    if code == parent or not code.startswith(parent):
        return False
    return code[len(parent)] in "._" or code[len(parent)].isdigit()


def _is_kipc_child(code: str, parent: str) -> bool:
    """KIPC hierarchy uses dot-separated segments only (01.1.1 != 01.1.10)."""
    return code.startswith(parent + ".")


def _leaf_codes(codes: set[str]) -> set[str]:
    return {code for code in codes if not any(_is_child(other, code) for other in codes)}


def _group_weight(weights: dict[str, dict], code: str) -> float:
    if code in weights:
        return float(weights[code]["weight"])
    leaves = [
        item["weight"]
        for key, item in weights.items()
        if _is_kipc_child(key, code)
        and not any(_is_kipc_child(other, key) for other in weights)
    ]
    return float(sum(leaves))


def _scope_weight(weights: dict[str, dict], groups: list[str]) -> float:
    total = 0.0
    for group in groups:
        value = _group_weight(weights, group)
        for sub in SUBTRACT_GROUPS.get(group, []):
            value -= _group_weight(weights, sub)
        total += value
    return total


def load_rules(path: Path) -> list[Rule]:
    rules: list[Rule] = []
    with path.open(encoding="utf-8", newline="") as fh:
        for row in csv.DictReader(fh):
            pattern = _text(row["pattern"])
            if row["match_type"] == "name_regex":
                re.compile(pattern)
            rules.append(
                Rule(
                    priority=int(row["priority"]),
                    match_type=row["match_type"],
                    pattern=pattern,
                    kipc_group=_text(row["kipc_group"]),
                    status=row["status"],
                    note=_text(row["note"]),
                )
            )
    return rules


def normalize_name(name: str) -> str:
    text = name.lower().replace("ё", "е").replace("\xa0", " ")
    text = re.sub(r"«[^»]*»|\"[^\"]*\"", " ", text)
    text = re.sub(r"\b(?:нд|пд|у)\b", " ", text)
    text = re.sub(r"[^a-zа-я0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def classify(code: str, name: str, rules: list[Rule]) -> Rule | None:
    exact = [rule for rule in rules if rule.match_type == "code_exact" and rule.pattern == code]
    if exact:
        return max(exact, key=lambda rule: rule.priority)
    normalized = normalize_name(name)
    matches = [
        rule
        for rule in rules
        if rule.match_type == "name_regex" and re.search(rule.pattern, normalized)
    ]
    if matches:
        return max(matches, key=lambda rule: (rule.priority, len(rule.pattern)))
    prefixes = [rule for rule in rules if rule.match_type == "code_prefix" and code.startswith(rule.pattern)]
    if prefixes:
        return max(prefixes, key=lambda rule: (len(rule.pattern), rule.priority))
    return None


def load_products() -> list[tuple[str, str]]:
    result: dict[str, str] = {}
    for _year, items in probe.build_products().items():
        for item in items:
            result.setdefault(item["code"], item["name"])
    return sorted(result.items())


def build_mapping(products: list[tuple[str, str]], weights: dict[str, dict], rules: list[Rule]) -> list[dict]:
    codes = {code for code, _ in products}
    leaves = _leaf_codes(codes)
    rows: list[dict] = []
    for code, name in products:
        rule = classify(code, name, rules)
        is_leaf = code in leaves
        status = rule.status if rule else "unmatched"
        group = rule.kipc_group if rule else ""
        note = rule.note if rule else ""
        if status == "mapped" and group not in weights:
            status = "unmatched"
            note = f"KIPC group {group} not found in {WEIGHT_YEAR} weights"
            group = ""
        if not is_leaf:
            status = "aggregate"
        rows.append(
            {
                "taxonomy_version": TAXONOMY_VERSION,
                "okpd2": code,
                "name": name,
                "is_leaf": "1" if is_leaf else "0",
                "kipc_group": group,
                "kipc_name": weights[group]["name"] if group in weights else "",
                "weight_2026": f"{weights[group]['weight']:.6f}" if group in weights else "",
                "match_type": rule.match_type if rule else "",
                "pattern": rule.pattern if rule else "",
                "status": status,
                "note": note,
            }
        )
    return rows


def write_mapping(rows: list[dict]) -> None:
    RESEARCH.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "taxonomy_version",
        "okpd2",
        "name",
        "is_leaf",
        "kipc_group",
        "kipc_name",
        "weight_2026",
        "match_type",
        "pattern",
        "status",
        "note",
    ]
    with MAPPING_FILE.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_weights_panel(weights: dict[str, dict]) -> None:
    panel = parse_kipc_panel(WEIGHTS_FILE)
    wanted = set(ALL_GROUPS) | set(SCOPE_GROUPS["utilities"]) | {
        sub for values in SUBTRACT_GROUPS.values() for sub in values
    }
    wanted |= {
        code for code in weights if code == "04.3.1.1" or _is_kipc_child(code, "04.3.1.1")
    }
    rows = [row for row in panel if row["code"] in wanted]
    rows.sort(key=lambda row: (row["period"], row["code"]))
    with WEIGHTS_PANEL_FILE.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=["period", "code", "name", "weight"], lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(rows)


def _mapped_leaves(rows: list[dict], groups: set[str] | None = None) -> list[str]:
    result = []
    for row in rows:
        if row["is_leaf"] != "1" or row["status"] != "mapped":
            continue
        if groups is not None and not any(
            row["kipc_group"] == group or _is_kipc_child(row["kipc_group"], group) for group in groups
        ):
            continue
        result.append(row["okpd2"])
    return sorted(result)


def _narrow_leaves(rows: list[dict]) -> list[str]:
    return sorted(
        row["okpd2"]
        for row in rows
        if row["is_leaf"] == "1"
        and row["status"] == "mapped"
        and row["okpd2"][:2] in NARROW_CLASSES
    )


def _by_group(rows: list[dict]) -> dict[str, list[str]]:
    result: dict[str, list[str]] = defaultdict(list)
    for row in rows:
        if row["is_leaf"] == "1" and row["status"] == "mapped":
            result[row["kipc_group"]].append(row["okpd2"])
    return {group: sorted(codes) for group, codes in result.items()}


def _dedupe_members(members: dict[str, list[str]]) -> dict[str, list[str]]:
    keys = set(members)
    return {
        group: codes
        for group, codes in members.items()
        if not any(_is_kipc_child(group, other) for other in keys if other != group)
    }


def _members_for_groups(
    members: dict[str, list[str]], groups: list[str]
) -> dict[str, list[str]]:
    selected = {
        group: codes
        for group, codes in members.items()
        if any(group == scope or _is_kipc_child(group, scope) for scope in groups)
    }
    return _dedupe_members(selected)


def _mean_by_date(codes: list[str], points: dict[str, dict[date, float]]) -> list[tuple[date, float]]:
    buckets: dict[date, list[float]] = defaultdict(list)
    for code in codes:
        for day, value in points.get(code, {}).items():
            buckets[day].append(value)
    return sorted((day, sum(values) / len(values)) for day, values in buckets.items())


def _group_series(
    members: dict[str, list[str]],
    weights: dict[str, dict],
    points: dict[str, dict[date, float]],
) -> list[tuple[date, float]]:
    buckets: dict[date, list[tuple[float, float]]] = defaultdict(list)
    for group, codes in members.items():
        if not codes or group not in weights:
            continue
        weight = _group_weight(weights, group)
        for sub in SUBTRACT_GROUPS.get(group, []):
            weight -= _group_weight(weights, sub)
        for day, value in _mean_by_date(codes, points):
            buckets[day].append((weight, value))
    result = []
    for day, pairs in sorted(buckets.items()):
        total_weight = sum(weight for weight, _ in pairs)
        if total_weight:
            result.append((day, sum(weight * value for weight, value in pairs) / total_weight))
    return result


def build_series(rows: list[dict], weights: dict[str, dict]) -> dict[str, list[tuple[date, float]]]:
    points = {}
    for year, items in probe.build_products().items():
        for item in items:
            points.setdefault(item["code"], {})[date(year, item["month"], 1)] = item["value"]

    members = _by_group(rows)
    food_codes = _mapped_leaves(rows, set(SCOPE_GROUPS["food"]))
    nonfood_codes = _mapped_leaves(rows, set(SCOPE_GROUPS["nonfood"]))
    all_codes = _mapped_leaves(rows, set(ALL_GROUPS))
    core_codes = _mapped_leaves(rows, set(CORE_GROUPS))
    narrow_codes = _narrow_leaves(rows)
    narrow_all_codes = sorted(set(food_codes) | set(narrow_codes))

    food_members = _members_for_groups(members, SCOPE_GROUPS["food"])
    nonfood_members = _members_for_groups(members, SCOPE_GROUPS["nonfood"])
    all_members = _members_for_groups(members, ALL_GROUPS)
    core_members = _members_for_groups(members, CORE_GROUPS)

    return {
        "equal/food": _mean_by_date(food_codes, points),
        "equal/nonfood": _mean_by_date(nonfood_codes, points),
        "equal/all": _mean_by_date(all_codes, points),
        "equal/core": _mean_by_date(core_codes, points),
        "kipc/food": _group_series(food_members, weights, points),
        "kipc/nonfood": _group_series(nonfood_members, weights, points),
        "kipc/all": _group_series(all_members, weights, points),
        "kipc/core": _group_series(core_members, weights, points),
        "narrow/food": _mean_by_date(food_codes, points),
        "narrow/nonfood": _mean_by_date(narrow_codes, points),
        "narrow/all": _mean_by_date(narrow_all_codes, points),
        "narrow/core": _mean_by_date(narrow_all_codes, points),
    }


def write_series(series: dict[str, list[tuple[date, float]]]) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "series_weighted.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(["series", "date", "level"])
        for name, points in series.items():
            for day, level in probe.chain(points):
                writer.writerow([name, day.isoformat(), f"{level:.6f}"])
    return path


def _counts(rows: list[dict]) -> dict:
    leaves = [row for row in rows if row["is_leaf"] == "1"]
    by_status = Counter(row["status"] for row in leaves)
    by_class: dict[str, dict] = defaultdict(lambda: Counter())
    for row in leaves:
        by_class[row["okpd2"][:2]]["total"] += 1
        by_class[row["okpd2"][:2]][row["status"]] += 1
    return {
        "products_total": len(rows),
        "leaves_total": len(leaves),
        "by_status": dict(sorted(by_status.items())),
        "by_class": {key: dict(sorted(value.items())) for key, value in sorted(by_class.items())},
    }


def _net_group_weight(weights: dict[str, dict], group: str) -> float:
    value = _group_weight(weights, group)
    for sub in SUBTRACT_GROUPS.get(group, []):
        value -= _group_weight(weights, sub)
    return value


def _mapped_weight_in_group(weights: dict[str, dict], mapped_codes: set[str], group: str) -> float:
    inside = {code for code in mapped_codes if code == group or _is_kipc_child(code, group)}
    top = {
        code
        for code in inside
        if not any(_is_kipc_child(code, other) for other in inside if other != code)
    }
    return sum(_net_group_weight(weights, code) for code in top)


def build_coverage(rows: list[dict], weights: dict[str, dict]) -> dict:
    counts = _counts(rows)
    mapped_codes = {
        row["kipc_group"]
        for row in rows
        if row["is_leaf"] == "1"
        and row["status"] in {"mapped", "out_of_scope"}
        and row["kipc_group"]
    }

    def group_report(groups: list[str]) -> dict:
        in_scope = sum(_net_group_weight(weights, group) for group in groups)
        mapped = sum(_mapped_weight_in_group(weights, mapped_codes, group) for group in groups)
        missing = []
        for group in groups:
            group_net = _net_group_weight(weights, group)
            group_mapped = _mapped_weight_in_group(weights, mapped_codes, group)
            if group_mapped + 1e-9 < group_net:
                missing.append(
                    {
                        "code": group,
                        "name": weights.get(group, {}).get("name", ""),
                        "weight": round(group_net - group_mapped, 6),
                        "weight_mapped": round(group_mapped, 6),
                    }
                )
        return {
            "weight_in_scope": round(in_scope, 6),
            "weight_mapped": round(mapped, 6),
            "coverage": round(mapped / in_scope, 6) if in_scope else None,
            "missing_groups": sorted(missing, key=lambda item: -item["weight"]),
        }

    food = group_report(SCOPE_GROUPS["food"])
    nonfood = group_report(SCOPE_GROUPS["nonfood"])
    core = group_report(CORE_GROUPS)
    all_report = group_report(ALL_GROUPS)
    utilities = group_report(SCOPE_GROUPS["utilities"])
    unmatched_products = [
        {"okpd2": row["okpd2"], "name": row["name"], "class": row["okpd2"][:2]}
        for row in rows
        if row["is_leaf"] == "1" and row["status"] == "unmatched"
    ]
    excluded_products = [
        {"okpd2": row["okpd2"], "name": row["name"], "class": row["okpd2"][:2]}
        for row in rows
        if row["is_leaf"] == "1" and row["status"] == "excluded"
    ]
    return {
        "taxonomy_version": TAXONOMY_VERSION,
        "weight_year": WEIGHT_YEAR,
        **counts,
        "weight_total": round(
            food["weight_in_scope"]
            + nonfood["weight_in_scope"]
            + _net_group_weight(weights, "02.1")
            + _net_group_weight(weights, "02.3")
            + _net_group_weight(weights, "07.2.2")
            + utilities["weight_in_scope"],
            6,
        ),
        "food": food,
        "nonfood": nonfood,
        "alcohol": group_report(SCOPE_GROUPS["alcohol"]),
        "tobacco": group_report(SCOPE_GROUPS["tobacco"]),
        "fuel": group_report(SCOPE_GROUPS["fuel"]),
        "core": core,
        "all": all_report,
        "utilities": utilities,
        "unmatched_products": sorted(unmatched_products, key=lambda item: item["okpd2"]),
        "excluded_products": sorted(excluded_products, key=lambda item: item["okpd2"]),
    }


def write_coverage(coverage: dict) -> None:
    RESEARCH.mkdir(parents=True, exist_ok=True)
    COVERAGE_FILE.write_text(json.dumps(coverage, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def run_r(input_path: Path, output_path: Path) -> None:
    env = dict(os.environ)
    env["R_LIBS_USER"] = str(BUILD / "r-lib")
    subprocess.run(
        ["Rscript", str(ROOT / "scripts" / "ppi_sa_x13.R"), str(input_path), str(output_path)],
        check=True,
        env=env,
        cwd=ROOT,
    )


def _float(value) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if text in {"", "NA", "NaN", "None"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _stats(values: list[dict]) -> dict:
    if not values:
        return {"n": 0}
    diffs = [item["diff"] for item in values]
    mean_anchor = sum(item["anchor"] for item in values) / len(values)
    mean_ours = sum(item["ours"] for item in values) / len(values)
    cov = sum((item["anchor"] - mean_anchor) * (item["ours"] - mean_ours) for item in values)
    var_a = sum((item["anchor"] - mean_anchor) ** 2 for item in values)
    var_o = sum((item["ours"] - mean_ours) ** 2 for item in values)
    corr = cov / (var_a * var_o) ** 0.5 if var_a > 0 and var_o > 0 else None
    return {
        "n": len(values),
        "mae": sum(abs(value) for value in diffs) / len(diffs),
        "max_abs": max(abs(value) for value in diffs),
        "bias": sum(diffs) / len(diffs),
        "corr": corr,
    }


def compute_parity(sa_path: Path, anchors_path: Path) -> tuple[list[dict], dict]:
    with sa_path.open(encoding="utf-8") as fh:
        sa_rows = list(csv.DictReader(fh))
    with anchors_path.open(encoding="utf-8") as fh:
        anchors = list(csv.DictReader(fh))
    sa: dict[str, dict[str, dict[str, dict]]] = defaultdict(lambda: defaultdict(dict))
    for row in sa_rows:
        sa[row["series"]][row["variant"]][row["date"]] = row
    variants = sorted({row["variant"] for row in sa_rows})
    results: list[dict] = []
    schemes = ("equal", "kipc", "narrow")
    for anchor in anchors:
        target = anchor["target"]
        expected = float(anchor["value"])
        anchor_series = anchor["series"]
        if anchor_series == "all":
            candidates = [(scheme, f"{scheme}/core", "core", True) for scheme in schemes] + [
                (scheme, f"{scheme}/all", "all", False) for scheme in schemes
            ]
        else:
            candidates = [
                (scheme, f"{scheme}/{anchor_series}", anchor_series, True) for scheme in schemes
            ]
        for scheme, series_name, series_kind, primary in candidates:
            for variant in variants:
                current = sa.get(series_name, {}).get(variant, {}).get(target)
                if current is None:
                    continue
                previous = sa.get(series_name, {}).get(variant, {}).get(
                    probe._shift_month(target, -3)
                )
                mean3 = _float(current.get("saar3"))
                level3 = None
                if previous is not None:
                    current_sa = _float(current.get("sa"))
                    previous_sa = _float(previous.get("sa"))
                    if current_sa is not None and previous_sa:
                        level3 = ((current_sa / previous_sa) ** 4 - 1) * 100
                for formula, ours in (("mean3_saar", mean3), ("level3_annualized", level3)):
                    if ours is None:
                        continue
                    results.append(
                        {
                            "scheme": scheme,
                            "series": series_name,
                            "series_kind": series_kind,
                            "variant": variant,
                            "formula": formula,
                            "issue": anchor["issue"],
                            "target": target,
                            "anchor_series": anchor_series,
                            "anchor": expected,
                            "ours": ours,
                            "diff": ours - expected,
                            "primary": primary,
                        }
                    )
    summary: dict[str, dict] = {}
    for scheme in schemes:
        for formula in ("mean3_saar", "level3_annualized"):
            subset = [
                item
                for item in results
                if item["scheme"] == scheme and item["formula"] == formula
            ]
            summary[f"{scheme}/{formula}"] = _stats(subset)
            summary[f"{scheme}/{formula}/primary"] = _stats(
                [item for item in subset if item["primary"]]
            )
            for kind in ("food", "nonfood", "all", "core"):
                summary[f"{scheme}/{formula}/{kind}"] = _stats(
                    [item for item in subset if item["series_kind"] == kind]
                )
    return results, summary


def write_parity(results: list[dict], summary: dict) -> None:
    RESEARCH.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "scheme",
        "series",
        "series_kind",
        "variant",
        "formula",
        "issue",
        "target",
        "anchor_series",
        "anchor",
        "ours",
        "diff",
        "primary",
    ]
    with PARITY_FILE.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(results)
    PARITY_SUMMARY_FILE.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def stage_mapping() -> None:
    weights = load_weights()
    rules = load_rules(RULES_FILE)
    products = load_products()
    rows = build_mapping(products, weights, rules)
    write_mapping(rows)
    write_weights_panel(weights)
    coverage = build_coverage(rows, weights)
    write_coverage(coverage)
    print(
        json.dumps(
            {
                "mapping": str(MAPPING_FILE.relative_to(ROOT)),
                "coverage": str(COVERAGE_FILE.relative_to(ROOT)),
                "products_total": coverage["products_total"],
                "leaves_total": coverage["leaves_total"],
                "by_status": coverage["by_status"],
                "core_coverage": coverage["core"],
                "unmatched": len(coverage["unmatched_products"]),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


def stage_series() -> None:
    weights = load_weights()
    rules = load_rules(RULES_FILE)
    rows = build_mapping(load_products(), weights, rules)
    series = build_series(rows, weights)
    path = write_series(series)
    print(
        json.dumps(
            {
                "series": str(path.relative_to(ROOT)),
                "points": {name: len(points) for name, points in series.items()},
            },
            ensure_ascii=False,
            indent=2,
        )
    )


def stage_pilot() -> None:
    input_path = OUT / "series_weighted.csv"
    output_path = OUT / "sa_weighted.csv"
    if not input_path.exists():
        raise SystemExit(f"run `series` first: {input_path} not found")
    run_r(input_path, output_path)
    print(json.dumps({"input": str(input_path), "output": str(output_path)}, ensure_ascii=False))


def stage_parity() -> None:
    anchors = OUT / "anchors.csv"
    sa_path = OUT / "sa_weighted.csv"
    if not anchors.exists():
        probe.parse_anchors()
    if not sa_path.exists():
        raise SystemExit(f"run `pilot` first: {sa_path} not found")
    results, summary = compute_parity(sa_path, anchors)
    write_parity(results, summary)
    primary = {
        key: value
        for key, value in summary.items()
        if key.endswith("/mean3_saar") or "/mean3_saar/" in key
    }
    print(json.dumps(primary, ensure_ascii=False, indent=2))


def stage_report() -> None:
    coverage = json.loads(COVERAGE_FILE.read_text(encoding="utf-8"))
    summary = json.loads(PARITY_SUMMARY_FILE.read_text(encoding="utf-8"))
    print("coverage:", json.dumps(coverage["core"], ensure_ascii=False))
    print("parity (mean3_saar):")
    for key in sorted(summary):
        if key.endswith("/mean3_saar") or "/mean3_saar/" in key:
            print(key, json.dumps(summary[key], ensure_ascii=False))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", choices=["mapping", "series", "pilot", "parity", "report"])
    args = parser.parse_args()
    if args.stage == "mapping":
        stage_mapping()
    elif args.stage == "series":
        stage_series()
    elif args.stage == "pilot":
        stage_pilot()
    elif args.stage == "parity":
        stage_parity()
    elif args.stage == "report":
        stage_report()
    return 0


if __name__ == "__main__":
    sys.exit(main())
