#!/usr/bin/env python3
"""Офлайн-тесты scripts/harvest.py: сеть не вызывается.

Проверяются решения, от которых зависит корректность раскладки в винтажный архив:
классификация заголовков (недельные/месячные/чужие релизы), декодирование cp1251/utf-8,
санитайзинг имён и ключ версии (повторяют archive.Key), композит отпечатка (повторяет
util.HeadFingerprint), upsert реестра (повторяет archive.upsert), разбор канонического
списка недель из Go и инспекция книги XLSX режима xlsx.

Запуск: python3 scripts/harvest_test.py
"""

import importlib.util
import json
import os
import tempfile
import types
import unittest
from pathlib import Path

spec = importlib.util.spec_from_file_location("harvest", Path(__file__).parent / "harvest.py")
harvest = importlib.util.module_from_spec(spec)
spec.loader.exec_module(harvest)


class ClassifyTest(unittest.TestCase):
    def test_weekly_plain(self):
        self.assertEqual(harvest.classify("Об оценке индекса потребительских цен с 3 по 9 августа 2021 года"),
                         "week:2021-08-09")

    def test_weekly_so_and_month_before_po(self):
        self.assertEqual(harvest.classify("Об оценке индекса потребительских цен со 2 по 9 марта 2021 года"),
                         "week:2021-03-09")
        self.assertEqual(harvest.classify("Об оценке индекса потребительских цен с 26 октября по 1 ноября 2017 года"),
                         "week:2017-11-01")

    def test_weekly_typo_cens(self):
        # Реальная опечатка Росстата 2021-05-31: «потребительских ценс 25 по 31 мая 2021 года»
        self.assertEqual(harvest.classify("Об оценке индекса потребительских ценс 25 по 31 мая 2021 года"),
                         "week:2021-05-31")

    def test_monthly(self):
        self.assertEqual(harvest.classify("Об индексе потребительских цен в апреле 2017 года"), "monthly")
        # «оценке» без диапазона — это месячный релиз
        self.assertEqual(harvest.classify("Об оценке индекса потребительских цен в декабре 2020 года"),
                         "monthly")

    def test_other_fuel_release_not_weekly(self):
        self.assertEqual(
            harvest.classify("Об объеме производства нефтепродуктов и потребительских ценах на них с 20 по 26 марта 2017 года"),
            "other")

    def test_week_without_parse(self):
        self.assertEqual(harvest.classify("Об оценке индекса потребительских цен с 5 по ? 2021 года"),
                         "week:??")


class DecodeTest(unittest.TestCase):
    def test_cp1251(self):
        raw = "Об оценке индекса потребительских цен".encode("cp1251")
        self.assertIn("потребительских", harvest.decode_title(raw))

    def test_utf8_placeholder_page(self):
        # Заглушка живого буфера приходит в utf-8; чтение «cp1251 первым» дало бы мохибейк
        raw = "Федеральная служба государственной статистики".encode("utf-8")
        self.assertIn("Федеральная", harvest.decode_title(raw))

    def test_title_of(self):
        body = "<html><head><title>Об оценке индекса потребительских цен с 1 по 11 января 2021 года</title></head>".encode("cp1251")
        self.assertEqual(harvest.classify(harvest.title_of(body)), "week:2021-01-11")

    def test_headline_prefers_title(self):
        body = ("<html><head><title>Об оценке индекса потребительских цен с 1 по 11 января 2021 года"
                "</title></head><body>x</body>").encode("cp1251")
        self.assertEqual(harvest.classify(harvest.headline_of(body)), "week:2021-01-11")

    def test_headline_from_body_without_title(self):
        # Word-документы без <title> (57.htm/77.htm/ind-zen24.htm 2017–2019): заголовок —
        # первая строка видимого текста, <body> — после блока стилей
        body = (b"<html><head><style>" + b".x{color:red}" * 4000 + b"</style></head><body>"
                + "<p>Об оценке индекса потребительских цен с 3 по 9 апреля 2017 года</p>".encode("cp1251"))
        self.assertEqual(harvest.classify(harvest.headline_of(body)), "week:2017-04-09")

    def test_headline_of_none(self):
        # Пустой ответ сети не должен ронять воркер (регресс: AttributeError на None)
        self.assertEqual(harvest.headline_of(None), "")
        self.assertEqual(harvest.headline_of(b""), "")

    def test_clean_title_strips_control_chars(self):
        # NUL в заголовке делал titles.csv нечитаемым для csv-модуля (регресс «line contains NUL»)
        self.assertEqual(harvest.clean_title("Об оценке\x00 цен; с 3 по 9"),
                         "Об оценке  цен  с 3 по 9")
        self.assertEqual(harvest.clean_title("Об оценке\r\nцен"), "Об оценке цен")

    def test_headline_foreign_release_is_other(self):
        body = (b"<html><head><style>" + b".x{}" * 4000 + b"</style></head><body>"
                + "<p>О динамике цен на бензин автомобильный и ресурсах нефтепродуктов в феврале 2017 года</p>".encode("cp1251"))
        self.assertEqual(harvest.classify(harvest.headline_of(body)), "other")


class KeyTest(unittest.TestCase):
    def test_basename_like_go_path_base(self):
        url = "https://web.archive.org/web/20201224012041id_/https://rosstat.gov.ru/bgd/free/B04_03/IssWWW.exe/Stg/d02/237.htm"
        self.assertEqual(harvest.basename(url), "237.htm")
        self.assertEqual(harvest.basename("http://x/y/56.htm?a=1"), "56.htm?a=1")

    def test_sanitize(self):
        self.assertEqual(harvest.sanitize_name('a/b\\c:d*e?f"g<h>i|j'), "a_b_c_d_e_f_g_h_i_j")
        self.assertEqual(harvest.sanitize_name("a\nb"), "a_b")
        self.assertEqual(harvest.sanitize_name(""), "file")

    def test_sanitize_truncates_on_rune_boundary(self):
        name = "я" * 150  # 300 байт cp1251-независимо: utf-8 — 2 байта на букву
        clean = harvest.sanitize_name(name)
        self.assertLessEqual(len(clean.encode("utf-8")), 200)
        self.assertTrue(clean.encode("utf-8").decode("utf-8"))  # непустая и валидная

    def test_key_prefix(self):
        key = harvest.archive_key("src", "http://x/56.htm", "0123456789abcdef0123", "ff" * 32)
        self.assertEqual(key, "src/0123456789abcdef_56.htm")

    def test_key_sha_fallback(self):
        key = harvest.archive_key("src", "http://x/56.htm", "", "ab" * 32)
        self.assertEqual(key, "src/sha-abababababababab_56.htm")

    def test_head_fingerprint_composite(self):
        headers = [("ETag", '"v1"'), ("Last-Modified", "Mon, 01 Jan 2020 00:00:00 GMT"),
                   ("Content-Length", "123"), ("Server", "nginx")]
        got = harvest.head_fingerprint_headers(headers, "http://x/56.htm")
        import hashlib
        want = hashlib.md5(("http://x/56.htm|56.htm|\"v1\"|Mon, 01 Jan 2020 00:00:00 GMT|123").encode()).hexdigest()
        self.assertEqual(got, want)

    def test_head_fingerprint_none(self):
        self.assertEqual(harvest.head_fingerprint_headers(None, "http://x/56.htm"), "")


class RegistryTest(unittest.TestCase):
    def entry(self, key, sha):
        return {"source": "s", "name": "56.htm", "fingerprint": "f", "sha256": sha,
                "size": 1, "url": "u", "key": key, "first_seen": "2026-10-16T00:00:00+05:00"}

    def test_append_new(self):
        rows = harvest.registry_upsert([], self.entry("s/1_56.htm", "a"))
        self.assertEqual(len(rows), 1)

    def test_same_version_noop(self):
        rows = [self.entry("s/1_56.htm", "a")]
        self.assertIsNone(harvest.registry_upsert(rows, self.entry("s/1_56.htm", "a")))

    def test_changed_sha_moves_to_end(self):
        first = self.entry("s/1_56.htm", "a")
        second = self.entry("s/2_57.htm", "b")
        rows = [first, second]
        out = harvest.registry_upsert(rows, self.entry("s/1_56.htm", "c"))
        self.assertEqual([r["key"] for r in out], ["s/2_57.htm", "s/1_56.htm"])
        self.assertEqual(out[-1]["sha256"], "c")


class HarvesterTest(unittest.TestCase):
    def make(self, tmp, dry_run=False):
        args = types.SimpleNamespace(state=os.path.join(tmp, "state"),
                                     archive=os.path.join(tmp, "archive"),
                                     threads=1, pace=0.0, dry_run=dry_run,
                                     sources="wayback", go_sources="x.go")
        return harvest.Harvester(args)

    def test_pending_rows_retries_err(self):
        h = self.make(tempfile.mkdtemp())
        h.titles[("t1", "u1")] = {"ts": "t1", "url": "u1", "digest": "d", "kind": "other", "title": ""}
        h.titles[("t2", "u2")] = {"ts": "t2", "url": "u2", "digest": "d", "kind": "err", "title": ""}
        h.titles[("t3", "u3")] = {"ts": "t3", "url": "u3", "digest": "d", "kind": "week:2020-01-06", "title": ""}
        h.titles[("t4", "u4")] = {"ts": "t4", "url": "u4", "digest": "d", "kind": "week:??", "title": ""}
        pending = h.pending_rows([("t1", "u1"), ("t2", "u2"), ("t3", "u3"), ("t4", "u4"), ("t5", "u5")])
        self.assertEqual(sorted(p[0] for p in pending), ["t2", "t4", "t5"])

    def test_load_titles_tolerates_nul(self):
        h = self.make(tempfile.mkdtemp())
        with open(os.path.join(h.state, "titles.csv"), "w", encoding="utf-8") as fh:
            fh.write("ts,url,digest,kind,title\n")
            fh.write('20200101000000,http://x/1.htm,d,other,"Оценке\x00цены"\n')
            fh.write("20200102000000,http://x/2.htm,d,week:2020-01-06,\n")
        h.load_titles()
        self.assertEqual(len(h.titles), 2)
        self.assertEqual(h.titles[("20200102000000", "http://x/2.htm")]["kind"], "week:2020-01-06")

    def test_gaps_compare_with_known(self):
        h = self.make(tempfile.mkdtemp())
        h.titles[("t1", "u1")] = {"ts": "t1", "url": "u1", "digest": "", "kind": "week:2020-08-10", "title": ""}
        h.titles[("t2", "u2")] = {"ts": "t2", "url": "u2", "digest": "", "kind": "week:2021-05-31", "title": ""}
        h.titles[("t3", "u3")] = {"ts": "t3", "url": "u3", "digest": "", "kind": "week:2022-01-10", "title": ""}
        holes, found = h.phase_gaps({"2021-05-31"})
        self.assertEqual(holes, ["2020-08-10"])
        self.assertEqual(len(found["2020-08-10"]), 1)

    def test_archive_upsert_writes_object_and_registry(self):
        h = self.make(tempfile.mkdtemp())
        h.head_headers = lambda url: None  # сеть в тесте не вызывается: отпечатка нет
        body = "Об оценке индекса потребительских цен с 3 по 9 августа 2021 года".encode("cp1251")
        entry = h.archive_upsert("https://web.archive.org/web/20210811id_/http://x/d02/144.htm", body)
        self.assertIsNotNone(entry)
        self.assertTrue(entry["key"].startswith("rosstat_ipc_weeks_bulletins/sha-"))
        obj = os.path.join(h.archive, entry["key"])
        self.assertTrue(os.path.exists(obj))
        with open(obj, "rb") as fh:
            self.assertEqual(fh.read(), body)
        with open(h.registry_path, encoding="utf-8") as fh:
            lines = [l for l in fh.read().splitlines() if l]
        self.assertEqual(len(lines), 1)
        self.assertIn('{"source":"rosstat_ipc_weeks_bulletins"', lines[0])  # компактный JSONL
        self.assertEqual(json.loads(lines[0])["sha256"], entry["sha256"])
        # повтор той же версии — без изменений
        self.assertIsNone(h.archive_upsert("https://web.archive.org/web/20210811id_/http://x/d02/144.htm", body))
        # другой документ — вторая строка в конце
        other = "<html><title>Об индексе потребительских цен в августе 2021 года</title></html>".encode("cp1251")
        h.archive_upsert("https://web.archive.org/web/20210901id_/http://x/d02/161.htm", other)
        with open(h.registry_path, encoding="utf-8") as fh:
            lines = [l for l in fh.read().splitlines() if l]
        self.assertEqual(len(lines), 2)
        self.assertTrue(json.loads(lines[1])["key"].endswith("_161.htm"))

    def test_known_weeks_from_go(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "x.go")
            with open(path, "w", encoding="utf-8") as fh:
                fh.write('{URL: "https://a/56.htm", Date: time.Date(2017, time.March, 20, 0, 0, 0, 0, time.UTC)},\n'
                         '{URL: "https://a/245.htm", Date: time.Date(2020, time.December, 28, 0, 0, 0, 0, time.UTC)},\n')
            weeks, urls = harvest.known_weeks_from_go(path)
        self.assertEqual(weeks, {"2017-03-20", "2020-12-28"})
        self.assertEqual(len(urls), 2)


class XlsxTest(unittest.TestCase):
    def build(self, sheets_headers):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        first = True
        for name, headers in sheets_headers:
            if first:
                ws = wb.active
                ws.title = name
                first = False
            else:
                ws = wb.create_sheet(title=name)
            for col, value in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=value)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_earliest_2022(self):
        body = self.build([("2022", ["Наименование", "на 10 января **", "на 17 января **"])])
        sheets, earliest, _ = harvest.inspect_xlsx(body)
        self.assertEqual(sheets, ["2022"])
        self.assertEqual(earliest, {"2022": "2022-01-10"})

    def test_violation_2021(self):
        body = self.build([("Содержание", ["Лист"]), ("2021", ["Наименование", "на 11 января"])])
        _, earliest, _ = harvest.inspect_xlsx(body)
        self.assertEqual(earliest, {"2021": "2021-01-11"})

    def test_year_in_cell_wins(self):
        body = self.build([("Листы", ["на 28 декабря 2021 года"])])
        _, earliest, _ = harvest.inspect_xlsx(body)
        self.assertEqual(earliest, {"Листы": "2021-12-28"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
