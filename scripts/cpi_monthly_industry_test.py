#!/usr/bin/env python3
"""Офлайн-тесты `scripts/cpi_monthly_industry.py` (issue #235, Stage 2).

Запуск из корня репозитория:

    make cpi-monthly-industry-test

Тесты не требуют ClickHouse и сети: книга xlsx собирается синтетически, строки
выборки — на синтетических значениях. Проверяются геометрия листов, as-of
отраслевых точек, NaN-политика, mapping, имена признаков и smoke моделей.
"""

import math
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

import cpi_monthly as cm
import cpi_monthly_industry as ind


SHEETS = ("Промышленность", "Добыча")
RAW_FROM = 2
SA_FROM = 15


def _write_workbook(path: Path, sheets=SHEETS, missing_costs: bool = False) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        worksheet = workbook.create_sheet(sheet)
        worksheet.cell(row=1, column=1, value="Данные мониторинга предприятий")
        worksheet.cell(row=2, column=1, value="(исходные и сезонно скорректированные данные)")
        worksheet.cell(row=3, column=RAW_FROM + 1, value="Исходные данные")
        worksheet.cell(row=3, column=SA_FROM + 1, value="Сезонно скорректированные данные")
        headers = [None] * 26
        headers[0] = "Отчетный период"
        headers[1] = "Месяц заполнения и сбора анкет (первая декада)"
        headers[RAW_FROM] = "Индикатор бизнес-климата Банка России"
        headers[RAW_FROM + 1] = "Индикатор бизнес-климата Банка России (факт)"
        headers[RAW_FROM + 2] = "Индикатор бизнес-климата Банка России (ожид)"
        headers[9] = "Ценовые ожидания предприятий на следующие 3 месяца"
        headers[11] = "Как изменились издержки производства, обращения?"
        headers[SA_FROM] = "Индикатор бизнес-климата Банка России"
        headers[SA_FROM + 1] = "Индикатор бизнес-климата Банка России (факт)"
        headers[SA_FROM + 2] = "Индикатор бизнес-климата Банка России (ожид)"
        headers[22] = "Ценовые ожидания предприятий на следующие 3 месяца"
        if not missing_costs:
            headers[24] = "Как изменились издержки производства, обращения?"
        for index, header in enumerate(headers, start=1):
            if header is not None:
                worksheet.cell(row=4, column=index, value=header)
        for offset, day in enumerate(
            (date(2024, 1, 1), date(2024, 2, 1), date(2024, 3, 1), date(2024, 4, 1)),
            start=5,
        ):
            value = float(offset - 4)
            worksheet.cell(row=offset, column=1, value=datetime(day.year, day.month, day.day))
            worksheet.cell(row=offset, column=2, value=datetime(day.year, day.month, day.day))
            worksheet.cell(row=offset, column=RAW_FROM + 1, value=value)
            worksheet.cell(row=offset, column=10, value=value * 2)
            worksheet.cell(row=offset, column=12, value=value * 3)
            worksheet.cell(row=offset, column=SA_FROM + 1, value=value * 10)
            worksheet.cell(row=offset, column=23, value=value * 20)
            if not missing_costs:
                worksheet.cell(row=offset, column=25, value=value * 30)
    workbook.save(path)


def _rows() -> list:
    names = ind.feature_names(component=None, include_mp=True)
    rows = []
    for index in range(24):
        month = date(2020 + index // 12, index % 12 + 1, 1)
        food = 0.1 + 0.01 * index
        nonfood = 0.2 + 0.005 * index
        services = 0.15 + 0.007 * index
        features = {
            name: (float("nan") if name.endswith("_chg") and index % 5 == 0 else 0.1 + 0.001 * index)
            for name in names
            if name != "month"
        }
        features["month"] = month.month
        rows.append(
            cm.Row(
                month=month,
                y={
                    cm.TARGET_ALL: 0.3 * food + 0.4 * nonfood + 0.3 * services,
                    "cpi_food": food,
                    "cpi_nonfood": nonfood,
                    "cpi_services": services,
                },
                features=features,
            )
        )
    return rows


class TestPanelLoader(unittest.TestCase):
    def test_reads_sa_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mp_survey_data.xlsx"
            _write_workbook(path)
            panel = ind.load_industry_panel(path, sheets=SHEETS)
        self.assertEqual(set(panel), set(SHEETS))
        self.assertEqual(
            panel["Промышленность"]["ibk"][0], (date(2024, 1, 1), 10.0)
        )
        self.assertEqual(
            panel["Промышленность"]["price_exp"][-1], (date(2024, 4, 1), 80.0)
        )
        self.assertEqual(
            panel["Промышленность"]["costs"][-1], (date(2024, 4, 1), 120.0)
        )

    def test_missing_sheet_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mp_survey_data.xlsx"
            _write_workbook(path)
            with self.assertRaises(ValueError):
                ind.load_industry_panel(path, sheets=("Промышленность", "Нет листа"))

    def test_missing_column_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mp_survey_data.xlsx"
            _write_workbook(path, sheets=("Промышленность",), missing_costs=True)
            with self.assertRaises(ValueError):
                ind.load_industry_panel(path, sheets=("Промышленность",))


class TestIndustryFeatures(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "mp_survey_data.xlsx"
        _write_workbook(self.path)
        self.panel = ind.load_industry_panel(self.path, sheets=SHEETS)

    def test_level_and_change(self):
        features = ind.industry_features(
            self.panel, date(2024, 4, 1), cm.cutoff_for(date(2024, 4, 1))
        )
        self.assertEqual(features["ind_industry_ibk_sa"], 30.0)
        self.assertEqual(features["ind_industry_ibk_sa_chg"], 10.0)
        self.assertEqual(features["ind_mining_price_exp_sa"], 60.0)
        self.assertEqual(features["ind_mining_price_exp_sa_chg"], 20.0)

    def test_point_m_is_not_used(self):
        features = ind.industry_features(
            self.panel, date(2024, 3, 1), cm.cutoff_for(date(2024, 3, 1))
        )
        self.assertEqual(features["ind_industry_ibk_sa"], 20.0)
        self.assertNotEqual(features["ind_industry_ibk_sa"], 30.0)

    def test_point_unavailable_at_cutoff_is_ignored(self):
        features = ind.industry_features(
            self.panel, date(2024, 3, 1), date(2024, 3, 15)
        )
        self.assertTrue(math.isnan(features["ind_industry_ibk_sa"]))
        self.assertTrue(math.isnan(features["ind_industry_ibk_sa_chg"]))

    def test_missing_sheet_is_nan_not_zero(self):
        features = ind.industry_features(
            self.panel, date(2024, 4, 1), cm.cutoff_for(date(2024, 4, 1))
        )
        self.assertTrue(math.isnan(features["ind_eegp_ibk_sa"]))
        self.assertTrue(math.isnan(features["ind_services_costs_sa"]))
        self.assertNotEqual(features["ind_eegp_ibk_sa"], 0.0)


class TestMappingAndNames(unittest.TestCase):
    def test_feature_names_are_unique_and_complete(self):
        names = ind.industry_feature_names()
        self.assertEqual(len(names), 96)
        self.assertEqual(len(set(names)), 96)
        self.assertTrue(all(name.startswith("ind_") for name in names))
        self.assertEqual(names[-1], "ind_services_costs_sa_chg")

    def test_component_mapping_covers_all_sheets(self):
        covered = set()
        for component in cm.BOTTOMUP_COMPONENTS:
            covered.update(ind.component_sheets(component))
        covered.update(ind.GENERAL_SHEETS)
        self.assertEqual(covered, set(ind.INDUSTRY_SHEETS))
        for component in cm.BOTTOMUP_COMPONENTS:
            self.assertTrue(set(ind.GENERAL_SHEETS).issubset(ind.component_sheets(component)))

    def test_feature_sets(self):
        food = ind.feature_names("cpi_food")
        self.assertEqual(food[-1], "month")
        self.assertEqual(
            len(food), len(ind.base_feature_names(True, False)) + 5 * 6
        )
        direct = ind.feature_names(None)
        self.assertEqual(
            len(direct), len(ind.base_feature_names(True, False)) + 96
        )
        all_sheets = ind.feature_names("cpi_food", all_sheets=True)
        self.assertEqual(len(all_sheets), len(direct))

    def test_no_weekly_and_no_mp_names(self):
        names = ind.feature_names(None, include_mp=False, include_weekly=True)
        self.assertIn("wk_share_rising", names)
        self.assertNotIn("price_expectations_sa", names)
        self.assertEqual(names[-1], "month")


class TestModels(unittest.TestCase):
    def test_catboost_smoke(self):
        rows = _rows()
        value = ind.fit_catboost_industry(
            rows[:-1],
            rows[-1],
            cm.TARGET_ALL,
            component="cpi_food",
            include_mp=True,
            iterations=5,
            thread_count=1,
        )
        self.assertTrue(math.isfinite(value))

    def test_enet_smoke_with_nan(self):
        rows = _rows()
        value = ind.fit_enet_industry(
            rows[:-1], rows[-1], component="cpi_food", include_mp=True
        )
        self.assertTrue(math.isfinite(value))

    def test_bottomup_wiring_and_weights(self):
        rows = _rows()
        with mock.patch.object(
            ind, "fit_catboost_industry", return_value=0.1
        ):
            value = ind.fit_bottomup_industry(rows[:-1], rows[-1], include_mp=True)
        self.assertAlmostEqual(value, 0.1, places=12)
        weights = ind.fit_bottomup_industry.last_weights
        self.assertEqual(len(weights), 3)
        self.assertAlmostEqual(sum(weights), 1.0, places=12)


class TestEvaluationHelpers(unittest.TestCase):
    def test_holm(self):
        self.assertEqual(ind.holm_adjust([0.01, 0.04]), [0.02, 0.04])
        self.assertEqual(ind.holm_adjust([0.5, 0.6]), [1.0, 1.0])

    def test_stage1_gate(self):
        summary = {
            "v2_bottomup": {"n": 10, "rmse": 0.362300},
            "v3_bottomup": {"n": 10, "rmse": 0.311487},
        }
        gate = ind.stage1_gate("weekly", summary, False)
        self.assertEqual(gate["v2_bottomup"]["status"], "ok")
        self.assertEqual(gate["v3_bottomup"]["status"], "ok")
        with self.assertRaises(SystemExit):
            ind.stage1_gate(
                "weekly",
                {
                    "v2_bottomup": {"n": 10, "rmse": 1.0},
                    "v3_bottomup": {"n": 10, "rmse": 0.311487},
                },
                False,
            )

    def test_decide(self):
        def pair(delta, p_value, ci, primary=True):
            return {
                "n": 10,
                "primary": primary,
                "delta": delta,
                "dm_p": p_value,
                "dm_p_holm": p_value,
                "delta_ci": ci,
                "rmse_left": 0.0,
                "rmse_right": 0.0,
                "mde_ratio": 0.1,
            }

        promote = {
            "v3_ind_bottomup vs v3_bottomup": pair(-0.05, 0.01, [-0.1, -0.01]),
            "ind_only_bottomup vs v2_bottomup": pair(-0.02, 0.1, [-0.06, 0.01]),
            "v3_ind_all_bottomup vs v3_ind_bottomup": pair(
                -0.01, 0.5, [-0.05, 0.03], False
            ),
        }
        self.assertEqual(ind.decide(promote)["verdict"], "promote")

        reject = {
            "v3_ind_bottomup vs v3_bottomup": pair(0.05, 0.01, [0.01, 0.1]),
            "ind_only_bottomup vs v2_bottomup": pair(0.02, 0.01, [0.01, 0.06]),
        }
        self.assertEqual(ind.decide(reject)["verdict"], "reject")

        iterate = {
            "v3_ind_bottomup vs v3_bottomup": pair(-0.01, 0.4, [-0.05, 0.03]),
            "ind_only_bottomup vs v2_bottomup": pair(-0.01, 0.6, [-0.05, 0.03]),
        }
        self.assertEqual(ind.decide(iterate)["verdict"], "iterate")


if __name__ == "__main__":
    unittest.main()
