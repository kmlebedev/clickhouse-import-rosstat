#!/usr/bin/env python3
"""Разбор PR #435: per-code МП, гармонизация SA и разрыв на h1 — research 2026-10-31.

Пункт `key-rate/phase-4-cb-comp-88-groups-live`, зонтик #419; отчёт —
`docs/research/2026-10-31-per-code-mp-reconstruction-review.md`. Скрипт
воспроизводит числа §2 и §6 отчёта (§7 — `cb_comp_h1_bs_combo.py`, Task 9 #436).

Шаги (`--step`):

* `distinct` — различимость 87 × 3 per-code рядов архива (§2.1);
* `forward` — лист против среднего его кодов: train/test, уровни, изменения,
  смещение (§2.2);
* `inverse` — per-code ridge на всех публичных SA-рядах трёх книг ЦБ против
  identity/offset/calib (§2.3);
* `noise` — AR-структура остатка «код − лист» (§2.4);
* `priors` — offset на скользящем окне и AR-затухание (§2.3);
* `sa-inputs` — NSA/SA листов × 3 индикатора для X-13 (§6.1); X-13 —
  `cb_comp_per_code_review_x13.R`;
* `sa-compare` — SA ЦБ / NSA / X-13 против среднего архивных кодов (§6.2);
* `panels` — панели `per_sheet` для `make cb-comp-parity` (§6.3);
* `dm` — парный DM-HLN прогонов `cb-comp-parity` против SA ЦБ (§6.3).

Входы — локальные файлы: архив per-code (`panel.csv`, `public_panel.csv`,
`sheet_map.csv` из `build/cb-comp-parity`) и книги ЦБ `mp_survey_data.xlsx`,
`survey_data_group.xlsx`, `mp_survey_data_gu.xlsx` (`make cb-comp-per-code-review-books`).
В ClickHouse не пишет и сеть не использует. Запуск —
`make cb-comp-per-code-review ARGS="--step ..."`, тесты — `...-test`.
"""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[3]

INDICATORS = {
    "Indikator biznes-klimata Banka Rossii (IBK)": "ibk",
    "Izmeneniye izderzhek proizvodstva (obrashcheniya): balans otvetov": "costs",
    "Ozhidaniya izmeneniya tsen na gotovuyu produktsiyu (uslugi) predpriyatiy: balans otvetov": "price_exp",
}
IND_ORDER = ("ibk", "costs", "price_exp")
SHEETS = {
    "agriculture": "Сельское хозяйство",
    "mining": "Добыча",
    "manufacturing_consumer": "Обработка_потребительские",
    "manufacturing_intermediate": "Обработка_промежуточные",
    "manufacturing_investment": "Обработка_инвестиционные",
    "construction": "Строительство",
    "trade_auto": "Торговля авто (45)",
    "trade_wholesale": "Торговля оптовая (46)",
    "trade_retail": "Торговля розничная (47)",
    "transport": "Транспортировка и хранение",
    "services": "Услуги",
}
# Лист книги по размеру предприятий, куда входит лист 87→11.
GROUP_SHEETS = {
    "agriculture": "Сельское хозяйство", "mining": "Добыча", "manufacturing_consumer": "Обработка",
    "manufacturing_intermediate": "Обработка", "manufacturing_investment": "Обработка",
    "construction": "Строительство", "trade_auto": "Торговля", "trade_wholesale": "Торговля",
    "trade_retail": "Торговля", "transport": "Транспортировка и хранение", "services": "Услуги",
}
# Вопрос книги для индикатора (SA-блок): точное имя или префикс.
QUESTION = {
    "ibk": lambda q: q == "Индикатор бизнес-климата Банка России",
    "costs": lambda q: q.startswith("Как изменились издержки"),
    "price_exp": lambda q: q.startswith("Ценовые ожидания предприятий"),
}
BOOKS = ("mp_survey_data.xlsx", "survey_data_group.xlsx", "mp_survey_data_gu.xlsx")
TRAIN = ("2010-01-01", "2022-12-01")
TEST_START = "2023-01-01"
RIDGE_LAMBDAS = (0.01, 0.03, 0.1, 0.3, 1.0, 3.0, 10.0)
OFFSET_WINDOWS = (156, 60, 24, 12, 6)
NEAR_DUP_CORR = 0.995
PARITY_MODEL = "bottomup_mp__per_sheet"
PARITY_HORIZONS = (1, 2, 3, 6)
DEFAULT_PARITY_DIR = ROOT / "build" / "cb-comp-parity"
DEFAULT_OUT = ROOT / "build" / "cb-comp-per-code-review"


# ---------------------------------------------------------------------------
# Данные
# ---------------------------------------------------------------------------


def load_codes(panel_csv: Path) -> dict:
    """Per-code МП архива: `{ind: DataFrame(date × code)}`."""
    frame = pd.read_csv(panel_csv, dtype={"code": str})
    frame = frame[frame["group_stat"] == "Мониторинг предприятий"].copy()
    frame["ind"] = frame["kind"].map(INDICATORS)
    frame = frame.dropna(subset=["ind"])
    frame["date"] = pd.to_datetime(frame["date"])
    return {ind: frame[frame["ind"] == ind].pivot_table(index="date", columns="code", values="value").sort_index()
            for ind in IND_ORDER}


def load_public_codes(public_csv: Path) -> dict:
    """Публичная панель (значение листа на каждом коде): `{ind: DataFrame(date × code)}`."""
    frame = pd.read_csv(public_csv, dtype={"code": str})
    frame["date"] = pd.to_datetime(frame["date"])
    return {ind: frame[frame["indicator"] == ind].pivot_table(index="date", columns="code", values="value").sort_index()
            for ind in IND_ORDER}


def load_sheet_map(map_csv: Path) -> dict:
    frame = pd.read_csv(map_csv, dtype={"code": str})
    return dict(zip(frame["code"], frame["industry_slug"]))


def parse_book_rows(tag: str, sheet: str, rows: list) -> dict:
    """Колонки одного листа книги ЦБ: `{"tag|sheet|SA/NSA|вопрос|подзаголовок": {date: value}}`.

    Геометрия книг: строка «Отчётный период» — заголовки вопросов, строка над
    ней — маркеры блоков («Исходные данные» / «Сезонно скорректированные»),
    строка под ней с пустой колонкой A — подзаголовки (размер предприятий).
    """
    head = next((i for i, r in enumerate(rows[:12]) if r and isinstance(r[0], str) and r[0].startswith("Отчетный")), None)
    if head is None:
        return {}

    def ffill(values):
        out, cur = [], None
        for value in values:
            if value:
                cur = str(value).strip()
            out.append(cur)
        return out

    blocks = ffill(rows[head - 1])
    names = ffill(rows[head])
    subs = rows[head + 1] if rows[head + 1][0] is None else [None] * len(rows[head])
    out = {}
    for row in rows[head + 1:]:
        if not isinstance(row[0], dt.datetime):
            continue
        month = pd.Timestamp(row[0])
        for j in range(2, len(row)):
            value = row[j]
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            block = blocks[j] if j < len(blocks) else None
            kind = "SA" if block and "езонн" in block else "NSA"
            sub = (subs[j] or "").strip() if j < len(subs) and isinstance(subs[j], str) else ""
            out.setdefault(f"{tag}|{sheet}|{kind}|{names[j]}|{sub}", {})[month] = float(value)
    return out


def parse_books(books_dir: Path) -> pd.DataFrame:
    """Все колонки трёх книг ЦБ (листы спецификаций и участников пропускаются)."""
    from openpyxl import load_workbook

    columns = {}
    for name in BOOKS:
        workbook = load_workbook(books_dir / name, read_only=True, data_only=True)
        try:
            tag = name.split(".")[0]
            for sheet in workbook.sheetnames:
                if sheet.startswith("Спецификации") or sheet == "Участники опроса":
                    continue
                columns.update(parse_book_rows(tag, sheet, [list(r) for r in workbook[sheet].iter_rows(values_only=True)]))
        finally:
            workbook.close()
    return pd.DataFrame(columns).sort_index()


def book_column(books: pd.DataFrame, sheet: str, kind: str, ind: str) -> pd.Series:
    cols = [c for c in books.columns
            if c.split("|")[0] == "mp_survey_data" and c.split("|")[1] == sheet and c.split("|")[2] == kind and QUESTION[ind](c.split("|")[3])]
    if len(cols) != 1:
        raise ValueError(f"{sheet}/{kind}/{ind}: найдено колонок {len(cols)}")
    return books[cols[0]]


# ---------------------------------------------------------------------------
# Метрики и модели
# ---------------------------------------------------------------------------


def r2(y, f) -> float:
    y, f = np.asarray(y, float), np.asarray(f, float)
    return float(1 - ((y - f) ** 2).sum() / ((y - y.mean()) ** 2).sum())


def corr(a: pd.Series, b: pd.Series) -> float:
    keep = a.notna() & b.notna()
    return float(np.corrcoef(a[keep], b[keep])[0, 1]) if keep.sum() > 10 else float("nan")


def ridge_fit(x: np.ndarray, y: np.ndarray, lam: float):
    """Ridge на стандартизованных признаках; λ масштабирована на n."""
    mu, sd = x.mean(0), x.std(0) + 1e-9
    z = (x - mu) / sd
    beta = np.linalg.solve(z.T @ z + lam * len(y) * np.eye(z.shape[1]), z.T @ (y - y.mean()))
    return lambda xn: ((xn - mu) / sd) @ beta + y.mean()


def ridge_cv_lambda(x: np.ndarray, y: np.ndarray, lambdas=RIDGE_LAMBDAS) -> float:
    """λ по rolling-origin внутри train: три окна валидации по 24 месяца."""
    n, best = len(y), None
    for lam in lambdas:
        errs = []
        for cut in (n - 72, n - 48, n - 24):
            fit = ridge_fit(x[:cut], y[:cut], lam)
            errs.append(float(((y[cut:cut + 24] - fit(x[cut:cut + 24])) ** 2).mean()))
        score = float(np.mean(errs))
        if best is None or score < best[0]:
            best = (score, lam)
    return best[1]


def distinct_series(frame: pd.DataFrame, start="2010-01-01") -> dict:
    """Точные дубли и почти-дубли (corr > 0,995) среди колонок."""
    part = frame.loc[start:]
    groups = {}
    for code in part.columns:
        key = tuple(part[code].dropna().round(6).values[:40])
        groups.setdefault(key, []).append(code)
    cmat = part.corr(min_periods=60)
    cols = list(part.columns)
    near = [(a, b) for i, a in enumerate(cols) for b in cols[i + 1:] if cmat.loc[a, b] > NEAR_DUP_CORR]
    return {"codes": len(cols), "distinct": len(groups), "duplicate_groups": [g for g in groups.values() if len(g) > 1],
            "near_duplicates": near}


def residual_stats(codes: pd.DataFrame, sheet: pd.DataFrame) -> dict:
    """Медианы AR(1)/AR(12) остатка «код − лист», доля медленной компоненты, доля вне листа."""
    resid = (codes - sheet[codes.columns]).loc["2010":"2024-09"]
    resid = resid - resid.mean()
    slow = resid.rolling(12, center=True).mean()
    test = resid.loc["2023":]
    return {
        "ar1": float(resid.apply(lambda s: s.autocorr(1)).median()),
        "ar12": float(resid.apply(lambda s: s.autocorr(12)).median()),
        "slow_share": float((slow.var() / resid.var()).median()),
        "resid_share_full": float((resid.var() / codes.loc["2010":"2024-09"].var()).median()),
        "resid_share_test": float((test.var() / codes.loc["2023":].var()).median()),
    }


def offset_forecast(code: pd.Series, sheet: pd.Series, window: int, train_end="2022-12-01") -> pd.Series:
    """Лист + средний сдвиг кода за последние `window` месяцев train."""
    end = pd.Timestamp(train_end)
    lo = end - pd.DateOffset(months=window - 1)
    span = (code.index >= lo) & (code.index <= end)
    return sheet + (code[span] - sheet[span]).mean()


def ar_decay_forecast(code: pd.Series, sheet: pd.Series, test_index, train_end="2022-12-01") -> pd.Series:
    """Лист + среднее + затухающий последний известный остаток: `mu + (e_T − mu)·ρ^h`."""
    resid = code - sheet
    mu = resid.loc["2010":train_end].mean()
    rho = min(max(resid.loc[:train_end].autocorr(1), 0.0), 0.99)
    steps = np.arange(1, len(test_index) + 1)
    return sheet.reindex(test_index) + mu + (resid.loc[train_end] - mu) * rho ** steps


def dm_hln_pvalue(err_a: np.ndarray, err_b: np.ndarray, horizon: int) -> float:
    """DM-HLN двусторонний p для квадратичной потери (как в `cb_comp_parity`)."""
    from scipy import stats

    d = np.asarray(err_a, float) ** 2 - np.asarray(err_b, float) ** 2
    n, m = len(d), d.mean()
    gamma = [np.mean((d[k:] - m) * (d[:n - k] - m)) for k in range(horizon)]
    var = (gamma[0] + 2 * sum(gamma[1:])) / n
    t = m / np.sqrt(var) * np.sqrt((n + 1 - 2 * horizon + horizon * (horizon - 1) / n) / n)
    return float(2 * stats.t.sf(abs(t), n - 1))


# ---------------------------------------------------------------------------
# Шаги
# ---------------------------------------------------------------------------


def step_distinct(args) -> dict:
    return {ind: distinct_series(frame) for ind, frame in load_codes(args.parity_dir / "panel.csv").items()}


def step_forward(args) -> dict:
    codes, public = load_codes(args.parity_dir / "panel.csv"), load_public_codes(args.parity_dir / "public_panel.csv")
    smap = load_sheet_map(args.parity_dir / "sheet_map.csv")
    out = {}
    for ind in IND_ORDER:
        x, p = codes[ind], public[ind]
        rows = {}
        for slug in sorted(set(smap.values())):
            members = [c for c in x.columns if smap.get(c) == slug]
            sheet = p[members[0]].reindex(x.index)
            mean = x[members].mean(axis=1)
            train, test = x.index < TEST_START, x.index >= TEST_START
            rows[slug] = {
                "codes": len(members),
                "corr_level_train": corr(sheet[train], mean[train]), "corr_level_test": corr(sheet[test], mean[test]),
                "corr_d3_test": corr(sheet.diff(3)[test], mean.diff(3)[test]),
                "bias_train": float(np.nanmean((mean - sheet)[train])), "bias_test": float(np.nanmean((mean - sheet)[test])),
            }
        out[ind] = rows
    return out


def step_inverse(args) -> dict:
    books = parse_books(args.books_dir)
    sa = books[[c for c in books.columns if c.split("|")[2] == "SA"]].loc["2009-01-01":"2024-09-01"]
    sa = sa.loc[:, sa.notna().all()]
    codes = load_codes(args.parity_dir / "panel.csv")
    smap = load_sheet_map(args.parity_dir / "sheet_map.csv")
    train = (sa.index >= TRAIN[0]) & (sa.index <= TRAIN[1])
    test = sa.index >= TEST_START
    out = {"features": sa.shape[1]}
    for ind in IND_ORDER:
        x = codes[ind].reindex(sa.index)
        rows = []
        for code in x.columns:
            y = x[code]
            ok = y.notna()
            a, b = train & ok, test & ok
            if a.sum() < 60 or b.sum() < 15:
                continue
            slug = smap[code]
            sheet = book_column(books, SHEETS[slug], "SA", ind).reindex(sa.index)
            own = [c for c in sa.columns if c.split("|")[0] == "mp_survey_data" and c.split("|")[1] == SHEETS[slug]]
            size = [c for c in sa.columns if c.split("|")[0] == "survey_data_group" and c.split("|")[1] == GROUP_SHEETS[slug]]
            same_q = [c for c in sa.columns if c.split("|")[0] == "mp_survey_data" and QUESTION[ind](c.split("|")[3])]
            row = {"code": code, "identity": r2(y[b], sheet[b]),
                   "offset": r2(y[b], sheet[b] + (y[a] - sheet[a]).mean())}
            coef = np.polyfit(sheet[a], y[a], 1)
            row["calib"] = r2(y[b], np.polyval(coef, sheet[b]))
            for name, cols in (("own_sheet_all", own), ("own_plus_size", own + size),
                               ("same_question_all_sheets", same_q), ("all_public", list(sa.columns))):
                xa = sa.loc[a, cols].values
                lam = ridge_cv_lambda(xa, y[a].values)
                row[name] = r2(y[b], ridge_fit(xa, y[a].values, lam)(sa.loc[b, cols].values))
            rows.append(row)
        frame = pd.DataFrame(rows).set_index("code")
        out[ind] = {"median_r2": frame.median().to_dict(), "share_r2_positive": (frame > 0).mean().to_dict(), "codes": len(frame)}
    return out


def step_noise(args) -> dict:
    codes, public = load_codes(args.parity_dir / "panel.csv"), load_public_codes(args.parity_dir / "public_panel.csv")
    return {ind: residual_stats(codes[ind], public[ind].reindex(codes[ind].index)) for ind in IND_ORDER}


def step_priors(args) -> dict:
    codes, public = load_codes(args.parity_dir / "panel.csv"), load_public_codes(args.parity_dir / "public_panel.csv")
    out = {}
    for ind in IND_ORDER:
        x = codes[ind]
        s = public[ind].reindex(x.index)[x.columns]
        test = x.index >= TEST_START
        cols = [c for c in x.columns if x.loc[test, c].notna().sum() > 15]
        res = {}
        for window in OFFSET_WINDOWS:
            scores = []
            for c in cols:
                y = x.loc[test, c]
                keep = y.notna()
                scores.append(r2(y[keep], offset_forecast(x[c], s[c], window)[test][keep]))
            res[f"offset_{window}m"] = float(np.median(scores))
        scores = []
        for c in cols:
            y = x.loc[test, c]
            keep = y.notna()
            scores.append(r2(y[keep], ar_decay_forecast(x[c], s[c], x.index[test])[keep]))
        res["ar_decay"] = float(np.median(scores))
        out[ind] = res
    return out


def step_sa_inputs(args) -> dict:
    books = parse_books(args.books_dir)
    rows = []
    for slug, sheet in SHEETS.items():
        for ind in IND_ORDER:
            for kind in ("NSA", "SA"):
                series = book_column(books, sheet, kind, ind).dropna()
                rows += [{"slug": slug, "indicator": ind, "kind": kind, "date": d.date().isoformat(), "value": v}
                         for d, v in series.items()]
    args.out.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_csv(args.out / "sheets_nsa_sa.csv", index=False)
    return {"rows": len(rows), "path": str(args.out / "sheets_nsa_sa.csv")}


def sheet_variants(out_dir: Path) -> pd.DataFrame:
    base = pd.read_csv(out_dir / "sheets_nsa_sa.csv")
    base["variant"] = base["kind"].map({"NSA": "nsa_raw", "SA": "cbr_sa"})
    x13 = pd.read_csv(out_dir / "sheets_x13.csv")
    frame = pd.concat([base.drop(columns="kind"), x13])
    frame["date"] = pd.to_datetime(frame["date"])
    return frame


def archive_sheet_means(parity_dir: Path) -> pd.DataFrame:
    codes = load_codes(parity_dir / "panel.csv")
    smap = load_sheet_map(parity_dir / "sheet_map.csv")
    parts = []
    for ind, frame in codes.items():
        long = frame.stack().rename("value").reset_index()
        long["slug"] = long["code"].map(smap)
        mean = long.groupby(["slug", "date"])["value"].mean().reset_index()
        mean["indicator"] = ind
        parts.append(mean)
    return pd.concat(parts)


def step_sa_compare(args) -> dict:
    variants = sheet_variants(args.out)
    arch = archive_sheet_means(args.parity_dir).rename(columns={"value": "arch"})
    joined = variants.merge(arch, on=["slug", "indicator", "date"])
    joined = joined[(joined["date"] >= "2010-01-01") & (joined["date"] <= "2024-09-01")]
    out = {}
    for (ind, variant), part in joined.groupby(["indicator", "variant"]):
        stats = []
        for _, g in part.groupby("slug"):
            g = g.sort_values("date")
            stats.append({"corr_level": g["value"].corr(g["arch"]),
                          "corr_d1": g["value"].diff().corr(g["arch"].diff()),
                          "corr_d3": g["value"].diff(3).corr(g["arch"].diff(3)),
                          "mae": (g["value"] - g["arch"]).abs().mean()})
        out.setdefault(ind, {})[variant] = pd.DataFrame(stats).median().to_dict()
    return out


def step_panels(args) -> dict:
    smap = load_sheet_map(args.parity_dir / "sheet_map.csv")
    smap_frame = pd.DataFrame({"code": list(smap), "slug": list(smap.values())})
    variants = sheet_variants(args.out)
    keep = variants[variants["variant"].isin(["cbr_sa", "x13_2009_2024-09", "x13_2009_2026-09"])]
    oracle = archive_sheet_means(args.parity_dir).assign(variant="oracle_arch_sheet")
    out = {}
    for variant, part in pd.concat([keep, oracle]).groupby("variant"):
        panel = part.merge(smap_frame, on="slug")[["code", "indicator", "date", "value"]].dropna()
        panel["date"] = pd.to_datetime(panel["date"]).dt.strftime("%Y-%m-%d")
        path = args.out / f"panel_{variant}.csv"
        panel.to_csv(path, index=False)
        out[variant] = {"rows": len(panel), "path": str(path)}
    return out


def step_dm(args) -> dict:
    def load(variant):
        frame = pd.read_csv(args.out / f"run_{variant}" / "cb_comp_parity_predictions.csv")
        return frame[(frame["model"] == PARITY_MODEL) & (frame["sample"] == "test") & (frame["normal"] == 1)]

    base = load("cbr_sa")
    out = {}
    for variant in ("x13_2009_2024-09", "x13_2009_2026-09", "oracle_arch_sheet"):
        other = load(variant)
        for h in PARITY_HORIZONS:
            joined = base[base["horizon"] == h].merge(other[other["horizon"] == h], on=["month", "horizon"], suffixes=("_a", "_b"))
            ea, eb = joined["error_a"].values, joined["error_b"].values
            out.setdefault(variant, {})[h] = {
                "n": len(joined), "rmse": float(np.sqrt((eb ** 2).mean())), "rmse_cbr_sa": float(np.sqrt((ea ** 2).mean())),
                "ratio": float(np.sqrt((eb ** 2).mean() / (ea ** 2).mean())), "dm_p": dm_hln_pvalue(eb, ea, h),
            }
    return out


STEPS = {
    "distinct": step_distinct, "forward": step_forward, "inverse": step_inverse, "noise": step_noise,
    "priors": step_priors, "sa-inputs": step_sa_inputs, "sa-compare": step_sa_compare, "panels": step_panels, "dm": step_dm,
}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--step", choices=sorted(STEPS), required=True)
    parser.add_argument("--parity-dir", type=Path, default=DEFAULT_PARITY_DIR,
                        help="panel.csv (архив), public_panel.csv, sheet_map.csv")
    parser.add_argument("--books-dir", type=Path, default=DEFAULT_OUT / "books", help="книги ЦБ (xlsx)")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args(argv)
    result = STEPS[args.step](args)
    args.out.mkdir(parents=True, exist_ok=True)
    path = args.out / f"{args.step}.json"
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str)[:4000])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
