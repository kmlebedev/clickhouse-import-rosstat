#!/usr/bin/env python3
"""Research: отраслевой потребительский ИЦП против product-level PPI SA."""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import subprocess
import tempfile
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl

from ppi_sa_probe import parse_weights
from ppi_sa_text_anchors import THRESHOLDS, evaluate, stats

ROOT = Path(__file__).resolve().parent
R_SCRIPT = ROOT / "ppi_sa_x13.R"
R_LIB = ROOT.parent / "build" / "r-lib"
ANCHORS_PATH = ROOT.parent / "docs" / "research" / "2026-10-21-ppi-sa-text-anchors-12.csv"
TOTAL_CODE = "1323500.029.31"
MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4,
    "май": 5, "июнь": 6, "июль": 7, "август": 8,
    "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}
MISSING = {"", "...", "…", "-", "—", "х", "x"}
CONSUMER_WORDS = (
    "потребительск", "пищев", "напит", "одежд", "текстиль", "кож",
    "табач", "мебел", "полиграф", "фармацевт", "бытов",
)
NONFOOD_CODES = tuple(str(code) for code in range(13, 34) if code != 19)
VARIANT_CODES = {
    "food": ("10", "11"),
    "nonfood": NONFOOD_CODES,
    "core": ("10", "11") + NONFOOD_CODES,
    "food_no_alcohol": ("10",),
    "core_no_alcohol": ("10",) + NONFOOD_CODES,
}


def _month(value: object) -> int | None:
    text = re.sub(r"\d+\)", "", str(value or "").strip().lower()).strip()
    return MONTHS.get(text)


def _number(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if text.lower() in MISSING:
        return None
    normalized = text.replace(",", ".")
    match = re.fullmatch(r"([+-]?\d+(?:\.\d+)?)\d+\)", normalized)
    if match:
        normalized = match.group(1)
    try:
        return float(normalized)
    except ValueError as exc:
        raise ValueError(f"нечисловое значение ИЦП: {value!r}") from exc


def _sheet_year(rows: list[tuple], sheet: str) -> int:
    for row in rows[:5]:
        for cell in row:
            match = re.search(r"\bв\s+((?:19|20)\d{2})\s*г\.", str(cell or ""))
            if match:
                return int(match.group(1))
    raise ValueError(f"{sheet}: год листа не найден")


def load_ved_mm(path: Path) -> list[dict]:
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result: list[dict] = []
    seen: set[tuple[date, str]] = set()
    try:
        for sheet in (name for name in wb.sheetnames if re.fullmatch(r"[3-8]\.1", name)):
            rows = list(wb[sheet].iter_rows(values_only=True))
            year = _sheet_year(rows, sheet)
            header = next((i for i, row in enumerate(rows) if any(
                str(cell or "").strip().upper().replace(" ", "") == "ОКВЭД2" for cell in row
            )), None)
            if header is None:
                raise ValueError(f"{sheet}: заголовок ОКВЭД2 не найден")
            month_cols = {i: m for i, cell in enumerate(rows[header]) if (m := _month(cell))}
            current_code = ""
            current_name = ""
            for row in rows[header + 1:]:
                name = str(row[0] or "").strip()
                code = str(row[1] or "").strip() if len(row) > 1 else ""
                direct = bool(name and code and not code.startswith("Код"))
                if direct:
                    current_code, current_name = code, name
                elif name != "Российская Федерация" or not current_code:
                    continue
                for col, month in month_cols.items():
                    value = _number(row[col]) if col < len(row) else None
                    if value is None:
                        continue
                    observed = date(year, month, 1)
                    key = (observed, current_code)
                    if key in seen:
                        raise ValueError(f"дубль наблюдения {observed} {current_code}")
                    seen.add(key)
                    result.append({
                        "date": observed, "code": current_code,
                        "name": current_name, "value": value,
                    })
                if not direct:
                    current_code, current_name = "", ""
    finally:
        wb.close()
    return sorted(result, key=lambda r: (r["code"], r["date"]))


def load_weights(path: Path) -> dict[tuple[int, str], float]:
    return {(r["year"], r["code"]): r["weight"] for r in parse_weights(Path(path), "3")}


def weight_map(weights: dict[tuple[int, str], float], codes: tuple[str, ...], year: int) -> dict[str, float]:
    selected = {code: weights.get((year, code)) for code in codes}
    missing = [code for code, value in selected.items() if value is None]
    if missing:
        raise ValueError(f"нет весов {year}: {','.join(missing)}")
    if any(value < 0 for value in selected.values()):
        raise ValueError("отрицательный вес")
    total = sum(selected.values())
    if total <= 0:
        raise ValueError("сумма весов должна быть положительной")
    return {code: value / total for code, value in selected.items()}


def audit_candidates(rows: list[dict], weights: dict[tuple[int, str], float]) -> list[dict]:
    grouped: dict[tuple[str, str], list[date]] = defaultdict(list)
    for row in rows:
        if row["code"] == TOTAL_CODE:
            continue
        if any(word in row["name"].lower() for word in CONSUMER_WORDS):
            grouped[(row["code"], row["name"])].append(row["date"])
    result = []
    for (code, name), dates in sorted(grouped.items()):
        expected = (max(dates).year - min(dates).year) * 12 + max(dates).month - min(dates).month + 1
        result.append({
            "code": code, "name": name, "from": min(dates).isoformat(),
            "to": max(dates).isoformat(), "n": len(dates), "missing": expected - len(set(dates)),
            "has_weights_2026": (2026, code) in weights,
        })
    return result


def write_candidates(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=["code", "name", "from", "to", "n", "missing", "has_weights_2026"],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def _add_month(day: date) -> date:
    month = day.month
    return date(day.year + month // 12, month % 12 + 1, 1)


def _validated_months(days: list[date], min_months: int = 1) -> None:
    if len(days) < min_months:
        raise ValueError(f"нужно минимум {min_months} месяцев, получено {len(days)}")
    if any(day.day != 1 for day in days):
        raise ValueError("даты должны быть первыми числами месяцев")
    for previous, current in zip(days, days[1:]):
        if _add_month(previous) != current:
            raise ValueError(f"нарушена непрерывность: {previous} → {current}")


def _validated_levels(levels: dict[date, float], min_months: int = 36) -> dict[date, float]:
    result = {day: float(value) for day, value in sorted(levels.items())}
    _validated_months(list(result), min_months)
    for day, value in result.items():
        if not math.isfinite(value) or value <= 0:
            raise ValueError(f"уровень должен быть положительным: {day}={value}")
    return result


def aggregate_mm(
    rows: list[dict],
    weights: dict[tuple[int, str], float],
    year: int = 2026,
) -> dict[date, float]:
    panel: dict[str, dict[date, float]] = defaultdict(dict)
    for row in rows:
        day = row["date"]
        code = row["code"]
        if day in panel[code]:
            raise ValueError(f"дубль наблюдения {day} {code}")
        panel[code][day] = float(row["value"])
    if not panel:
        raise ValueError("нет рядов для агрегата")
    codes = tuple(sorted(panel))
    dates = set(panel[codes[0]])
    for code in codes[1:]:
        if set(panel[code]) != dates:
            raise ValueError(f"неполный набор компонент: {code}")
    normalized = weight_map(weights, codes, year)
    return {
        day: sum(normalized[code] * panel[code][day] for code in codes)
        for day in sorted(dates)
    }


def levels_from_mm(points: dict[date, float], base: float = 100.0) -> dict[date, float]:
    days = sorted(points)
    _validated_months(days, 36)
    values = {day: float(points[day]) for day in days}
    for day, value in values.items():
        if not math.isfinite(value):
            raise ValueError(f"нечисловой m/m: {day}")
    level = float(base)
    result: dict[date, float] = {}
    previous = None
    for day in days:
        if previous is not None:
            level *= values[day] / 100
        if level <= 0 or not math.isfinite(level):
            raise ValueError(f"уровень должен быть положительным: {day}={level}")
        result[day] = level
        previous = day
    return result


def mm_from_levels(levels: dict[date, float]) -> dict[date, float | None]:
    checked = _validated_levels(levels, 1)
    days = list(checked)
    result: dict[date, float | None] = {days[0]: None}
    for previous, current in zip(days, days[1:]):
        result[current] = (checked[current] / checked[previous] - 1) * 100
    return result


def saar_from_mm(mm_sa: dict[date, float | None]) -> dict[date, float]:
    result = {}
    for day in sorted(mm_sa):
        value = mm_sa[day]
        if value is None:
            continue
        result[day] = ((1 + value / 100) ** 12 - 1) * 100
    return result


def saar3_from_mm(mm_sa: dict[date, float | None]) -> dict[date, float]:
    saar = saar_from_mm(mm_sa)
    days = sorted(saar)
    result = {}
    for index, day in enumerate(days):
        window = days[max(0, index - 2):index + 1]
        result[day] = sum(saar[item] for item in window) / len(window)
    return result


def _strict_float(value: object, label: str) -> float:
    text = str(value or "").strip()
    if text in {"", "NA", "NaN", "None"}:
        raise ValueError(f"{label}: пустое или нечисловое значение {value!r}")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{label}: нечисловое значение {value!r}") from exc


def seasonal_adjust(levels: dict[date, float]) -> dict[date, float]:
    checked = _validated_levels(levels)
    with tempfile.TemporaryDirectory(prefix="ppi-industry-x13-") as tmp:
        input_path = Path(tmp) / "series.csv"
        output_path = Path(tmp) / "sa.csv"
        with input_path.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=["series", "date", "level"], lineterminator="\n")
            writer.writeheader()
            for day, value in checked.items():
                writer.writerow({"series": "industry", "date": day.isoformat(), "level": value})
        env = dict(os.environ)
        env["R_LIBS_USER"] = str(R_LIB)
        try:
            subprocess.run(
                ["Rscript", str(R_SCRIPT), str(input_path), str(output_path)],
                check=True,
                env=env,
                cwd=ROOT,
            )
        except (subprocess.CalledProcessError, OSError) as exc:
            raise RuntimeError(f"X-13 seats_log_cal завершился ошибкой: {exc}") from exc
        if not output_path.exists():
            raise ValueError("X-13 вернул пустой результат")
        values: dict[date, float] = {}
        with output_path.open(encoding="utf-8", newline="") as fh:
            for row in csv.DictReader(fh):
                if row["series"] != "industry" or row["variant"] != "seats_log_cal":
                    continue
                day = date.fromisoformat(row["date"])
                values[day] = _strict_float(row["sa"], f"X-13 {day}")
    if not values:
        raise ValueError("X-13 вернул пустой seats_log_cal")
    missing = set(checked) - set(values)
    if missing:
        raise ValueError(f"X-13 не вернул даты: {min(missing)}")
    return _validated_levels(values)


def _rows_for_codes(rows: list[dict], codes: tuple[str, ...]) -> list[dict]:
    selected = [row for row in rows if row["code"] in codes]
    found = {row["code"] for row in selected}
    missing = set(codes) - found
    if missing:
        raise ValueError(f"нет отраслевых рядов: {','.join(sorted(missing))}")
    return selected


def _saar3_from_sa(levels: dict[date, float]) -> dict[date, float]:
    return saar3_from_mm(mm_from_levels(levels))


def direct_saar3(
    rows: list[dict],
    weights: dict[tuple[int, str], float],
    codes: tuple[str, ...],
    year: int = 2026,
    adjust=seasonal_adjust,
) -> dict[date, float]:
    aggregated = aggregate_mm(_rows_for_codes(rows, codes), weights, year)
    levels = levels_from_mm(aggregated)
    return _saar3_from_sa(adjust(levels))


def indirect_saar3(
    rows: list[dict],
    weights: dict[tuple[int, str], float],
    codes: tuple[str, ...],
    year: int = 2026,
    adjust=seasonal_adjust,
    sa_cache: dict[str, dict[date, float]] | None = None,
) -> dict[date, float]:
    normalized = weight_map(weights, codes, year)
    selected = _rows_for_codes(rows, codes)
    components: dict[str, dict[date, float]] = {}
    for code in codes:
        if sa_cache is not None and code in sa_cache:
            components[code] = sa_cache[code]
            continue
        component_rows = [row for row in selected if row["code"] == code]
        levels = levels_from_mm({row["date"]: float(row["value"]) for row in component_rows})
        adjusted = _validated_levels(adjust(levels))
        components[code] = adjusted
        if sa_cache is not None:
            sa_cache[code] = adjusted
    dates = set(components[codes[0]])
    for code in codes[1:]:
        if set(components[code]) != dates:
            raise ValueError(f"неполный набор SA компонент: {code}")
    aggregate = {
        day: sum(normalized[code] * components[code][day] for code in codes)
        for day in sorted(dates)
    }
    return _saar3_from_sa(aggregate)


def build_variants(
    rows: list[dict],
    weights: dict[tuple[int, str], float],
    year: int = 2026,
    adjust=seasonal_adjust,
) -> dict[str, dict[str, dict[date, float]]]:
    cache: dict[str, dict[date, float]] = {}
    direct = {
        name: direct_saar3(rows, weights, VARIANT_CODES[name], year, adjust)
        for name in ("food", "nonfood", "core")
    }
    indirect = {
        name: indirect_saar3(rows, weights, VARIANT_CODES[name], year, adjust, cache)
        for name in ("food", "nonfood", "core")
    }
    direct_no_alcohol = {
        "food": direct_saar3(rows, weights, VARIANT_CODES["food_no_alcohol"], year, adjust),
        "nonfood": direct["nonfood"],
        "core": direct_saar3(rows, weights, VARIANT_CODES["core_no_alcohol"], year, adjust),
    }
    indirect_no_alcohol = {
        "food": indirect_saar3(rows, weights, VARIANT_CODES["food_no_alcohol"], year, adjust, cache),
        "nonfood": indirect["nonfood"],
        "core": indirect_saar3(rows, weights, VARIANT_CODES["core_no_alcohol"], year, adjust, cache),
    }
    return {
        "direct": direct,
        "indirect": indirect,
        "direct_no_alcohol": direct_no_alcohol,
        "indirect_no_alcohol": indirect_no_alcohol,
    }


ANCHOR_FIELDS = (
    "issue", "url", "target", "window", "series", "kind", "text",
    "anchor_value", "ours", "diff", "threshold", "result", "note",
)
QUALITATIVE_KINDS = {"near_zero", "near_4", "below_4", "declining", "almost_flat", "direction_up"}


def load_anchors(path: Path) -> list[dict]:
    path = Path(path)
    with path.open(encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None or set(ANCHOR_FIELDS) - set(reader.fieldnames):
            raise ValueError("контракт якорей: нет обязательных колонок")
        rows = []
        for raw in reader:
            kind = raw["kind"]
            if kind not in QUALITATIVE_KINDS | {"exact", "not_stated"}:
                raise ValueError(f"контракт якорей: неизвестный kind {kind!r}")
            if raw["series"] not in {"food", "nonfood", "core"}:
                raise ValueError(f"контракт якорей: неизвестный series {raw['series']!r}")
            target = date.fromisoformat(raw["target"])
            if target.day != 1:
                raise ValueError(f"контракт якорей: target не первый день месяца {target}")
            row = dict(raw)
            row["target"] = target.isoformat()
            row["ours"] = _strict_float(raw["ours"], "anchor ours")
            row["anchor_value"] = (
                _strict_float(raw["anchor_value"], "anchor value") if kind == "exact" else ""
            )
            rows.append(row)
    exact = sum(row["kind"] == "exact" for row in rows)
    qualitative = sum(row["kind"] in QUALITATIVE_KINDS for row in rows)
    not_stated = sum(row["kind"] == "not_stated" for row in rows)
    if (len(rows), exact, qualitative, not_stated) != (37, 11, 14, 12):
        raise ValueError(
            "контракт якорей: ожидалось 37/11 exact/14 qualitative/12 not_stated, "
            f"получено {len(rows)}/{exact}/{qualitative}/{not_stated}"
        )
    return rows


def _previous_available(points: dict[date, float], target: date) -> float | None:
    earlier = [day for day in points if day < target]
    return points[max(earlier)] if earlier else None


def _comparison_rows(anchors: list[dict], values: dict[str, dict[date, float]]) -> list[dict]:
    rows = []
    for anchor in anchors:
        series = anchor["series"]
        target = anchor["target"]
        target_day = date.fromisoformat(target)
        points = values.get(series)
        if points is None:
            raise ValueError(f"нет saar3: series {series}")
        current = points.get(target_day)
        if current is None:
            raise ValueError(f"нет saar3: {series} {target}")
        previous = _previous_available(points, target_day) if anchor["kind"] == "direction_up" else None
        result, threshold = evaluate(anchor["kind"], current, previous)
        rows.append({
            "variant": "",
            "issue": anchor["issue"],
            "url": anchor["url"],
            "target": target,
            "window": anchor["window"],
            "series": series,
            "kind": anchor["kind"],
            "text": anchor["text"],
            "anchor_value": anchor["anchor_value"],
            "ours": current,
            "diff": current - anchor["anchor_value"] if anchor["kind"] == "exact" else "",
            "threshold": threshold,
            "result": result,
            "note": anchor["note"],
        })
    return rows


def _metric_block(rows: list[dict]) -> dict:
    exact = [row for row in rows if row["kind"] == "exact"]
    qualitative = [row for row in rows if row["kind"] in QUALITATIVE_KINDS]
    summary = {"pass": 0, "fail": 0, "n/a": 0}
    for row in qualitative:
        summary[row["result"]] = summary.get(row["result"], 0) + 1
    return {
        "rows": rows,
        "exact": {
            "rows": exact,
            "by_series": {
                series: stats([row for row in exact if row["series"] == series])
                for series in ("food", "nonfood", "core")
            },
            "overall": stats(exact),
        },
        "qualitative": {"rows": qualitative, "summary": summary},
    }


def baseline_metrics(anchors: list[dict]) -> dict:
    values: dict[str, dict[date, float]] = {}
    for anchor in anchors:
        values.setdefault(anchor["series"], {})[date.fromisoformat(anchor["target"])] = anchor["ours"]
    return _metric_block(_comparison_rows(anchors, values))


def compare_anchors(
    variants: dict[str, dict[str, dict[date, float]]],
    anchors_path: Path,
) -> dict:
    anchors = load_anchors(anchors_path)
    report = {
        "thresholds": THRESHOLDS,
        "baseline": baseline_metrics(anchors),
        "variants": {},
    }
    for name, values in variants.items():
        rows = _comparison_rows(anchors, values)
        for row in rows:
            row["variant"] = name
        report["variants"][name] = _metric_block(rows)
    return report


def write_comparison_csv(path: Path, comparison: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "variant", "target", "series", "kind", "anchor_value", "ours", "diff",
        "result", "issue", "url", "window", "text", "threshold", "note",
    ]
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for block in comparison["variants"].values():
            writer.writerows(block["rows"])


def write_comparison_json(path: Path, comparison: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(comparison, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ved-xlsx", type=Path, default=Path("build/ppi-industry-sensitivity/raw/ved.xlsx"))
    parser.add_argument("--weights-xlsx", type=Path, default=Path("build/ppi-industry-sensitivity/raw/vesa.xlsx"))
    parser.add_argument("--anchors", type=Path, default=ANCHORS_PATH)
    parser.add_argument("--output-prefix", type=Path, default=Path("docs/research/2026-10-21-ppi-industry-sensitivity"))
    parser.add_argument("--print", action="store_true")
    args = parser.parse_args()
    rows = load_ved_mm(args.ved_xlsx)
    weights = load_weights(args.weights_xlsx)
    candidates = audit_candidates(rows, weights)
    write_candidates(args.output_prefix.parent / "2026-10-21-ppi-ved-candidates.csv", candidates)
    variants = build_variants(rows, weights)
    comparison = compare_anchors(variants, args.anchors)
    write_comparison_csv(args.output_prefix.with_suffix(".csv"), comparison)
    write_comparison_json(args.output_prefix.with_suffix(".json"), comparison)
    if args.print:
        print(f"VED rows: {len(rows)}; candidates: {len(candidates)}")
        baseline = comparison["baseline"]
        print(
            "baseline: exact MAE "
            f"{baseline['exact']['overall']['mae']:.3f}; "
            f"qualitative {baseline['qualitative']['summary']}"
        )
        for name, block in comparison["variants"].items():
            print(
                f"{name}: exact MAE {block['exact']['overall']['mae']:.3f}; "
                f"qualitative {block['qualitative']['summary']}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
