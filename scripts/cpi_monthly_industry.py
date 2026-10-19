#!/usr/bin/env python3
"""CPI monthly industry: отраслевые МП-признаки в компонентном nowcast ИПЦ.

Stage 2 пункта `key-rate/phase-4-ml-components` (issue #235, родитель #76),
спека `docs/superpowers/specs/2026-10-19-cpi-industry-mp-features-design.md`.

Идея: 16 публичных отраслевых листов `mp_survey_data.xlsx` (IBK, ценовые
ожидания, издержки) дают месячные признаки за M−1, доступные к cutoff
8-го M+1 — той же границе, что у агрегатного `cbr_survey`. Скрипт добавляет
96 SA-признаков (16 листов × 3 индикатора × уровень/изменение) к v2/v3
моделям и сравнивает aggregate vs industry на общих фолдах.

Выборки:
* `long` — 2013-10…2026-08, start 48, 107 фолдов (primary: отраслевая панель
  полна, разреженности недельных рядов нет);
* `weekly` — 2022-01…2026-08, start 24, 32 фолда (secondary: сопоставимость
  со Stage 1).

As-of отраслевой точки `P` — конец `P+1` (`cpi_monthly.available_from_survey`):
на cutoff(M) доступны `M−1` и раньше, точка `M` не участвует. Отсутствующая
точка/лист — NaN, не 0; строка выборки не выбрасывается.

Скрипт только читает ClickHouse и публичный xlsx; в БД данные не пишет.
Запуск — `make cpi-monthly-industry` (env подключает Makefile), тесты —
`scripts/cpi_monthly_industry_test.py`.
"""

import argparse
import dataclasses
import json
import math
import os
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

import numpy as np
from catboost import CatBoostRegressor
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sklearn.exceptions import ConvergenceWarning
from sklearn.linear_model import ElasticNetCV
from sklearn.model_selection import TimeSeriesSplit
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import cpi_monthly as cm
import cpi_monthly_v3 as cmv3
import eval_utils as eu
import ml_experiment as mlx
from cpi_public_panel_audit import (
    INDUSTRY_SHEETS,
    MP_INDICATORS,
    MP_URL,
    resolve_mp_survey,
    sha256_file,
)

SEED = cm.SEED
SAMPLES = ("long", "weekly")
DEFAULT_START = {"long": 48, "weekly": 24}
BOOTSTRAP_BLOCK = 3
BOOTSTRAP_N = 2000
NAN = float("nan")

INDUSTRY_SLUGS = {
    "Промышленность": "industry",
    "Добыча": "mining",
    "Обработка": "manufacturing",
    "Обработка_потребительские": "mfg_consumer",
    "Обработка_инвестиционные": "mfg_investment",
    "Обработка_промежуточные": "mfg_intermediate",
    "ЭЭГП": "eegp",
    "Водоснабжение": "water",
    "Сельское хозяйство": "agriculture",
    "Строительство": "construction",
    "Торговля (45+46+47)": "trade_all",
    "Торговля авто (45)": "trade_auto",
    "Торговля оптовая (46)": "trade_wholesale",
    "Торговля розничная (47)": "trade_retail",
    "Транспортировка и хранение": "transport",
    "Услуги": "services",
}

INDICATOR_KEYS = ("ibk", "price_exp", "costs")
INDICATOR_DEFS = {
    "ibk": MP_INDICATORS["ibk"],
    "price_exp": MP_INDICATORS["price_expectations"],
    "costs": MP_INDICATORS["costs"],
}

COMPONENT_SHEETS = {
    "cpi_food": (
        "Сельское хозяйство",
        "Обработка_потребительские",
        "Торговля розничная (47)",
    ),
    "cpi_nonfood": (
        "Промышленность",
        "Добыча",
        "Обработка",
        "Обработка_инвестиционные",
        "Обработка_промежуточные",
        "Строительство",
        "Торговля авто (45)",
    ),
    "cpi_services": (
        "ЭЭГП",
        "Водоснабжение",
        "Транспортировка и хранение",
        "Услуги",
    ),
}
GENERAL_SHEETS = (
    "Торговля (45+46+47)",
    "Торговля оптовая (46)",
)

PRIMARY_PAIRS = (
    ("v3_ind_bottomup", "v3_bottomup"),
    ("ind_only_bottomup", "v2_bottomup"),
)
SECONDARY_PAIRS = (
    ("v3_ind_direct", "v3_direct"),
    ("ind_direct", "v2_baseline"),
    ("ind_bottomup", "v2_bottomup"),
    ("v3_ind_all_bottomup", "v3_ind_bottomup"),
    ("B3_ind", "B3"),
)

STAGE1_REFERENCE = {
    "weekly": {"v2_bottomup": 0.362300, "v3_bottomup": 0.311487},
    "long": {"v2_bottomup": 0.688869, "v3_bottomup": 0.642997},
}
REGRESSION_TOLERANCE = 1e-6

DEFAULT_MODELS = {
    "weekly": (
        "v2_baseline",
        "v2_bottomup",
        "v3_direct",
        "v3_bottomup",
        "ind_direct",
        "ind_bottomup",
        "ind_only_bottomup",
        "v3_ind_direct",
        "v3_ind_bottomup",
        "v3_ind_all_bottomup",
        "B1",
        "B2",
        "B3",
        "B3_weekly",
        "B3_ind",
    ),
    "long": (
        "v2_bottomup",
        "v3_bottomup",
        "ind_only_bottomup",
        "ind_bottomup",
        "v3_ind_bottomup",
        "v3_ind_all_bottomup",
        "B1",
    ),
}


# ---------------------------------------------------------------------------
# Загрузка отраслевой панели
# ---------------------------------------------------------------------------


def cell_date(value) -> date:
    """Дата из ячейки xlsx: `datetime`/`date`/Excel-serial; иначе None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            result = from_excel(value)
        except (TypeError, ValueError, OverflowError):
            return None
        if isinstance(result, datetime):
            return result.date()
        if isinstance(result, date):
            return result
    return None


def marker_positions(header_row) -> tuple:
    """Индексы блоков «Исходные данные» / «Сезонно скорректированные данные»."""
    raw_from = -1
    sa_from = -1
    for index, value in enumerate(header_row):
        text = str(value).strip() if value is not None else ""
        if text == "Исходные данные" and raw_from < 0:
            raw_from = index
        if text == "Сезонно скорректированные данные" and sa_from < 0:
            sa_from = index
    return raw_from, sa_from


def find_column(header: list, start: int, end: int, needle: str, exclude=()) -> int:
    """Индекс колонки с `needle` в блоке `[start, end)`, без `exclude`."""
    needle = needle.lower()
    for index in range(start, min(end, len(header))):
        text = str(header[index]).strip().lower() if header[index] is not None else ""
        if needle in text and not any(item in text for item in exclude):
            return index
    raise ValueError(f"нет колонки {needle!r} в блоке {start}:{end}")


def extract_points(rows: list, column: int) -> list:
    """`[(month_start, value)]` из строк данных; нечисловые/пустые пропускаются."""
    points = {}
    for row in rows:
        if not row or row[0] is None or column >= len(row):
            continue
        day = cell_date(row[0])
        if day is None:
            continue
        value = row[column]
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        if isinstance(value, float) and math.isnan(value):
            continue
        points[day.replace(day=1)] = float(value)
    return sorted(points.items())


def load_industry_panel(path, sheets=INDUSTRY_SHEETS) -> dict:
    """`{sheet: {indicator: [(month_start, value)]}}` из SA-блока книги.

    Геометрия публичной книги: строка 3 — маркеры блоков, строка 4 — заголовки,
    данные — с 5-й строки. Дата — колонка A (Excel-serial или datetime).
    Отсутствие листа, маркера или колонки — `ValueError`: это структура книги,
    а не пропуск отдельной точки.
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        available = set(workbook.sheetnames)
        missing = [sheet for sheet in sheets if sheet not in available]
        if missing:
            raise ValueError(f"{path}: нет листов {missing}")
        panel = {}
        for sheet in sheets:
            rows = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
            if len(rows) < 5:
                raise ValueError(f"{path}:{sheet}: меньше 5 строк")
            raw_from, sa_from = marker_positions(rows[2])
            if raw_from < 0 or sa_from <= raw_from:
                raise ValueError(f"{path}:{sheet}: не найдены маркеры блоков")
            header = rows[3]
            indicators = {}
            for key, (needle, exclude) in INDICATOR_DEFS.items():
                column = find_column(header, sa_from, len(header), needle, exclude)
                points = extract_points(rows[4:], column)
                if not points:
                    raise ValueError(f"{path}:{sheet}: пустой ряд {key!r}")
                indicators[key] = points
            panel[sheet] = indicators
        return panel
    finally:
        workbook.close()


def panel_meta(panel: dict, path: Path, source: str) -> dict:
    """Метаданные панели для отчёта и паспорта run."""
    dates = [
        day
        for sheet in panel.values()
        for points in sheet.values()
        for day, _ in points
    ]
    return {
        "url": MP_URL,
        "path": str(path),
        "source": source,
        "sha256": sha256_file(path),
        "n_sheets": len(panel),
        "n_months": len({day for day in dates}),
        "min_date": min(dates).isoformat() if dates else None,
        "max_date": max(dates).isoformat() if dates else None,
    }


# ---------------------------------------------------------------------------
# Признаки и mapping
# ---------------------------------------------------------------------------


def industry_feature_names(sheets=INDUSTRY_SHEETS) -> list:
    """96 имён: лист × индикатор × (уровень, изменение)."""
    names = []
    for sheet in sheets:
        slug = INDUSTRY_SLUGS[sheet]
        for key in INDICATOR_KEYS:
            names.append(f"ind_{slug}_{key}_sa")
            names.append(f"ind_{slug}_{key}_sa_chg")
    return names


def component_sheets(component=None) -> tuple:
    """Листы компонента: mapped-группа + general; `None` — все 16."""
    if component is None:
        return tuple(INDUSTRY_SHEETS)
    return tuple(COMPONENT_SHEETS[component]) + tuple(GENERAL_SHEETS)


def base_feature_names(include_mp: bool, include_weekly: bool) -> list:
    """Имена базы v2/v3 с `month` последним."""
    names = cm.feature_names(with_mp=include_mp)
    if include_weekly:
        names = names[:-1] + list(cmv3.WEEKLY_FEATURES) + [names[-1]]
    return names


def feature_names(
    component=None,
    include_mp: bool = True,
    include_weekly: bool = False,
    all_sheets: bool = False,
) -> list:
    """База + industry-признаки + `month` последним.

    `component=None` или `all_sheets=True` — все 16 листов (direct и
    robustness-вариант); иначе — mapped-группа компонента + general.
    """
    base = base_feature_names(include_mp, include_weekly)
    sheets = (
        tuple(INDUSTRY_SHEETS)
        if all_sheets or component is None
        else component_sheets(component)
    )
    return base[:-1] + industry_feature_names(sheets) + [base[-1]]


def eligible_point(points: dict, day: date, cutoff: date):
    """Значение точки `day`, если она доступна к `cutoff`; иначе None."""
    value = points.get(day)
    if value is None:
        return None
    if cm.available_from_survey(day) > cutoff:
        return None
    return value


def industry_features(panel: dict, month: date, cutoff: date) -> dict:
    """96 признаков месяца `M`: уровень `M−1`, изменение `M−1 − M−2`.

    Отсутствующая точка/лист — NaN, не 0. Точка `M` не участвует: уровень
    берётся строго за предыдущий месяц, а доступность `M` (конец `M+1`) позже
    cutoff. Это правило и проверяет `assert_no_industry_leakage`.
    """
    level_month = cm.previous_month(month)
    base_month = cm.previous_month(level_month)
    result = {}
    for sheet in INDUSTRY_SHEETS:
        slug = INDUSTRY_SLUGS[sheet]
        data = panel.get(sheet, {})
        for key in INDICATOR_KEYS:
            points = dict(data.get(key, []))
            level = eligible_point(points, level_month, cutoff)
            base = eligible_point(points, base_month, cutoff)
            result[f"ind_{slug}_{key}_sa"] = level if level is not None else NAN
            result[f"ind_{slug}_{key}_sa_chg"] = (
                level - base
                if level is not None and base is not None
                else NAN
            )
    return result


def build_sample_industry(sample: str, panel: dict) -> tuple:
    """Строки v3 + industry-признаки; industry-NaN строку не выбрасывает."""
    if sample not in SAMPLES:
        raise ValueError(f"неизвестная выборка {sample!r}")
    rows, skipped = cmv3.build_sample_v3(sample)
    enriched = []
    for row in rows:
        extra = industry_features(panel, row.month, cm.cutoff_for(row.month))
        enriched.append(
            dataclasses.replace(row, features={**row.features, **extra})
        )
    return enriched, skipped


def assert_no_industry_leakage(rows: list, panel: dict) -> None:
    """Ни одна точка `date >= M`, доступная к cutoff(M), не должна участвовать."""
    for row in rows:
        cutoff = cm.cutoff_for(row.month)
        for sheet in INDUSTRY_SHEETS:
            data = panel.get(sheet, {})
            for key in INDICATOR_KEYS:
                for day, _ in data.get(key, []):
                    if day >= row.month and cm.available_from_survey(day) <= cutoff:
                        raise AssertionError(
                            f"утечка: {sheet}/{key} {day} доступна к cutoff "
                            f"{cutoff} месяца {row.month}"
                        )


# ---------------------------------------------------------------------------
# Модели
# ---------------------------------------------------------------------------


def feature_matrix(rows: list, names: list) -> np.ndarray:
    return np.array(
        [[row.features[name] for name in names] for row in rows], dtype=object
    )


def fit_catboost_industry(
    train: list,
    row: cm.Row,
    target: str,
    component=None,
    include_mp: bool = True,
    include_weekly: bool = False,
    all_sheets: bool = False,
    **overrides,
) -> float:
    """CatBoost на выбранном наборе; `month` категориальным, NaN нативно."""
    names = feature_names(component, include_mp, include_weekly, all_sheets)
    params = {
        "loss_function": "RMSE",
        "random_seed": SEED,
        "verbose": False,
        "allow_writing_files": False,
        "cat_features": [names.index("month")],
    }
    params.update(overrides)
    model = CatBoostRegressor(**params)
    model.fit(
        feature_matrix(train, names),
        np.asarray([float(item.y[target]) for item in train], dtype=float),
    )
    return float(model.predict(feature_matrix([row], names))[0])


def fit_bottomup_industry(
    train: list,
    row: cm.Row,
    include_mp: bool = True,
    include_weekly: bool = False,
    all_sheets: bool = False,
) -> float:
    """`Σ w_j·ŷ_j`: компонентные CatBoost + NNLS-веса обучающего окна."""
    weights = cm.nnls_weights(train)
    components = tuple(
        fit_catboost_industry(
            train,
            row,
            code,
            component=code,
            include_mp=include_mp,
            include_weekly=include_weekly,
            all_sheets=all_sheets,
        )
        for code in cm.BOTTOMUP_COMPONENTS
    )
    fit_bottomup_industry.last_weights = tuple(weights)
    fit_bottomup_industry.last_components = components
    return float(
        sum(weight * value for weight, value in zip(weights, components))
    )


fit_bottomup_industry.last_weights = ()
fit_bottomup_industry.last_components = ()


def fit_ind_direct(train: list, row: cm.Row) -> float:
    return fit_catboost_industry(
        train, row, cm.TARGET_ALL, component=None, include_mp=True
    )


def fit_ind_bottomup(train: list, row: cm.Row) -> float:
    value = fit_bottomup_industry(
        train, row, include_mp=True, include_weekly=False
    )
    fit_ind_bottomup.last_weights = fit_bottomup_industry.last_weights
    fit_ind_bottomup.last_components = fit_bottomup_industry.last_components
    return value


def fit_ind_only_bottomup(train: list, row: cm.Row) -> float:
    value = fit_bottomup_industry(
        train, row, include_mp=False, include_weekly=False
    )
    fit_ind_only_bottomup.last_weights = fit_bottomup_industry.last_weights
    fit_ind_only_bottomup.last_components = fit_bottomup_industry.last_components
    return value


def fit_v3_ind_direct(train: list, row: cm.Row) -> float:
    return fit_catboost_industry(
        train, row, cm.TARGET_ALL, component=None, include_mp=True,
        include_weekly=True,
    )


def fit_v3_ind_bottomup(train: list, row: cm.Row) -> float:
    value = fit_bottomup_industry(
        train, row, include_mp=True, include_weekly=True
    )
    fit_v3_ind_bottomup.last_weights = fit_bottomup_industry.last_weights
    fit_v3_ind_bottomup.last_components = fit_bottomup_industry.last_components
    return value


def fit_v3_ind_all_bottomup(train: list, row: cm.Row) -> float:
    value = fit_bottomup_industry(
        train, row, include_mp=True, include_weekly=True, all_sheets=True
    )
    fit_v3_ind_all_bottomup.last_weights = fit_bottomup_industry.last_weights
    fit_v3_ind_all_bottomup.last_components = fit_bottomup_industry.last_components
    return value


for _function in (
    fit_ind_bottomup,
    fit_ind_only_bottomup,
    fit_v3_ind_bottomup,
    fit_v3_ind_all_bottomup,
):
    _function.last_weights = ()
    _function.last_components = ()


def enet_impute(train_matrix: np.ndarray, row_matrix: np.ndarray) -> tuple:
    """Импутация NaN медианой train; колонка без наблюдений — 0.0.

    Медианы считаются ТОЛЬКО по обучающему окну фолда; тестовая строка в них
    не участвует. Возврат — `(train_imputed, row_imputed, medians)`; медианы
    нужны тесту NaN-политики.
    """
    medians = np.nanmedian(train_matrix, axis=0)
    medians = np.where(np.isnan(medians), 0.0, medians)
    return (
        np.where(np.isnan(train_matrix), medians, train_matrix),
        np.where(np.isnan(row_matrix), medians, row_matrix),
        medians,
    )


def fit_enet_industry(
    train: list,
    row: cm.Row,
    component=None,
    include_mp: bool = False,
    include_weekly: bool = False,
    all_sheets: bool = False,
) -> float:
    """ElasticNetCV на стандартизованных признаках; NaN — медиана train."""
    names = [
        name
        for name in feature_names(
            component, include_mp, include_weekly, all_sheets
        )
        if name != "month"
    ]
    if len(train) < cm.ENET_MIN_TRAIN:
        raise cm.DegenerateWindowError(
            f"B3 (ElasticNetCV): обучающее окно {len(train)} месяцев короче "
            f"{cm.ENET_MIN_TRAIN}"
        )
    train_matrix = np.asarray(
        [[float(item.features[name]) for name in names] for item in train],
        dtype=float,
    )
    row_matrix = np.asarray(
        [[float(row.features[name]) for name in names]], dtype=float
    )
    train_imputed, row_imputed, medians = enet_impute(train_matrix, row_matrix)
    fit_enet_industry.last_medians = medians
    targets = np.asarray(
        [float(item.y[cm.TARGET_ALL]) for item in train], dtype=float
    )
    model = Pipeline(
        [
            ("scaler", StandardScaler()),
            (
                "enet",
                ElasticNetCV(
                    alphas=cm.ENET_ALPHAS,
                    l1_ratio=list(cm.ENET_L1_RATIOS),
                    cv=TimeSeriesSplit(cm.ENET_CV_SPLITS),
                    random_state=SEED,
                ),
            ),
        ]
    )
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        model.fit(train_imputed, targets)
    non_converged = any(
        issubclass(warning.category, ConvergenceWarning) for warning in caught
    )
    for warning in caught:
        if not issubclass(warning.category, ConvergenceWarning):
            warnings.warn(str(warning.message), warning.category, stacklevel=2)
    fit_enet_industry.fold_notes = {cm.NON_CONVERGED: int(non_converged)}
    return float(model.predict(row_imputed)[0])


fit_enet_industry.fold_notes = {}
fit_enet_industry.last_medians = None


def fit_b3_ind(train: list, row: cm.Row) -> float:
    return fit_enet_industry(
        train, row, component=None, include_mp=False, include_weekly=False
    )


MODELS = {
    "v2_baseline": cm.fit_catboost_direct_mp,
    "v2_bottomup": cm.fit_catboost_bottomup_mp,
    "v3_direct": cmv3.fit_v3_direct,
    "v3_bottomup": cmv3.fit_v3_bottomup,
    "ind_direct": fit_ind_direct,
    "ind_bottomup": fit_ind_bottomup,
    "ind_only_bottomup": fit_ind_only_bottomup,
    "v3_ind_direct": fit_v3_ind_direct,
    "v3_ind_bottomup": fit_v3_ind_bottomup,
    "v3_ind_all_bottomup": fit_v3_ind_all_bottomup,
    "B1": cm.predict_last,
    "B2": cm.fit_ar1,
    "B3": cm.fit_enet_base,
    "B3_weekly": cmv3.fit_b3_weekly,
    "B3_ind": fit_b3_ind,
}


# ---------------------------------------------------------------------------
# Оценка
# ---------------------------------------------------------------------------


def _errors(predictions: list, model: str) -> tuple:
    group = [item for item in predictions if item.model == model]
    group.sort(key=lambda item: item.month)
    return (
        [item.month for item in group],
        [item.y_pred - item.y_true for item in group],
        group,
    )


def _mde_ratio(err_a: list, err_b: list) -> float:
    if len(err_a) < 2:
        return NAN
    differences = [a * a - b * b for a, b in zip(err_a, err_b)]
    sd = float(np.std(differences, ddof=1))
    rmse_a = eu.rmse(err_a)
    rmse_b = eu.rmse(err_b)
    mde_mse = 2.802 * sd / math.sqrt(len(differences))
    return mde_mse / (rmse_a + rmse_b) if (rmse_a + rmse_b) else NAN


def _pair(predictions: list, left: str, right: str, primary: bool) -> dict:
    months_left, errors_left, _ = _errors(predictions, left)
    months_right, errors_right, _ = _errors(predictions, right)
    common = sorted(set(months_left) & set(months_right))
    if not common:
        return {"n": 0, "primary": primary}
    by_month_left = dict(zip(months_left, errors_left))
    by_month_right = dict(zip(months_right, errors_right))
    err_left = [by_month_left[month] for month in common]
    err_right = [by_month_right[month] for month in common]
    dm_stat, dm_p = eu.diebold_mariano(err_left, err_right, 1)
    boot = eu.block_bootstrap_rmse_diff(
        err_left, err_right, BOOTSTRAP_BLOCK, BOOTSTRAP_N, SEED
    )
    return {
        "n": len(common),
        "primary": primary,
        "rmse_left": eu.rmse(err_left),
        "rmse_right": eu.rmse(err_right),
        "delta": boot.rmse_diff,
        "delta_ci": [boot.low, boot.high],
        "dm_stat": dm_stat,
        "dm_p": dm_p,
        "mde_ratio": _mde_ratio(err_left, err_right),
    }


def holm_adjust(p_values: list) -> list:
    """Поправка Холма; NaN считается незначимым (1.0)."""
    values = [1.0 if value is None or math.isnan(value) else float(value)
              for value in p_values]
    order = sorted(range(len(values)), key=lambda index: values[index])
    adjusted = [1.0] * len(values)
    running = 0.0
    for rank, index in enumerate(order):
        value = min(1.0, values[index] * (len(values) - rank))
        running = max(running, value)
        adjusted[index] = running
    return adjusted


def evaluate(rows: list, start: int, embargo: int, models: dict) -> tuple:
    predictions = cm.walk_forward(rows, models, start, embargo)
    summary = {}
    for model in models:
        months, errors, group = _errors(predictions, model)
        if not group:
            summary[model] = {"n": 0}
            continue
        summary[model] = {
            "n": len(group),
            "rmse": cm.rmse_pp(group),
            "mae": cm.mae_pp(group),
            "months": [month.isoformat() for month in months],
        }

    pairs = {}
    primary_names = []
    for left, right in PRIMARY_PAIRS:
        if left not in models or right not in models:
            continue
        name = f"{left} vs {right}"
        pairs[name] = _pair(predictions, left, right, True)
        primary_names.append(name)
    for left, right in SECONDARY_PAIRS:
        if left not in models or right not in models:
            continue
        name = f"{left} vs {right}"
        pairs[name] = _pair(predictions, left, right, False)

    if primary_names:
        adjusted = holm_adjust([pairs[name].get("dm_p") for name in primary_names])
        for name, value in zip(primary_names, adjusted):
            pairs[name]["dm_p_holm"] = value
    return summary, pairs, predictions


def stage1_gate(
    sample: str,
    summary: dict,
    start: int = None,
    embargo: int = 0,
) -> tuple:
    """Регрессионный гейт: v2/v3 bottom-up воспроизводят числа Stage 1.

    Гейт применим только к окну Stage 1 (`start` по умолчанию для выборки,
    `embargo=0`); на сокращённом/отладочном окне он помечается `skipped`,
    иначе числа сравнивались бы с чужим числом фолдов. Возврат —
    `(gate, failed)`: расхождение не уничтожает уже посчитанные артефакты,
    решение о стопе принимает вызывающий.
    """
    if start is not None and (
        start != DEFAULT_START[sample] or embargo != 0
    ):
        return (
            {
                model: {"status": "skipped"}
                for model in STAGE1_REFERENCE[sample]
            },
            [],
        )
    gate = {}
    failed = []
    for model, expected in STAGE1_REFERENCE[sample].items():
        item = summary.get(model)
        if not item or not item.get("n"):
            gate[model] = {"status": "missing"}
            continue
        actual = item["rmse"]
        delta = actual - expected
        ok = abs(delta) <= REGRESSION_TOLERANCE
        gate[model] = {
            "status": "ok" if ok else "differs",
            "expected": expected,
            "actual": actual,
            "delta": delta,
        }
        if not ok:
            failed.append(model)
    return gate, failed


def _significant(item: dict) -> bool:
    if not item.get("n") or item.get("dm_p") is None:
        return False
    low, high = item["delta_ci"]
    return item["dm_p"] < 0.05 or low > 0.0 or high < 0.0


def decide(pairs: dict) -> dict:
    """Вердикт по pre-registered правилу спеки §8."""
    primary = {
        name: item
        for name, item in pairs.items()
        if item.get("primary") and item.get("n")
    }
    if len(primary) < 2:
        return {
            "verdict": "iterate",
            "reasons": ["primary-контрасты не посчитаны полностью"],
        }
    improved = [
        name for name, item in primary.items()
        if item["delta"] < 0.0 and _significant(item)
    ]
    degraded = [
        name for name, item in primary.items()
        if item["delta"] > 0.0 and _significant(item)
    ]
    both_not_worse = all(item["delta"] <= 0.0 for item in primary.values())
    mapping = pairs.get("v3_ind_all_bottomup vs v3_ind_bottomup", {})
    mapping_ok = not mapping.get("n") or mapping.get("delta", 0.0) <= 0.0
    reasons = []
    if improved:
        reasons.append("значимо лучше: " + ", ".join(improved))
    if degraded:
        reasons.append("значимо хуже: " + ", ".join(degraded))
    if not mapping_ok:
        reasons.append("mapping-all вариант хуже mapped")
    if both_not_worse and improved and not degraded and mapping_ok:
        return {"verdict": "promote", "reasons": reasons}
    if degraded and not improved:
        return {"verdict": "reject", "reasons": reasons}
    if not improved:
        reasons.append("значимого улучшения нет")
    return {"verdict": "iterate", "reasons": reasons}


# ---------------------------------------------------------------------------
# Артефакты
# ---------------------------------------------------------------------------


def write_report(path: Path, payload: dict) -> None:
    lines = ["# CPI monthly industry: отраслевые МП-признаки — отчёт прогона", ""]
    lines.append(
        "> `scripts/cpi_monthly_industry.py`; пункт "
        "`key-rate/phase-4-ml-components` (issue #235), Stage 2. "
        "Числа не редактируются руками."
    )
    lines.append("")
    lines.append(
        f"- Выборка: `{payload['sample']}`; строк: {payload['n_rows']}; "
        f"start={payload['start']}; embargo={payload['embargo']}"
    )
    lines.append(f"- Пропущено: {payload['n_skipped']}")
    if payload.get("run_id"):
        lines.append(f"- run_id: `{payload['run_id']}`")
    lines.append(
        f"- МП: {payload['mp_survey']['n_sheets']} листов, "
        f"{payload['mp_survey']['n_months']} месяцев, "
        f"{payload['mp_survey']['min_date']}…{payload['mp_survey']['max_date']}; "
        f"sha256 `{payload['mp_survey']['sha256']}`"
    )
    lines.append("")
    lines.append("## Модели")
    lines.append("")
    lines.append("| Модель | n | RMSE | MAE |")
    lines.append("|---|---:|---:|---:|")
    for model, item in payload["summary"].items():
        if not item.get("n"):
            lines.append(f"| `{model}` | 0 | — | — |")
            continue
        lines.append(
            f"| `{model}` | {item['n']} | {item['rmse']:.6f} | "
            f"{item['mae']:.6f} |"
        )
    lines.append("")
    lines.append("## Парные сравнения")
    lines.append("")
    lines.append(
        "| Пара | primary | n | RMSE left | RMSE right | Δ | 95 % CI Δ | "
        "DM p | Holm p | MDE (ratio) |"
    )
    lines.append("|---|---|---:|---:|---:|---:|---|---:|---:|---:|")
    for name, item in payload["pairs"].items():
        if not item.get("n"):
            lines.append(f"| `{name}` | {'да' if item.get('primary') else 'нет'} | 0 | — | — | — | — | — | — | — |")
            continue
        holm = item.get("dm_p_holm")
        lines.append(
            f"| `{name}` | {'да' if item.get('primary') else 'нет'} | "
            f"{item['n']} | {item['rmse_left']:.6f} | "
            f"{item['rmse_right']:.6f} | {item['delta']:+.6f} | "
            f"[{item['delta_ci'][0]:+.6f}; {item['delta_ci'][1]:+.6f}] | "
            f"{item['dm_p']:.4f} | "
            f"{'—' if holm is None else f'{holm:.4f}'} | "
            f"{item['mde_ratio']:.4f} |"
        )
    lines.append("")
    lines.append("## Регрессионный гейт Stage 1")
    lines.append("")
    lines.append("| Модель | Ожидалось | Получено | Δ | Статус |")
    lines.append("|---|---:|---:|---:|---|")
    for model, item in payload["regression_gate"].items():
        if item.get("status") in ("missing", "skipped"):
            lines.append(
                f"| `{model}` | — | — | — | {item['status']} |"
            )
            continue
        lines.append(
            f"| `{model}` | {item['expected']:.6f} | {item['actual']:.6f} | "
            f"{item['delta']:+.9f} | {item['status']} |"
        )
    lines.append("")
    lines.append("## Вердикт")
    lines.append("")
    lines.append(f"`{payload['decision']['verdict']}`")
    for reason in payload["decision"]["reasons"]:
        lines.append(f"- {reason}")
    lines.append("")
    lines.append("## Что не проверено")
    lines.append("")
    lines.append(
        "- Mapping отрасль → компонент — pre-registered prior, не публичные "
        "веса; robustness-вариант `all` назван отдельно."
    )
    lines.append(
        "- Используется SA-блок; raw-блок в этом пункте не проверялся."
    )
    lines.append(
        "- Нет винтажного архива и фактических дат публикации: as-of — "
        "правило (конец P+1), а не PIT-слой."
    )
    lines.append(
        "- CB_COMP напрямую не сравнивается; h2/h3/h6 не считаются."
    )
    lines.append(
        "- Production-импорт 16 листов в ClickHouse — отдельный Go-пункт "
        "Task 6."
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def write_json(path: Path, payload: dict) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Прогон
# ---------------------------------------------------------------------------


def v3_sample_hash(sample: str) -> str:
    """Hash v3-выборки (без industry) для сверки с паспортом Stage 1.

    Stage 1 писал `sample_hash` по `{month, y, features}` с признаками
    `cpi_monthly_v3.feature_names_v3()`; здесь те же поля и порядок
    канонического JSON, поэтому hash сравним с `v_ml_runs.sample_hash`.
    """
    rows, _ = cmv3.build_sample_v3(sample)
    names = cmv3.feature_names_v3()
    payload = [
        {
            "month": row.month.isoformat(),
            "y": mlx.json_safe(row.y),
            "features": mlx.json_safe(
                {name: row.features[name] for name in names}
            ),
        }
        for row in rows
    ]
    return mlx.sha256_text(mlx.canonical_json(payload))


def run(
    sample: str,
    out_dir: Path,
    start: int = None,
    embargo: int = 0,
    models: list = None,
    mp_survey: str = None,
    allow_data_revision: bool = False,
) -> dict:
    if sample not in SAMPLES:
        raise SystemExit(f"неизвестная выборка {sample!r}")
    start = DEFAULT_START[sample] if start is None else start
    selected_names = list(DEFAULT_MODELS[sample] if models is None else models)
    unknown = set(selected_names) - set(MODELS)
    if unknown:
        raise SystemExit(
            f"неизвестные модели: {sorted(unknown)}; доступны {list(MODELS)}"
        )
    selected = {name: MODELS[name] for name in selected_names}

    cache_dir = ROOT / "build" / "cpi-monthly-industry"
    mp_path, mp_source = resolve_mp_survey(mp_survey, cache_dir)
    panel = load_industry_panel(mp_path)
    rows, skipped = build_sample_industry(sample, panel)
    if not rows:
        raise SystemExit(f"пустая выборка {sample!r}")
    assert_no_industry_leakage(rows, panel)
    print(
        f"Выборка {sample}: {len(rows)} строк "
        f"({rows[0].month}…{rows[-1].month}), пропущено {len(skipped)}",
        flush=True,
    )

    summary, pairs, predictions = evaluate(rows, start, embargo, selected)
    gate, gate_failed = stage1_gate(
        sample, summary, start=start, embargo=embargo
    )
    decision = decide(pairs)
    payload = {
        "sample": sample,
        "n_rows": len(rows),
        "n_skipped": len(skipped),
        "start": start,
        "embargo": embargo,
        "models": selected_names,
        "seed": SEED,
        "mapping": {
            "components": {
                key: list(value) for key, value in COMPONENT_SHEETS.items()
            },
            "general": list(GENERAL_SHEETS),
        },
        "industry_features": industry_feature_names(),
        "mp_survey": panel_meta(panel, mp_path, mp_source),
        "summary": summary,
        "pairs": pairs,
        "regression_gate": gate,
        "gate_failed": bool(gate_failed),
        "allow_data_revision": bool(allow_data_revision),
        "decision": decision,
    }

    sample_dir = Path(out_dir) / sample
    sample_dir.mkdir(parents=True, exist_ok=True)
    report_path = sample_dir / "cpi_monthly_industry_report.md"
    json_path = sample_dir / "cpi_monthly_industry_result.json"
    write_report(report_path, payload)
    write_json(json_path, payload)

    if os.environ.get("ML_LOG") == "1":
        try:
            logger = mlx.start_run(
                experiment="cpi_monthly_industry",
                run_name=(
                    f"cpi_monthly_industry__{sample}__start{start}"
                    f"__embargo{embargo}"
                ),
                config={
                    "sample": sample,
                    "start": start,
                    "embargo": embargo,
                    "seed": SEED,
                    "models": selected_names,
                    "mapping": payload["mapping"],
                    "industry_features": payload["industry_features"],
                },
                data_asof=cm.month_end(rows[-1].month),
                sample_meta={
                    "rows": len(rows),
                    "first": rows[0].month.isoformat(),
                    "last": rows[-1].month.isoformat(),
                    "skipped": len(skipped),
                    "mp_survey_sha256": payload["mp_survey"]["sha256"],
                },
                sample=[
                    {
                        "month": row.month.isoformat(),
                        "y": mlx.json_safe(row.y),
                        "features": mlx.json_safe(row.features),
                    }
                    for row in rows
                ],
                spec_path=(
                    "docs/superpowers/specs/"
                    "2026-10-19-cpi-industry-mp-features-design.md"
                ),
                enabled=True,
            )
            logger.log_metrics(
                {
                    "sample": sample,
                    "summary": summary,
                    "pairs": pairs,
                    "regression_gate": gate,
                    "decision": decision,
                }
            )
            logger.log_predictions(
                [
                    mlx.PredictionRow(
                        model=item.model,
                        window="h1",
                        target_date=item.month,
                        predicted=item.y_pred,
                        actual=item.y_true,
                        unit="pp",
                        meta={"sample": sample},
                    )
                    for item in predictions
                ]
            )
            logger.log_artifacts([report_path, json_path])
            logger.finish("finished", f"sample={sample}; rows={len(rows)}")
            payload["run_id"] = logger.run_id
            write_json(json_path, payload)
            write_report(report_path, payload)
            print(f"MLflow run_id: {payload['run_id']}", flush=True)
        except Exception as exc:
            print(
                f"Журнал эксперимента не записан: {type(exc).__name__}: {exc}",
                flush=True,
            )
    else:
        logger = None

    print("Модели:", flush=True)
    for model, item in summary.items():
        if item.get("n"):
            print(
                f"  {model}: RMSE {item['rmse']:.6f}, "
                f"MAE {item['mae']:.6f}",
                flush=True,
            )
    print("Primary-пары:", flush=True)
    for name, item in pairs.items():
        if item.get("primary") and item.get("n"):
            print(
                f"  {name}: Δ {item['delta']:+.6f} "
                f"[{item['delta_ci'][0]:+.6f}; {item['delta_ci'][1]:+.6f}], "
                f"DM p={item['dm_p']:.4f}, Holm p={item['dm_p_holm']:.4f}, "
                f"MDE {item['mde_ratio']:.4f}",
                flush=True,
            )
    print(f"Вердикт: {decision['verdict']} ({'; '.join(decision['reasons'])})")
    print(f"Артефакты: {report_path}, {json_path}", flush=True)
    if gate_failed and not allow_data_revision:
        print(
            "ВНИМАНИЕ: регрессионный гейт Stage 1 не пройден: "
            + ", ".join(gate_failed)
            + "; артефакты записаны, для официального прогона разберите "
            "расхождение или передайте --allow-data-revision",
            flush=True,
        )
    return payload


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="CPI monthly industry: отраслевые МП-признаки (issue #235)"
    )
    parser.add_argument("--sample", choices=SAMPLES, default="long")
    parser.add_argument("--start", type=int, default=None)
    parser.add_argument("--embargo", type=int, default=0)
    parser.add_argument(
        "--models",
        default=None,
        help="через запятую; по умолчанию набор выборки (см. DEFAULT_MODELS)",
    )
    parser.add_argument(
        "--mp-survey",
        default=None,
        help="путь к mp_survey_data.xlsx; по умолчанию публичный URL/кэш",
    )
    parser.add_argument(
        "--out-dir",
        default=str(ROOT / "build" / "cpi-monthly-industry"),
    )
    parser.add_argument(
        "--allow-data-revision",
        action="store_true",
        help="не падать, если регрессионный гейт Stage 1 разошёлся",
    )
    parser.add_argument(
        "--sample-hash",
        action="store_true",
        help="напечатать hash v3-выборки (без industry) и выйти",
    )
    arguments = parser.parse_args(argv)
    if arguments.sample_hash:
        print(v3_sample_hash(arguments.sample))
        return 0
    payload = run(
        sample=arguments.sample,
        out_dir=Path(arguments.out_dir),
        start=arguments.start,
        embargo=arguments.embargo,
        models=arguments.models.split(",") if arguments.models else None,
        mp_survey=arguments.mp_survey,
        allow_data_revision=arguments.allow_data_revision,
    )
    if payload.get("gate_failed") and not arguments.allow_data_revision:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
