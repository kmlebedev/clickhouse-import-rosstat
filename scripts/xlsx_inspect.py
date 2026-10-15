#!/usr/bin/env python3
"""Обзор листов в выпусках ЦБ: подписи, даты, объём — чем выбирать набор рядов.

Зачем этот скрипт
-----------------
Объявленный набор рядов импортёра (`series_catalog`, §11.4 стандарта) должен
опираться на то, что в источнике реально есть, а не на догадку. Бюллетени ЦБ
«О чём говорят тренды» приходят одной книгой на 28–47 листов, и по имени листа
(`2`, `23`, `64`) невозможно понять, какой это ряд. Скрипт печатает по каждому
листу то, что нужно для решения:

* имя листа и значение `A1` — заголовок таблицы;
* строку шапки и подписи колонок (`A=Дата`, `B=Все товары и услуги`, …);
* первую и последнюю дату столбца-даты и число точек;
* сколько строк заполнено и сколько всего колонок.

Прогон по каталогу скачанных выпусков показывает, какие листы и подписи есть в
каких выпусках, — то есть ровно ту работу, которой стоило объявленное множество
рядов. Дальше выбранные листы урезаются в фикстуры через `scripts/xlsx_trim.py`.

Как определяется структура (эвристики, а не догадки о конкретном файле)
-----------------------------------------------------------------------
* **колонка-дата** — колонка `A`: шапка бюллетеней подписана «Дата»/«Период»,
  а данные начинаются со строки, где в `A` стоит дата;
* **шапка** — последняя заполненная строка перед первой строкой с датой в `A`
  (в выпуске 2021-06 подпись `A` пустая, поэтому «строка над данными» надёжнее
  поиска слова «Дата»);
* **точки** — строки, где в `A` стоит дата; **первая/последняя дата** — крайние
  такие значения;
* датой считается `datetime`/`date` из `openpyxl`, а также текст вида
  `ДД.ММ.ГГГГ` и `ГГГГ-ММ-ДД`.

Если дат в колонке `A` нет (справочный лист), скрипт честно это печатает и
показывает превью первых непустых строк, а не выдумывает шапку.

Требует `openpyxl` — он нужен здесь по существу (разбор книги), в отличие от
`scripts/xlsx_trim.py`, который живёт на стандартной библиотеке. Ставится
`pip install openpyxl`; конвейер импорта от `openpyxl` НЕ зависит, это ручной
инструмент исследования.

Запуск (из корня репозитория)
-----------------------------
    # один выпуск, все листы подробно
    python3 scripts/xlsx_inspect.py /tmp/dip/f_62343_bulletin_26-06.xlsx

    # каталог скачанных выпусков, коротко по каждому листу
    python3 scripts/xlsx_inspect.py /tmp/dip --brief

    # только интересные листы, по всем выпускам
    python3 scripts/xlsx_inspect.py /tmp/dip/*.xlsx --sheets 2,23

Прогон по пятидесяти выпускам — это минута-две: книга читается потоково, но
`openpyxl` разбирает её на Python.

Чего скрипт НЕ делает
---------------------
* не скачивает выпуски и не ходит в сеть;
* ничего не пишет: ни файлов, ни ClickHouse (только читает книги);
* не урезает книги — это `scripts/xlsx_trim.py`;
* не решает, какие листы объявлять рядами: он даёт факты для решения.
"""

import argparse
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

DATE_TEXT_RE = re.compile(r"^\s*(\d{2})\.(\d{2})\.(\d{4})\s*$")
ISO_TEXT_RE = re.compile(r"^\s*(\d{4})-(\d{2})-(\d{2})")
TITLE_LIMIT = 90
LABEL_LIMIT = 30
LABELS_MAX = 12
HEAD_ROWS = 40


def as_date(value):
    """Значение ячейки -> `ГГГГ-ММ-ДД` или None, если это не дата."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        found = ISO_TEXT_RE.match(value)
        if found:
            return "%s-%s-%s" % found.groups()
        found = DATE_TEXT_RE.match(value)
        if found:
            return "%s-%s-%s" % (found.group(3), found.group(2), found.group(1))
    return None


def clip(text, limit):
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) > limit:
        return text[: limit - 1] + "…"
    return text


def cell_text(value):
    """Значение ячейки -> короткая строка для печати."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return as_date(value) or str(value)
    if isinstance(value, float):
        return "%g" % value
    return clip(str(value), LABEL_LIMIT)


def column_letter(index):
    """0 -> A, 25 -> Z, 26 -> AA."""
    name = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        name = chr(ord("A") + rest) + name
    return name


def sheet_summary(ws):
    """Один потоковый проход по листу -> словарь фактов для печати."""
    facts = {
        "title": "",
        "header_row": None,
        "labels": [],
        "first_data_row": None,
        "last_data_row": None,
        "points": 0,
        "first_date": None,
        "last_date": None,
        "filled_rows": 0,
        "columns": 0,
        "n_rows": 0,
        "preview": [],
    }
    head = []
    previous_filled = None

    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        facts["n_rows"] = row_number
        if row_number <= HEAD_ROWS:
            head.append(row)
        filled = [value for value in row if value is not None and value != ""]
        if filled:
            facts["filled_rows"] += 1
            facts["columns"] = max(facts["columns"], len(row))
        marker = as_date(row[0]) if row else None
        if marker:
            if facts["points"] == 0:
                facts["header_row"] = previous_filled
                facts["first_data_row"] = row_number
            facts["points"] += 1
            facts["last_data_row"] = row_number
            facts["first_date"] = facts["first_date"] or marker
            facts["last_date"] = marker
        if filled:
            previous_filled = row_number

    # заголовок: A1, иначе первое непустое значение первого непустого ряда
    title = head[0][0] if head and head[0] else None
    if title is None or title == "":
        for row in head:
            for value in row if row else ():
                if value is not None and value != "":
                    title = value
                    break
            if title is not None and title != "":
                break
    facts["title"] = cell_text(title)

    if facts["first_data_row"]:
        header_row = facts["header_row"]
        if header_row and header_row <= len(head):
            for index, value in enumerate(head[header_row - 1]):
                if value is not None and value != "":
                    facts["labels"].append("%s=%s" % (column_letter(index), cell_text(value)))
    else:
        preview = []
        for number, row in enumerate(head[:6], 1):
            text = " | ".join(cell_text(v) for v in row if v is not None and v != "")
            if text:
                preview.append("r%d: %s" % (number, text))
        facts["preview"] = preview[:4]
    return facts


def print_sheet(sheet_name, facts, brief):
    """Печать фактов по листу: подробно или одной строкой (--brief)."""
    if brief:
        dates = "дат нет"
        if facts["points"]:
            dates = "%d точек %s..%s" % (facts["points"], facts["first_date"], facts["last_date"])
        print("  %-6s %-46s подписей %-3d %-34s строк %d"
              % (repr(sheet_name), clip(facts["title"], 46), len(facts["labels"]),
                 dates, facts["filled_rows"]))
        return

    print("  ── лист %s" % repr(sheet_name))
    if facts["title"]:
        print("     заголовок: «%s»" % clip(facts["title"], TITLE_LIMIT))
    else:
        print("     заголовок: (пусто)")
    if facts["labels"]:
        shown = facts["labels"][:LABELS_MAX]
        tail = ""
        if len(facts["labels"]) > LABELS_MAX:
            tail = " …(+%d)" % (len(facts["labels"]) - LABELS_MAX)
        print("     шапка r%d: %s%s" % (facts["header_row"], "  ".join(shown), tail))
    elif facts["points"]:
        print("     шапка: подписей не найдено (r%s)" % facts["header_row"])
    if facts["points"]:
        print("     данные: r%d..r%d, точек %d, даты %s … %s"
              % (facts["first_data_row"], facts["last_data_row"], facts["points"],
                 facts["first_date"], facts["last_date"]))
    else:
        print("     данных-времени нет: в колонке A ни одной даты (справочный лист?)")
        for line in facts["preview"]:
            print("     %s" % clip(line, TITLE_LIMIT + 40))
    print("     заполнено строк %d из %d, колонок до %d"
          % (facts["filled_rows"], facts["n_rows"], facts["columns"]))


def collect_files(targets):
    """Аргументы -> список xlsx: каталоги разворачиваются в *.xlsx (по алфавиту)."""
    files = []
    for target in targets:
        path = Path(target)
        if path.is_dir():
            files.extend(sorted(str(item) for item in path.glob("*.xlsx")))
        elif path.suffix.lower() == ".xlsx":
            files.append(str(path))
        else:
            print("ПРОПУСК (не xlsx): %s" % target)
    return files


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Обзор листов в XLSX-выпусках ЦБ (требуется openpyxl).",
    )
    parser.add_argument("targets", nargs="+", help="файлы .xlsx или каталоги с ними")
    parser.add_argument("--sheets", default=None,
                        help="только эти листы, через запятую (по умолчанию все)")
    parser.add_argument("--brief", action="store_true",
                        help="одна строка на лист — для прогона по десяткам выпусков")
    args = parser.parse_args(argv)

    try:
        import openpyxl
    except ImportError:
        print("ОШИБКА: нужен openpyxl — без него книгу не разобрать.\n"
              "        поставьте: pip install openpyxl\n"
              "        (конвейер импорта от openpyxl не зависит: это ручной инструмент)",
              file=sys.stderr)
        return 2

    wanted = None
    if args.sheets:
        wanted = [name.strip() for name in args.sheets.split(",") if name.strip()]

    files = collect_files(args.targets)
    if not files:
        print("ОШИБКА: не нашлось ни одного .xlsx", file=sys.stderr)
        return 2

    for path in files:
        try:
            book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            # обзорный инструмент: важен сам факт «не открылась», а не трасса
            print("=== %s: НЕ ОТКРЫЛСЯ: %s: %s" % (path, type(exc).__name__, exc))
            continue
        try:
            sheet_names = list(book.sheetnames)
            print("=== %s — листов %d, %d Б" % (path, len(sheet_names), os.path.getsize(path)))
            print("    листы: %s" % ", ".join(sheet_names))
            if wanted:
                absent = [name for name in wanted if name not in sheet_names]
                if absent:
                    print("    в этом выпуске нет листов: %s" % ", ".join(absent))
            for name in sheet_names:
                if wanted and name not in wanted:
                    continue
                print_sheet(name, sheet_summary(book[name]), args.brief)
        finally:
            book.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
