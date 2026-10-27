#!/usr/bin/env python3
"""Офлайн-тесты research-разбора отраслевого ИЦП."""

from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

import ppi_sa_industry as m


class WorkbookCase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    @staticmethod
    def ved_sheet(wb: Workbook, sheet: str, year: int, bad: object = None) -> None:
        ws = wb.create_sheet(sheet)
        ws.append([f"Индексы цен производителей в {year} г."])
        ws.append(["", "ОКВЭД 2", "январь", "февраль"])
        ws.append(["Производство пищевых продуктов", "10"])
        ws.append(["Российская Федерация", "", "100,5", bad if bad is not None else "…"])
        ws.append(["Производство напитков", "11"])
        ws.append(["Российская Федерация", "", "101,0", "-"])

    def ved_path(self, duplicate: bool = False, bad: object = None) -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        self.ved_sheet(wb, "8.1", 2026, bad)
        if duplicate:
            self.ved_sheet(wb, "7.1", 2026)
        path = self.root / "ved.xlsx"
        wb.save(path)
        return path

    def direct_ved_path(self) -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("3.1")
        ws.append(["Индексы цен производителей в 2021 г."])
        ws.append(["", "ОКВЭД2", "январь", "февраль"])
        ws.append(["Производство пищевых продуктов", 10, "100,5", "101,0"])
        ws.append(["Производство напитков", "11", "99,5", "…"])
        path = self.root / "ved-direct.xlsx"
        wb.save(path)
        return path

    def weights_path(self, values: tuple[object, object] = ("60,0", "40,0")) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "3"
        ws.append(["Удельный вес"])
        ws.append(["", "Код", "2022", "2023", "2024", "2025", "2026"])
        ws.append(["Пищевые продукты", "10", 1, 1, 1, 1, values[0]])
        ws.append(["Напитки", "11", 1, 1, 1, 1, values[1]])
        path = self.root / "weights.xlsx"
        wb.save(path)
        return path


class VedLoadTest(WorkbookCase):
    def test_loads_code_name_date_and_skips_missing(self) -> None:
        rows = m.load_ved_mm(self.ved_path())
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], {
            "date": date(2026, 1, 1), "code": "10",
            "name": "Производство пищевых продуктов", "value": 100.5,
        })
        self.assertEqual(rows[1]["code"], "11")
        self.assertNotIn(date(2026, 2, 1), {r["date"] for r in rows})

    def test_loads_pre2025_rows_without_region_row(self) -> None:
        rows = m.load_ved_mm(self.direct_ved_path())
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0], {
            "date": date(2021, 1, 1), "code": "10",
            "name": "Производство пищевых продуктов", "value": 100.5,
        })
        self.assertIn(date(2021, 2, 1), {r["date"] for r in rows if r["code"] == "10"})

    def test_rejects_nonnumeric_value(self) -> None:
        with self.assertRaisesRegex(ValueError, "нечислов"):
            m.load_ved_mm(self.ved_path(bad="ошибка"))

    def test_rejects_duplicate_date_code(self) -> None:
        with self.assertRaisesRegex(ValueError, "дубль"):
            m.load_ved_mm(self.ved_path(duplicate=True))


class WeightsTest(WorkbookCase):
    def test_loads_and_normalizes_weights(self) -> None:
        weights = m.load_weights(self.weights_path())
        self.assertEqual(weights[(2026, "10")], 60.0)
        self.assertEqual(m.weight_map(weights, ("10", "11"), 2026), {"10": 0.6, "11": 0.4})

    def test_rejects_zero_weight_sum(self) -> None:
        weights = m.load_weights(self.weights_path((0, 0)))
        with self.assertRaisesRegex(ValueError, "сумма весов"):
            m.weight_map(weights, ("10", "11"), 2026)

    def test_rejects_negative_weight(self) -> None:
        weights = m.load_weights(self.weights_path((-1, 2)))
        with self.assertRaisesRegex(ValueError, "отрицатель"):
            m.weight_map(weights, ("10", "11"), 2026)


class AuditTest(WorkbookCase):
    def test_candidates_exclude_total_and_report_coverage(self) -> None:
        rows = m.load_ved_mm(self.ved_path())
        rows.append({"date": date(2026, 1, 1), "code": "1323500.029.31", "name": "Промышленность", "value": 100.0})
        weights = m.load_weights(self.weights_path())
        candidates = m.audit_candidates(rows, weights)
        self.assertEqual([r["code"] for r in candidates], ["10", "11"])
        self.assertEqual(candidates[0]["from"], "2026-01-01")
        self.assertEqual(candidates[0]["to"], "2026-01-01")
        self.assertEqual(candidates[0]["n"], 1)
        self.assertTrue(candidates[0]["has_weights_2026"])

    def test_candidates_report_internal_gaps(self) -> None:
        rows = [
            {"date": date(2026, month, 1), "code": "10", "name": "Производство пищевых продуктов", "value": 100.0}
            for month in (1, 3)
        ]
        candidates = m.audit_candidates(rows, {})
        self.assertEqual(candidates[0].get("missing"), 1)

    def test_preregistered_variant_membership(self) -> None:
        nonfood = tuple(str(code) for code in range(13, 34) if code != 19)
        self.assertEqual(getattr(m, "VARIANT_CODES", {}), {
            "food": ("10", "11"),
            "nonfood": nonfood,
            "core": ("10", "11") + nonfood,
            "food_no_alcohol": ("10",),
            "core_no_alcohol": ("10",) + nonfood,
        })


if __name__ == "__main__":
    unittest.main(verbosity=2)
