#!/usr/bin/env python3
"""Аудит публичной панели ИПЦ: ряды, mapping на веса, покрытие, PIT, компаундирование.

Task 1 пункта `key-rate/phase-4-ml-components` (issue #234). Скрипт только
читает ClickHouse и публичные xlsx; в БД не пишет. Внутренняя отраслевая панель
ЦБ не используется: mapping строится по публичным весам `cpi_weights_tovar.xlsx`
и словарю `dictionary_cpi.xlsx` из архива CB_COMP.

Запуск — через `make cpi-public-audit` (env подключает Makefile), тесты —
`make cpi-public-audit-test`.
"""

import argparse
import bisect
import hashlib
import json
import math
import os
import re
import subprocess
import sys
import urllib.request
import warnings
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from extrapolation_aisc import ch_query

ARCHIVE_ROOT_NAME = "BusinessSurveyCBRCpiForecastingSourceCode"
MP_URL = "https://www.cbr.ru/Content/Document/File/135603/mp_survey_data.xlsx"

WEEKLY_AGGREGATE = "Индекс потребительских цен (недельная оценка)"

MONTHLY_SERIES = {
    "cpi_all": "Индексы потребительских цен на товары и услуги по Российской Федерации",
    "cpi_food": "Индексы потребительских цен на продовольственные товары по Российской Федерации",
    "cpi_nonfood": "Индексы потребительских цен на непродовольственные товары по Российской Федерации",
    "cpi_services": "Индексы потребительских цен на услуги по Российской Федерации",
}

WEEKLY_SQL = """
SELECT series, date, value
FROM v_rosstat_macro
WHERE series IN (
    SELECT series FROM v_series_catalog
    WHERE source = 'rosstat' AND frequency = 'W'
)
ORDER BY date, series
FORMAT JSON
"""

MONTHLY_SQL = """
SELECT series, date, value
FROM v_rosstat_macro
WHERE series IN (%s)
ORDER BY series, date
FORMAT JSON
"""

INDUSTRY_SHEETS = (
    "Промышленность",
    "Добыча",
    "Обработка",
    "Обработка_потребительские",
    "Обработка_инвестиционные",
    "Обработка_промежуточные",
    "ЭЭГП",
    "Водоснабжение",
    "Сельское хозяйство",
    "Строительство",
    "Торговля (45+46+47)",
    "Торговля авто (45)",
    "Торговля оптовая (46)",
    "Торговля розничная (47)",
    "Транспортировка и хранение",
    "Услуги",
)

MP_INDICATORS = {
    "ibk": ("индикатор бизнес-климата банка россии", ("(факт)", "(ожид)")),
    "price_expectations": ("ценовые ожидания предприятий", ()),
    "costs": ("издержки производства", ()),
}

WEIGHT_YEARS = tuple(range(2016, 2025))
DEFAULT_WEIGHTS_YEAR = 2024
COVERAGE_GATE = 0.50

SAMPLE_FIRST = date(2013, 10, 1)
SAMPLE_LAST = date(2026, 8, 1)
CUTOFF_DAY = 8
WEEKLY_PUBLICATION_LAG_DAYS = 2

UNIT_RE = re.compile(
    r"\b(шт|кг|л|мл|г|м3|м2|пара|пачка|рулон|день|сутки|поездка|услуга|"
    r"человека|таблеток|таблетки|драже)\b",
    re.IGNORECASE,
)

COMPONENT_KEYS = {
    "6": "food",
    "7": "nonfood",
    "9000": "services",
}


# ---------------------------------------------------------------------------
# Даты, коды, нормализация
# ---------------------------------------------------------------------------


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def month_start(value: date) -> date:
    return value.replace(day=1)


def shift_month(value: date, months: int) -> date:
    index = value.year * 12 + (value.month - 1) + months
    return date(index // 12, index % 12 + 1, 1)


def month_end(value: date) -> date:
    return shift_month(month_start(value), 1) - timedelta(days=1)


def cutoff_for(month: date) -> date:
    following = shift_month(month_start(month), 1)
    return date(following.year, following.month, CUTOFF_DAY)


def available_from_weekly(day: date) -> date:
    return day + timedelta(days=WEEKLY_PUBLICATION_LAG_DAYS)


def available_from_mes(day: date) -> date:
    return month_end(shift_month(month_start(day), 1))


def sql_quote(value: str) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def normalize_code(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if "->" in text:
        text = text.split("->")[-1].strip()
    return text


def normalize_name(name, mode: str = "extended") -> str:
    """Нормализация имени ряда/веса.

    raw — точное совпадение; basic — регистр/ё/пунктуация; extended — плюс
    скобочные уточнения и единицы. Порядок операций зафиксирован: он даёт
    118 совпадений и покрытие 39,299 % на весах 2024.
    """
    if mode == "raw":
        return str(name).strip()
    text = str(name).strip().lower().replace("ё", "е")
    if mode == "extended":
        text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    if mode == "extended":
        text = UNIT_RE.sub(" ", text)
    return " ".join(text.split())


# ---------------------------------------------------------------------------
# ClickHouse
# ---------------------------------------------------------------------------


def load_weekly() -> dict:
    rows = ch_query(WEEKLY_SQL)
    weekly = defaultdict(list)
    for row in rows:
        weekly[str(row["series"])].append(
            (parse_date(row["date"]), float(row["value"]))
        )
    if not weekly:
        raise RuntimeError("v_rosstat_macro: нет недельных рядов rosstat/W")
    return {series: sorted(points) for series, points in weekly.items()}


def load_monthly() -> dict:
    sql = MONTHLY_SQL % ", ".join(sql_quote(name) for name in MONTHLY_SERIES.values())
    rows = ch_query(sql)
    names = {name: key for key, name in MONTHLY_SERIES.items()}
    monthly = defaultdict(list)
    for row in rows:
        series = str(row["series"])
        key = names.get(series)
        if key is None:
            continue
        monthly[key].append(
            (parse_date(row["date"]), float(row["value"]))
        )
    missing = set(MONTHLY_SERIES) - set(monthly)
    if missing:
        raise RuntimeError(f"v_rosstat_macro: нет месячных рядов: {sorted(missing)}")
    return {key: sorted(points) for key, points in monthly.items()}


# ---------------------------------------------------------------------------
# xlsx: веса и словарь
# ---------------------------------------------------------------------------


def weights_path(archive: Path) -> Path:
    return (
        archive
        / "Models/CpiComponents/StfCpi_External_data/databases/weights/cpi_weights_tovar.xlsx"
    )


def dictionary_path(archive: Path) -> Path:
    return (
        archive
        / "Models/CpiComponents/StfCpi_External_data/dictionaries/dictionary_cpi.xlsx"
    )


def excel_rows(path: Path, sheet: str) -> list:
    if not path.exists():
        raise FileNotFoundError(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"{path}: нет листа {sheet!r}")
        return [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()


def column_index(header: list, *names: str) -> int:
    normalized = [str(value).strip() if value is not None else "" for value in header]
    for name in names:
        for index, value in enumerate(normalized):
            if value == name:
                return index
    raise ValueError(f"нет колонки {names!r} в заголовке {normalized!r}")


def load_weights(archive: Path, year: int) -> list:
    path = weights_path(archive)
    rows = excel_rows(path, str(year))
    if not rows:
        raise ValueError(f"{path}: лист {year} пуст")
    header = rows[0]
    i_name = column_index(header, "наименование_товара_услуги")
    i_weight = column_index(header, "вес_товара")
    i_code = column_index(header, "код_товара_услуги")
    i_okato = column_index(header, "окато_код")
    result = []
    for number, row in enumerate(rows[1:], start=2):
        if len(row) <= max(i_name, i_weight, i_code, i_okato):
            continue
        if row[i_code] is None:
            continue
        if normalize_code(row[i_okato]) != "643":
            continue
        try:
            weight = float(row[i_weight])
        except (TypeError, ValueError) as error:
            raise ValueError(
                f"{path}:{year}: строка {number}: вес {row[i_weight]!r} не число"
            ) from error
        result.append(
            {
                "code": normalize_code(row[i_code]),
                "name": str(row[i_name]).strip(),
                "weight": weight,
            }
        )
    if not result:
        raise ValueError(f"{path}: лист {year}: нет строк с окато_код=643")
    return result


def load_dictionary(archive: Path, year: int) -> dict:
    path = dictionary_path(archive)
    rows = excel_rows(path, str(year))
    if not rows:
        raise ValueError(f"{path}: лист {year} пуст")
    header = rows[0]
    i_code = column_index(header, "код_товара_услуги")
    i_name = column_index(header, "наименование_товара_услуги")
    i_level = column_index(header, "уровень")
    i_l1 = column_index(header, "уровень_1")
    i_l2 = column_index(header, "уровень_2")
    i_l3 = column_index(header, "уровень_3")
    result = {}
    for row in rows[1:]:
        if len(row) <= max(i_code, i_name, i_level, i_l1, i_l2, i_l3):
            continue
        code = normalize_code(row[i_code])
        if not code:
            continue
        try:
            level = int(row[i_level])
        except (TypeError, ValueError):
            continue
        result[code] = {
            "code": code,
            "name": str(row[i_name]).strip(),
            "level": level,
            "l1": normalize_code(row[i_l1]),
            "l2": normalize_code(row[i_l2]),
            "l3": normalize_code(row[i_l3]),
        }
    l3_names = {
        code: value["name"]
        for code, value in result.items()
        if value["level"] == 3
    }
    for value in result.values():
        value["group_name"] = l3_names.get(value["l3"], value["l3"])
        value["component_key"] = COMPONENT_KEYS.get(value["l2"], "other")
    return result


def weights_by_code(weights: list) -> dict:
    result = defaultdict(float)
    for item in weights:
        result[item["code"]] += item["weight"]
    return dict(result)


# ---------------------------------------------------------------------------
# Mapping и покрытие
# ---------------------------------------------------------------------------


def match_mode(products: list, weights: list, mode: str) -> tuple:
    index = defaultdict(list)
    for item in weights:
        index[normalize_name(item["name"], mode)].append(item)
    matched = []
    unmatched = []
    for series in products:
        entries = index.get(normalize_name(series, mode), [])
        if entries:
            for item in entries:
                matched.append({"series": series, "weight": item})
        else:
            unmatched.append(series)
    return matched, unmatched


def coverage_base(weights: list, dictionary: dict) -> tuple:
    by_code = weights_by_code(weights)
    leaf_codes = {
        code for code, value in dictionary.items() if value["level"] == 7
    }
    total = sum(by_code.get(code, 0.0) for code in leaf_codes)
    return by_code, leaf_codes, total


def coverage_for_codes(by_code: dict, leaf_codes: set, codes: set) -> dict:
    used = sorted(code for code in codes if code in leaf_codes)
    numerator = sum(by_code.get(code, 0.0) for code in used)
    denominator = sum(by_code.get(code, 0.0) for code in leaf_codes)
    return {
        "matched_codes": len(used),
        "numerator": numerator,
        "denominator": denominator,
        "coverage": numerator / denominator if denominator else None,
    }


def component_coverage(
    by_code: dict, dictionary: dict, codes: set
) -> dict:
    result = {}
    for component in sorted(set(COMPONENT_KEYS.values())):
        component_dict_codes = {
            code
            for code, value in dictionary.items()
            if value["level"] == 7 and value["component_key"] == component
        }
        denominator = sum(by_code.get(code, 0.0) for code in component_dict_codes)
        numerator = sum(
            by_code.get(code, 0.0)
            for code in codes
            if code in component_dict_codes
        )
        result[component] = {
            "matched_codes": sum(1 for code in codes if code in component_dict_codes),
            "numerator": numerator,
            "denominator": denominator,
            "coverage": numerator / denominator if denominator else None,
        }
    return result


def group_coverage(by_code: dict, dictionary: dict, codes: set) -> list:
    groups = defaultdict(lambda: {"codes": 0, "weight": 0.0})
    for code in codes:
        value = dictionary.get(code)
        if not value:
            continue
        key = value["l3"]
        groups[key]["codes"] += 1
        groups[key]["weight"] += by_code.get(code, 0.0)
        groups[key]["name"] = value["group_name"]
    return [
        {
            "code": code,
            "name": item.get("name", code),
            "matched_codes": item["codes"],
            "weight": item["weight"],
        }
        for code, item in sorted(
            groups.items(), key=lambda pair: (-pair[1]["weight"], pair[0])
        )
    ]


def build_mapping(
    archive: Path, products: list, weights: list, dictionary: dict
) -> dict:
    matched, unmatched = match_mode(products, weights, "extended")
    raw_matched, _ = match_mode(products, weights, "raw")
    basic_matched, _ = match_mode(products, weights, "basic")
    by_code, leaf_codes, total = coverage_base(weights, dictionary)
    codes = {item["weight"]["code"] for item in matched}
    rows = []
    for item in matched:
        weight = item["weight"]
        value = dictionary.get(weight["code"], {})
        rows.append(
            {
                "series": item["series"],
                "weight_code": weight["code"],
                "weight_name": weight["name"],
                "weight_2024": weight["weight"],
                "component": value.get("component_key", "other"),
                "component_code": value.get("l2", ""),
                "component_name": {
                    "6": "Продовольственные товары",
                    "7": "Непродовольственные товары",
                    "9000": "Услуги",
                }.get(value.get("l2", ""), ""),
                "group_code": value.get("l3", ""),
                "group_name": value.get("group_name", ""),
                "match_kind": "extended",
            }
        )
    rows.sort(key=lambda item: (item["component"], item["group_code"], item["series"]))
    unmatched_rows = [
        {
            "series": series,
            "reason": "нет совпадения с именем веса после extended-нормализации "
            "(дженерик/бренд/снятая позиция — поштучно не проверялось)",
        }
        for series in sorted(unmatched)
    ]
    coverage = coverage_for_codes(by_code, leaf_codes, codes)
    return {
        "normalizer": "raw < basic < extended",
        "raw_exact_series": len({item["series"] for item in raw_matched}),
        "basic_normalized_series": len({item["series"] for item in basic_matched}),
        "extended_normalized_series": len({item["series"] for item in matched}),
        "matched_weight_codes": len(codes),
        "matched_rows": rows,
        "unmatched": unmatched_rows,
        "coverage_2024": coverage,
        "coverage_by_component_2024": component_coverage(
            by_code, dictionary, codes
        ),
        "coverage_by_group_2024": group_coverage(by_code, dictionary, codes),
        "weights_total_2024": total,
    }


def coverage_by_year(
    archive: Path, matched_codes: set
) -> list:
    result = []
    for year in WEIGHT_YEARS:
        weights = load_weights(archive, year)
        dictionary = load_dictionary(archive, year)
        by_code, leaf_codes, total = coverage_base(weights, dictionary)
        coverage = coverage_for_codes(by_code, leaf_codes, matched_codes)
        result.append(
            {
                "year": year,
                "matched_codes": coverage["matched_codes"],
                "numerator": coverage["numerator"],
                "denominator": coverage["denominator"],
                "coverage": coverage["coverage"],
            }
        )
    return result


# ---------------------------------------------------------------------------
# Недельные и месячные инвентари
# ---------------------------------------------------------------------------


def series_summary(points: list) -> dict:
    dates = sorted(day for day, _ in points)
    gaps = [
        (right - left).days
        for left, right in zip(dates, dates[1:])
    ]
    return {
        "n": len(points),
        "min_date": dates[0].isoformat() if dates else None,
        "max_date": dates[-1].isoformat() if dates else None,
        "max_gap_days": max(gaps) if gaps else 0,
    }


def weekly_summary(weekly: dict) -> dict:
    products = {
        series: points
        for series, points in weekly.items()
        if series != WEEKLY_AGGREGATE
    }
    all_dates = [day for points in weekly.values() for day, _ in points]
    by_year = defaultdict(
        lambda: {"series": set(), "products": set(), "points": 0}
    )
    per_series = []
    for series, points in sorted(weekly.items()):
        summary = series_summary(points)
        summary["series"] = series
        summary["is_aggregate"] = series == WEEKLY_AGGREGATE
        per_series.append(summary)
        for day, _ in points:
            by_year[day.year]["series"].add(series)
            if series != WEEKLY_AGGREGATE:
                by_year[day.year]["products"].add(series)
            by_year[day.year]["points"] += 1
    max_date = max(all_dates) if all_dates else None
    return {
        "n_catalog_series": len(weekly),
        "n_product_series": len(products),
        "n_points": len(all_dates),
        "min_date": min(all_dates).isoformat() if all_dates else None,
        "max_date": max_date.isoformat() if max_date else None,
        "aggregate": series_summary(weekly.get(WEEKLY_AGGREGATE, [])),
        "by_year": [
            {
                "year": year,
                "n_series": len(value["series"]),
                "n_product_series": len(value["products"]),
                "n_points": value["points"],
            }
            for year, value in sorted(by_year.items())
        ],
        "per_series": per_series,
        "n_series_ending_before_global_max": sum(
            1
            for item in per_series
            if item["max_date"]
            and max_date
            and item["max_date"] < max_date.isoformat()
        ),
        "n_series_with_gap_gt_14": sum(
            1 for item in per_series if item["max_gap_days"] > 14
        ),
    }


def monthly_summary(monthly: dict) -> dict:
    result = []
    for key, series in MONTHLY_SERIES.items():
        points = monthly.get(key, [])
        summary = series_summary(points)
        months = sorted({month_start(day) for day, _ in points})
        missing = 0
        if months:
            cursor = months[0]
            known = set(months)
            while cursor <= months[-1]:
                if cursor not in known:
                    missing += 1
                cursor = shift_month(cursor, 1)
        summary.update(
            {
                "key": key,
                "series": series,
                "n_missing_months": missing,
            }
        )
        result.append(summary)
    return {
        "source_table": "rosstat_ipc_mes via v_rosstat_macro",
        "series": result,
    }


# ---------------------------------------------------------------------------
# Отраслевой МП
# ---------------------------------------------------------------------------


def _marker_positions(row: list) -> tuple:
    raw_from = -1
    sa_from = -1
    for index, value in enumerate(row):
        text = str(value).strip() if value is not None else ""
        if text == "Исходные данные" and raw_from < 0:
            raw_from = index
        if text == "Сезонно скорректированные данные" and sa_from < 0:
            sa_from = index
    return raw_from, sa_from


def _find_header(
    header: list, start: int, end: int, needle: str, exclude=()
) -> int:
    for index in range(start, min(end, len(header))):
        text = str(header[index]).strip().lower() if header[index] is not None else ""
        if needle in text and not any(item in text for item in exclude):
            return index
    raise ValueError(f"нет колонки {needle!r} в блоке {start}:{end}")


def _month_gaps(dates: list) -> int:
    months = sorted({month_start(day) for day in dates})
    if not months:
        return 0
    known = set(months)
    cursor = months[0]
    missing = 0
    while cursor <= months[-1]:
        if cursor not in known:
            missing += 1
        cursor = shift_month(cursor, 1)
    return missing


def _indicator_stats(rows: list, column: int) -> dict:
    values = []
    for row in rows:
        if column >= len(row) or row[0] is None:
            continue
        try:
            day = parse_date(row[0]) if not isinstance(row[0], (int, float)) else from_excel(row[0]).date()
        except (TypeError, ValueError, OverflowError):
            continue
        value = row[column]
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        if isinstance(value, float) and math.isnan(value):
            continue
        values.append((day, float(value)))
    dates = [day for day, _ in values]
    return {
        "n": len(values),
        "min_date": min(dates).isoformat() if dates else None,
        "max_date": max(dates).isoformat() if dates else None,
        "n_missing_months": _month_gaps(dates),
    }


def audit_mp_survey(path: Path) -> dict:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        industry = []
        for sheet in INDUSTRY_SHEETS:
            if sheet not in sheet_names:
                raise ValueError(f"{path}: нет отраслевого листа {sheet!r}")
            rows = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
            if len(rows) < 5:
                raise ValueError(f"{path}:{sheet}: меньше 5 строк")
            raw_from, sa_from = _marker_positions(rows[2])
            if raw_from < 0 or sa_from <= raw_from:
                raise ValueError(f"{path}:{sheet}: не найдены маркеры блоков")
            header = rows[3]
            indicators = {}
            for key, (needle, exclude) in MP_INDICATORS.items():
                raw_col = _find_header(header, raw_from, sa_from, needle, exclude)
                sa_col = _find_header(header, sa_from, len(header), needle, exclude)
                indicators[f"{key}_raw"] = _indicator_stats(rows[4:], raw_col)
                indicators[f"{key}_sa"] = _indicator_stats(rows[4:], sa_col)
            first = indicators["ibk_raw"]
            industry.append(
                {
                    "sheet": sheet,
                    "n_months": first["n"],
                    "min_date": first["min_date"],
                    "max_date": first["max_date"],
                    "n_missing_months": first["n_missing_months"],
                    "indicators": indicators,
                    "weight": "—",
                    "basis": "feature, not weighted",
                }
            )
        other = [sheet for sheet in sheet_names if sheet not in INDUSTRY_SHEETS]
        return {
            "url": MP_URL,
            "n_sheets": len(sheet_names),
            "n_industry_sheets": len(industry),
            "industry_sheets": industry,
            "other_sheets": other,
            "n_months": industry[0]["n_months"] if industry else 0,
        }
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# PIT/as-of
# ---------------------------------------------------------------------------


def sample_months(first: date, last: date) -> list:
    months = []
    cursor = month_start(first)
    while cursor <= month_start(last):
        months.append(cursor)
        cursor = shift_month(cursor, 1)
    return months


def audit_pit(weekly: dict, monthly: dict, months: list) -> dict:
    products = {
        series: points
        for series, points in weekly.items()
        if series != WEEKLY_AGGREGATE
    }
    weekly_dates = sorted({day for points in products.values() for day, _ in points})
    target_dates = {day for day, _ in monthly.get("cpi_all", [])}
    violations = []
    excluded_total = 0
    max_excluded = 0
    missing_targets = []
    for month in months:
        cutoff = cutoff_for(month)
        boundary = cutoff - timedelta(days=WEEKLY_PUBLICATION_LAG_DAYS)
        position = bisect.bisect_right(weekly_dates, boundary)
        used = weekly_dates[max(0, position - 4) : position]
        excluded = [
            day
            for day in weekly_dates
            if day <= cutoff and available_from_weekly(day) > cutoff
        ]
        excluded_total += len(excluded)
        max_excluded = max(max_excluded, len(excluded))
        used_ok = all(available_from_weekly(day) <= cutoff for day in used)
        monthly_ok = available_from_mes(month_end(month)) > cutoff
        target_ok = month_end(month) in target_dates
        if not target_ok:
            missing_targets.append(month.isoformat())
        if not used_ok or not monthly_ok:
            violations.append(
                {
                    "month": month.isoformat(),
                    "cutoff": cutoff.isoformat(),
                    "used_weekly_ok": used_ok,
                    "monthly_availability_ok": monthly_ok,
                }
            )
    return {
        "cutoff_rule": f"{CUTOFF_DAY}-е число M+1",
        "weekly_available_from": f"date + {WEEKLY_PUBLICATION_LAG_DAYS}",
        "monthly_available_from": "конец M+1",
        "sample_first": months[0].isoformat() if months else None,
        "sample_last": months[-1].isoformat() if months else None,
        "n_months": len(months),
        "n_weekly_dates": len(weekly_dates),
        "n_excluded_weekly_points": excluded_total,
        "max_excluded_weekly_points_per_month": max_excluded,
        "n_missing_targets": len(missing_targets),
        "missing_targets": missing_targets,
        "n_violations": len(violations),
        "violations": violations,
    }


# ---------------------------------------------------------------------------
# Компаундирование недель
# ---------------------------------------------------------------------------


def compound_month(points: list) -> float:
    index = 1.0
    for _, value in points:
        index *= value / 100.0
    return index


def _stats(values: list) -> dict:
    if not values:
        return {"n": 0}
    absolute = [abs(value) for value in values]
    return {
        "n": len(values),
        "mae_pp": sum(absolute) / len(absolute),
        "rmse_pp": math.sqrt(sum(value * value for value in values) / len(values)),
        "max_abs_pp": max(absolute),
        "mean_pp": sum(values) / len(values),
    }


def audit_compounding(
    weekly: dict, monthly: dict, mapping: dict
) -> dict:
    aggregate = weekly.get(WEEKLY_AGGREGATE, [])
    by_month = defaultdict(list)
    for day, value in aggregate:
        by_month[month_start(day)].append((day, value))
    monthly_index = {
        month_start(day): value / 100.0
        for day, value in monthly.get("cpi_all", [])
    }
    common = sorted(set(by_month) & set(monthly_index))
    errors = [
        (compound_month(by_month[month]) - monthly_index[month]) * 100.0
        for month in common
    ]
    span = sample_months(common[0], common[-1]) if common else []
    known = set(common)
    missing_months = [month.isoformat() for month in span if month not in known]
    by_year = defaultdict(list)
    for month, error in zip(common, errors):
        by_year[month.year].append(error)
    aggregate_result = {
        "method": "Π(value/100) по неделям месяца",
        "n_months": len(common),
        "min_month": common[0].isoformat() if common else None,
        "max_month": common[-1].isoformat() if common else None,
        "n_missing_months_in_span": len(missing_months),
        "missing_months_in_span": missing_months,
        "n_weeks_min": min((len(by_month[m]) for m in common), default=0),
        "n_weeks_max": max((len(by_month[m]) for m in common), default=0),
        **_stats(errors),
        "by_year": [
            {"year": year, **_stats(values)}
            for year, values in sorted(by_year.items())
        ],
    }

    # Диагностика 2022+: агрегатного недельного ряда нет, поэтому месячный m/m
    # собирается только по 118 совпавшим товарным рядам с нормировкой весов.
    matched = {}
    for row in mapping["matched_rows"]:
        matched.setdefault(row["series"], []).append(row)
    multi = sum(1 for rows in matched.values() if len(rows) > 1)
    total_weight = sum(
        max(row["weight_2024"] for row in rows) for rows in matched.values()
    )
    diagnostic = []
    for month in sample_months(date(2022, 1, 1), date(2026, 8, 1)):
        weighted = 0.0
        used_weight = 0.0
        n_series = 0
        for series, rows in matched.items():
            points = [
                (day, value)
                for day, value in weekly.get(series, [])
                if month_start(day) == month
            ]
            if len(points) < 3:
                continue
            weight = max(row["weight_2024"] for row in rows)
            weighted += weight * (compound_month(points) - 1.0)
            used_weight += weight
            n_series += 1
        if used_weight <= 0:
            continue
        official = monthly_index.get(month)
        if official is None:
            continue
        diagnostic.append(
            {
                "month": month.isoformat(),
                "n_series": n_series,
                "matched_weight_coverage": used_weight / total_weight if total_weight else 0.0,
                "error_pp": (weighted / used_weight - (official - 1.0)) * 100.0,
            }
        )
    diagnostic_errors = [item["error_pp"] for item in diagnostic]
    matched_share = (mapping.get("coverage_2024") or {}).get("coverage") or 0.0
    mean_matched = (
        sum(item["matched_weight_coverage"] for item in diagnostic) / len(diagnostic)
        if diagnostic
        else None
    )
    return {
        "aggregate": aggregate_result,
        "product_panel_2022_plus": {
            "method": "coverage-limited diagnostic: Σ w·(Π(value/100)−1) / Σ w",
            "n_months": len(diagnostic),
            "n_series_multi_mapping": multi,
            "matched_weight_share_of_total": matched_share,
            "mean_matched_weight_coverage": mean_matched,
            "mean_total_weight_coverage": (
                mean_matched * matched_share if mean_matched is not None else None
            ),
            **_stats(diagnostic_errors),
            "by_month": diagnostic,
        },
    }


# ---------------------------------------------------------------------------
# Решение, архив, MP, JSON
# ---------------------------------------------------------------------------


def decision(coverage_2024: float, by_component: dict) -> dict:
    bottomup_full = coverage_2024 >= COVERAGE_GATE
    return {
        "coverage_gate": COVERAGE_GATE,
        "coverage_2024": coverage_2024,
        "bottomup_full": bottomup_full,
        "mode": "bottomup_full" if bottomup_full else "hybrid_4_components_plus_features",
        "full_88_groups": False,
        "partial_bottomup": "robustness_only" if not bottomup_full else "allowed",
        "reason": (
            f"покрытие весов {coverage_2024:.3%} "
            f"{'≥' if bottomup_full else '<'} гейта {COVERAGE_GATE:.0%}; "
            f"услуги {by_component.get('services', {}).get('coverage', 0):.1%}; "
            "полный 88-групповой bottom-up на публичных данных не заявляется"
        ),
    }


def resolve_archive(explicit: str = None) -> Path:
    candidates = []
    if explicit:
        candidates.append(Path(explicit))
    env = os.environ.get("CB_COMP_ARCHIVE")
    if env:
        candidates.append(Path(env))
    candidates.extend(
        [
            ROOT.parent
            / "cb-comp-archive"
            / "build"
            / "cb-comp"
            / "extracted"
            / ARCHIVE_ROOT_NAME,
            ROOT
            / "build"
            / "cb-comp"
            / "extracted"
            / ARCHIVE_ROOT_NAME,
            ROOT.parent / ARCHIVE_ROOT_NAME,
        ]
    )
    for candidate in candidates:
        path = candidate
        if path.name != ARCHIVE_ROOT_NAME:
            path = path / ARCHIVE_ROOT_NAME
        if weights_path(path).exists() and dictionary_path(path).exists():
            return path
    raise SystemExit(
        "архив CB_COMP не найден: передайте --archive PATH или CB_COMP_ARCHIVE; "
        "скачивание и распаковка — docs/research/2026-10-18-cb-comp-archive.md §2"
    )


def resolve_mp_survey(explicit: str, out_dir: Path) -> tuple:
    if explicit:
        path = Path(explicit)
        if not path.exists():
            raise FileNotFoundError(path)
        return path, "path"
    cached = out_dir / "mp_survey_data.xlsx"
    if cached.exists():
        return cached, "cache"
    out_dir.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(MP_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        if response.status != 200:
            raise RuntimeError(f"mp_survey: HTTP {response.status}")
        data = response.read()
    if not data.startswith(b"PK"):
        raise RuntimeError("mp_survey: ответ не похож на xlsx (нет PK-сигнатуры)")
    cached.write_bytes(data)
    return cached, "download"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def git_sha() -> str:
    try:
        return subprocess.check_output(
            ["git", "-C", str(ROOT), "rev-parse", "HEAD"], text=True
        ).strip()
    except (OSError, subprocess.CalledProcessError):
        return ""


def display_path(path: Path, base: Path) -> str:
    try:
        return str(Path(path).relative_to(base))
    except ValueError:
        return Path(path).name


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def build_payload(
    weekly: dict,
    monthly: dict,
    archive: Path,
    mp_path: Path,
    mp_source: str,
) -> dict:
    products = sorted(
        series for series in weekly if series != WEEKLY_AGGREGATE
    )
    weights = load_weights(archive, DEFAULT_WEIGHTS_YEAR)
    dictionary = load_dictionary(archive, DEFAULT_WEIGHTS_YEAR)
    mapping = build_mapping(archive, products, weights, dictionary)
    matched_codes = {
        row["weight_code"] for row in mapping["matched_rows"]
    }
    yearly = coverage_by_year(archive, matched_codes)
    mp = audit_mp_survey(mp_path)
    mp["sha256"] = sha256_file(mp_path)
    mp["source"] = mp_source
    pit = audit_pit(
        weekly,
        monthly,
        sample_months(SAMPLE_FIRST, SAMPLE_LAST),
    )
    compounding = audit_compounding(weekly, monthly, mapping)
    component_2024 = mapping["coverage_by_component_2024"]
    resolution = decision(
        mapping["coverage_2024"]["coverage"], component_2024
    )
    return {
        "schema": "cpi-public-data-audit/v1",
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "git_sha": git_sha(),
        "inputs": {
            "archive": archive.name,
            "weekly": {
                "sql": "v_rosstat_macro + v_series_catalog (rosstat/W)",
                "n_catalog_series": len(weekly),
                "n_product_series": len(products),
                "aggregate_series": WEEKLY_AGGREGATE,
            },
            "monthly": {
                "source": "rosstat_ipc_mes via v_rosstat_macro",
                "series": list(MONTHLY_SERIES.values()),
            },
            "weights": {
                "path": display_path(weights_path(archive), archive),
                "sha256": sha256_file(weights_path(archive)),
                "default_year": DEFAULT_WEIGHTS_YEAR,
                "years": list(WEIGHT_YEARS),
            },
            "dictionary": {
                "path": display_path(dictionary_path(archive), archive),
                "sha256": sha256_file(dictionary_path(archive)),
                "default_year": DEFAULT_WEIGHTS_YEAR,
            },
            "mp_survey": {
                "path": display_path(mp_path, ROOT),
                "source": mp_source,
                "url": MP_URL,
                "sha256": mp["sha256"],
            },
        },
        "weekly": weekly_summary(weekly),
        "monthly_components": monthly_summary(monthly),
        "mapping": mapping
        | {
            "coverage_by_year": yearly,
        },
        "mp_survey": mp,
        "pit": pit,
        "weekly_compounding": compounding,
        "decision": resolution,
    }


def parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Audit public CPI panel (issue #234)"
    )
    parser.add_argument("--archive", help="распакованный архив CB_COMP")
    parser.add_argument("--mp-survey", help="путь к mp_survey_data.xlsx")
    parser.add_argument(
        "--json",
        default=str(ROOT / "docs/research/2026-10-19-cpi-public-data-audit.json"),
        help="куда записать JSON",
    )
    parser.add_argument(
        "--out-dir",
        default=str(ROOT / "build/cpi-public-audit"),
        help="каталог кэша/артефактов",
    )
    return parser.parse_args(argv)


def main(argv=None) -> int:
    warnings.filterwarnings("ignore", category=UserWarning)
    args = parse_args(argv)
    archive = resolve_archive(args.archive)
    out_dir = Path(args.out_dir)
    mp_path, mp_source = resolve_mp_survey(args.mp_survey, out_dir)
    weekly = load_weekly()
    monthly = load_monthly()
    payload = build_payload(weekly, monthly, archive, mp_path, mp_source)
    json_path = Path(args.json)
    write_json(json_path, payload)
    mapping = payload["mapping"]
    coverage = mapping["coverage_2024"]["coverage"]
    print(
        f"W: {payload['weekly']['n_catalog_series']} "
        f"({payload['weekly']['n_product_series']} товарных + агрегат); "
        f"точек {payload['weekly']['n_points']}"
    )
    print(
        f"Mapping: {mapping['extended_normalized_series']} совпадений, "
        f"покрытие {coverage:.3%} (2024), "
        f"unmatched {len(mapping['unmatched'])}"
    )
    print(
        "Компоненты 2024: "
        + ", ".join(
            f"{key} {value['coverage']:.1%}"
            for key, value in mapping["coverage_by_component_2024"].items()
        )
    )
    print(
        f"PIT: {payload['pit']['n_violations']} нарушений; "
        f"компаундирование: {payload['weekly_compounding']['aggregate']['n_months']} мес., "
        f"MAE {payload['weekly_compounding']['aggregate']['mae_pp']:.4f} п.п."
    )
    print(
        f"Решение: {payload['decision']['mode']} "
        f"(bottomup_full={payload['decision']['bottomup_full']})"
    )
    print(f"JSON: {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
